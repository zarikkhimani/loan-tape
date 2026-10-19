"""Owned Excel writing and exact saved-file validation for reviewed text exports."""

from __future__ import annotations

import importlib
import json
import sys
import tempfile
import time
import zipfile
from pathlib import Path
from threading import Event
from typing import Any
from uuid import uuid4

from defusedxml.ElementTree import fromstring  # type: ignore[import-untyped]
from openpyxl import load_workbook  # type: ignore[import-untyped]

from loan_tape._date_json import canonical
from loan_tape.excel_process import start_excel, start_in_job
from loan_tape.inspection import InspectionError
from loan_tape.ooxml_safety import validate_ooxml_archive
from loan_tape.workbook_index import check_cancelled

MAX_EXPORT_BYTES = 256 * 1024 * 1024


def validate_payload(payload: dict[str, Any]) -> bytes:
    """Reject unrepresentable cells before starting Excel; never truncate evidence."""
    for sheet in payload["sheets"]:
        rows = sheet["rows"]
        if not 1 <= len(rows) <= 1_048_576:
            raise ValueError("Export exceeds the worksheet row limit.")
        width = max(map(len, rows))
        if not 1 <= width <= 16_384:
            raise ValueError("Export exceeds the worksheet column limit.")
        for row in rows:
            for value in row:
                if value is not None and type(value) not in (str, int, bool):
                    raise ValueError("Export values require exact literal text or bounded counts.")
                if isinstance(value, str) and (
                    len(value.encode("utf-16-le")) // 2 > 32_766 or value.count("\n") > 253
                ):
                    raise ValueError(
                        "A value exceeds Excel's literal-cell limit; no text was truncated."
                    )
    encoded = canonical(payload).encode("utf-8")
    if len(encoded) > MAX_EXPORT_BYTES:
        raise ValueError("Text export payload exceeds the 256 MiB limit.")
    return encoded


def validate_output(
    path: Path, payload: dict[str, Any], *, cancelled: Event | None = None
) -> tuple[tuple[str, int], ...]:
    check_cancelled(cancelled)
    if not path.is_file() or not 0 < path.stat().st_size <= MAX_EXPORT_BYTES:
        raise InspectionError("Excel output is absent, empty or exceeds the export limit.")
    try:
        archive = zipfile.ZipFile(path)
    except zipfile.BadZipFile as error:
        raise InspectionError("Excel output is not a valid workbook archive.") from error
    with archive:
        validate_ooxml_archive(
            archive, error_type=InspectionError, cancelled=lambda: check_cancelled(cancelled)
        )
        if any(
            name.startswith(("xl/externalLinks/", "xl/connections", "xl/queryTables/"))
            or name.endswith(".bin")
            for name in archive.namelist()
        ):
            raise InspectionError("Export contains unexpected linked or executable content.")
        for name in archive.namelist():
            if name.endswith(".rels"):
                relationships = fromstring(archive.read(name), forbid_dtd=True)
                if any(item.get("TargetMode", "").lower() == "external" for item in relationships):
                    raise InspectionError("Export contains an unexpected external link.")
    book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        if book.sheetnames != [s["name"] for s in payload["sheets"]]:
            raise InspectionError("Export worksheet order or population differs from the payload.")
        counts = []
        for sheet, spec in zip(book.worksheets, payload["sheets"], strict=True):
            rows = spec["rows"]
            width = max(map(len, rows))
            if (sheet.max_row, sheet.max_column) != (len(rows), width):
                raise InspectionError(f"{sheet.title} dimensions do not reconcile.")
            for row_number, (actual, expected) in enumerate(
                zip(sheet.iter_rows(), rows, strict=True), 1
            ):
                check_cancelled(cancelled)
                padded = expected + [None] * (width - len(expected))
                for cell, value in zip(actual, padded, strict=True):
                    # Empty/absent source distinctions are retained in Source Evidence.
                    wanted = None if value == "" else value
                    if cell.data_type in {"f", "e"}:
                        raise InspectionError("Export contains a formula or executable error cell.")
                    if cell.value != wanted or type(cell.value) is not type(wanted):
                        raise InspectionError(
                            f"{sheet.title}!{cell.coordinate} differs from its exact export value."
                        )
                    if isinstance(wanted, str) and cell.data_type not in {"s", "inlineStr"}:
                        raise InspectionError("Source-derived text did not remain literal text.")
                if row_number > 1_048_576:
                    raise InspectionError("Export contains excess rows.")
            counts.append((sheet.title, len(rows)))
        return tuple(counts)
    finally:
        book.close()


def write_excel(
    payload: dict[str, Any], staging: Path, *, cancelled: Event | None, timeout: float
) -> str:
    """Start only job-owned Excel/Python processes; never open a source workbook."""
    encoded = validate_payload(payload)
    workspace = Path.cwd() / ".artifacts" / "text-exports"
    workspace.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="write-", dir=workspace) as directory:
        work = Path(directory)
        payload_path, result_path = work / "payload.json", work / "result.json"
        payload_path.write_bytes(encoded)
        jobs = importlib.import_module("win32job")
        events = importlib.import_module("win32event")
        processes = importlib.import_module("win32process")
        job = jobs.CreateJobObject(None, "LoanTapeTextExport-" + uuid4().hex)
        excel_process = worker_process = None
        try:
            limits = jobs.QueryInformationJobObject(job, jobs.JobObjectExtendedLimitInformation)
            limits["BasicLimitInformation"]["LimitFlags"] |= jobs.JOB_OBJECT_LIMIT_KILL_ON_JOB_CLOSE
            jobs.SetInformationJobObject(job, jobs.JobObjectExtendedLimitInformation, limits)
            check_cancelled(cancelled)
            excel_process, excel_pid = start_excel(job)
            python = Path(sys.executable)
            if python.name.lower() == "pythonw.exe":
                python = python.with_name("python.exe")
            worker_process, _ = start_in_job(
                str(python),
                [
                    "-m",
                    "loan_tape.text_export_worker",
                    str(payload_path),
                    str(staging),
                    str(result_path),
                    str(excel_pid),
                ],
                job,
            )
            started = time.monotonic()
            while True:
                check_cancelled(cancelled)
                if time.monotonic() - started > timeout:
                    raise InspectionError("Text export timed out; no workbook was published.")
                if events.WaitForSingleObject(worker_process, 200) == 0:
                    break
            if processes.GetExitCodeProcess(worker_process):
                raise InspectionError("The owned Excel writer failed; no workbook was published.")
            result = json.loads(result_path.read_text(encoding="utf-8"))
            if "error" in result:
                raise InspectionError(f"Excel export failed: {result['error']}")
            if result.get("writer") != "xlwings" or not result.get("writer_version"):
                raise InspectionError("Excel writer did not return its expected identity.")
            return str(result["writer_version"])
        finally:
            try:
                jobs.TerminateJobObject(job, 0)
                for handle in (worker_process, excel_process):
                    if handle is not None and events.WaitForSingleObject(handle, 5000) != 0:
                        raise InspectionError("An owned text-export process did not shut down.")
            finally:
                for handle in (worker_process, excel_process):
                    if handle is not None:
                        handle.Close()
                job.Close()
