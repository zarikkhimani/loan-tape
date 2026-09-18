"""Bounded, read-only previews. No mapping, cleaning, or loan validation."""

from __future__ import annotations

import csv
import io
from contextlib import ExitStack
from dataclasses import dataclass
from datetime import date, datetime, time
from itertools import islice, zip_longest
from pathlib import Path
from threading import Event
from typing import Any
from zipfile import ZipFile

from openpyxl import load_workbook  # type: ignore[import-untyped]

from loan_tape.ooxml_safety import validate_ooxml_archive
from loan_tape.ranges import CellRange

MAX_ROWS = 20
MAX_COLUMNS = 10
ENCODINGS = {"Automatic": None, "UTF-8": "utf-8-sig", "Windows-1252": "cp1252", "UTF-16": "utf-16"}
DELIMITERS = {"Automatic": None, "Comma": ",", "Semicolon": ";", "Tab": "\t", "Pipe": "|"}


class InspectionError(ValueError):
    """The source cannot be previewed with the selected reader settings."""


def check_read_cancelled(cancelled: Event | None) -> None:
    if cancelled is not None and cancelled.is_set():
        raise InspectionError("Reading cancelled; no complete result was produced.")


class CancellableReader(io.BufferedReader):
    """Check cancellation during ZIP/XML loading, including before a row is yielded."""

    def __init__(self, path: Path, cancelled: Event | None) -> None:
        self.cancelled = cancelled
        check_read_cancelled(cancelled)
        super().__init__(io.FileIO(path, "r"))

    def read(self, size: int | None = -1) -> bytes:
        check_read_cancelled(self.cancelled)
        return super().read(size)

    def readinto(self, buffer: Any) -> int:
        check_read_cancelled(self.cancelled)
        return super().readinto(buffer)

    def read1(self, size: int = -1) -> bytes:
        check_read_cancelled(self.cancelled)
        return super().read1(size)


def preflight_ooxml(path: Path, cancelled: Event | None = None) -> None:
    """Apply the same package budget before any openpyxl reader sees a workbook."""
    with CancellableReader(path, cancelled) as source, ZipFile(source) as archive:
        validate_ooxml_archive(
            archive,
            error_type=InspectionError,
            cancelled=lambda: check_read_cancelled(cancelled),
        )


@dataclass(frozen=True)
class PreviewCell:
    text: str
    formula: str | None = None
    number_format: str | None = None
    source_path: str | None = None
    presence: str | None = None


@dataclass(frozen=True)
class FilePreview:
    rows: tuple[tuple[PreviewCell, ...], ...]
    column_count: int
    more_rows: bool
    more_columns: bool
    sheets: tuple[str, ...] = ()
    selected_sheet: str | None = None
    reading_note: str = ""
    start_row: int = 1
    start_column: int = 1
    column_labels: tuple[str, ...] = ()
    record_groups: tuple[str, ...] = ()
    selected_record_group: str | None = None
    record_count: int | None = None
    record_group_error: str | None = None
    field_count: int | None = None
    source_sha256: str | None = None
    record_group_paths: tuple[tuple[str, ...], ...] = ()


def read_preview(
    path: Path,
    *,
    sheet: str | None = None,
    encoding: str = "Automatic",
    delimiter: str = "Automatic",
    area: CellRange | None = None,
    cancelled: Event | None = None,
    xml_group: str | None = None,
    xml_start_row: int = 1,
    xml_start_column: int = 1,
    xml_group_path: tuple[str, ...] | None = None,
    xml_source_sha256: str | None = None,
) -> FilePreview:
    """Read a bounded page of a saved source; never write to it."""
    check_read_cancelled(cancelled)
    extension = path.suffix.lower()
    if extension in {".xls", ".xlsb"}:
        raise InspectionError(
            "Preview currently supports CSV, .xlsx, and .xlsm. "
            "For this workbook, save a separate .xlsx copy in Excel and add that copy."
        )
    if extension == ".csv":
        return _read_csv(path, encoding, delimiter, cancelled)
    if extension == ".xml":
        from loan_tape.xml_preview import read_xml_preview

        return read_xml_preview(
            path,
            xml_group,
            cancelled,
            start_row=xml_start_row,
            start_column=xml_start_column,
            group_path=xml_group_path,
            source_sha256=xml_source_sha256,
        )
    if extension in {".xlsx", ".xlsm"}:
        return _read_excel(path, sheet, area, cancelled)
    raise InspectionError("Choose a CSV, XML, .xlsx, or .xlsm file to inspect.")


def _read_csv(
    path: Path, encoding: str, delimiter: str, cancelled: Event | None = None
) -> FilePreview:
    if encoding not in ENCODINGS or delimiter not in DELIMITERS:
        raise InspectionError("Choose one of the listed CSV reading options.")
    codec = ENCODINGS[encoding]
    if codec is None:
        with path.open("rb") as source:
            prefix = source.read(4)
        codec = "utf-16" if prefix.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig"
    try:
        with io.TextIOWrapper(
            CancellableReader(path, cancelled), encoding=codec, newline=""
        ) as source:
            sample = source.read(65536)
            separator = DELIMITERS[delimiter]
            if separator is None:
                try:
                    separator = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
                except csv.Error as error:
                    if any(character in sample for character in ",;\t|"):
                        raise InspectionError(
                            "Could not identify the CSV separator. Choose it above and try again."
                        ) from error
                    separator = ","  # A single-column file has no separator to detect.
            source.seek(0)
            records = list(
                islice(csv.reader(source, delimiter=separator, strict=True), MAX_ROWS + 1)
            )
    except UnicodeError as error:
        raise InspectionError(
            "Could not read the CSV text with this encoding. "
            "Choose another encoding above, or save a separate CSV UTF-8 copy in Excel."
        ) from error
    except csv.Error as error:
        raise InspectionError(
            "Could not read the CSV structure. Check its separator, quoting, and field lengths."
        ) from error
    visible = records[:MAX_ROWS]
    width = max((len(row) for row in visible), default=0)
    label = next(name for name, value in DELIMITERS.items() if value == separator)
    return FilePreview(
        rows=tuple(tuple(PreviewCell(value) for value in row[:MAX_COLUMNS]) for row in visible),
        column_count=min(width, MAX_COLUMNS),
        more_rows=len(records) > MAX_ROWS,
        more_columns=width > MAX_COLUMNS,
        reading_note=f"{codec} · {label.lower()} separator · Values read as text.",
    )


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _read_excel(
    path: Path, sheet: str | None, area: CellRange | None = None, cancelled: Event | None = None
) -> FilePreview:
    # Keep formulas alongside saved results; never calculate, refresh, or save.
    preflight_ooxml(path, cancelled)
    with ExitStack() as stack:
        source = stack.enter_context(CancellableReader(path, cancelled))
        book = load_workbook(source, read_only=True, data_only=False, keep_links=False)
        stack.callback(book.close)
        sheets = tuple(str(worksheet.title) for worksheet in book.worksheets)
        if not sheets:
            raise InspectionError("This workbook contains no worksheets to preview.")
        selected = sheets[0] if sheet is None else sheet
        if selected not in sheets:
            raise InspectionError("This worksheet is no longer available. Reopen Inspect file.")
        cached_source = stack.enter_context(CancellableReader(path, cancelled))
        cached = load_workbook(cached_source, read_only=True, data_only=True, keep_links=False)
        stack.callback(cached.close)
        worksheet, cached_sheet = book[selected], cached[selected]
        # Some producers record incorrect dimensions; stream actual rows instead.
        worksheet.reset_dimensions()
        cached_sheet.reset_dimensions()
        rows: list[tuple[PreviewCell, ...]] = []
        width = 0
        options = (
            {}
            if area is None
            else {
                "min_row": area.min_row,
                "max_row": min(area.max_row, area.min_row + MAX_ROWS - 1),
                "min_col": area.min_column,
                "max_col": min(area.max_column, area.min_column + MAX_COLUMNS - 1),
            }
        )
        pairs = zip_longest(worksheet.iter_rows(**options), cached_sheet.iter_rows(**options))
        for source_row, cached_row in islice(pairs, MAX_ROWS + 1):
            check_read_cancelled(cancelled)
            if source_row is None or cached_row is None:
                raise InspectionError("The workbook changed while reading. Add it again.")
            width = max(width, len(source_row))
            cells: list[PreviewCell] = []
            for cell, saved in islice(zip(source_row, cached_row, strict=True), MAX_COLUMNS):
                value: object = cell.value
                formula: str | None = None
                if cell.data_type == "f":
                    formula = (
                        str(value)
                        if isinstance(value, str)
                        else str(getattr(value, "text", None) or "[Formula]")
                    )
                    value = saved.value if saved.value is not None else formula
                cells.append(
                    PreviewCell(
                        _text(value),
                        formula,
                        str(cell.number_format) if cell.number_format else None,
                    )
                )
            rows.append(tuple(cells[:MAX_COLUMNS]))
        visible_width = min(area.column_count if area else width, MAX_COLUMNS)
        if area:
            # Explicit ranges include their blank tail even after the stored XML ends.
            missing = min(area.row_count, MAX_ROWS) - len(rows)
            rows.extend((PreviewCell(""),) * visible_width for _ in range(max(0, missing)))
        return FilePreview(
            rows=tuple(
                row[:visible_width] + (PreviewCell(""),) * max(0, visible_width - len(row))
                for row in rows[:MAX_ROWS]
            ),
            column_count=visible_width,
            more_rows=area.row_count > MAX_ROWS if area else len(rows) > MAX_ROWS,
            more_columns=area.column_count > MAX_COLUMNS if area else width > MAX_COLUMNS,
            start_row=area.min_row if area else 1,
            start_column=area.min_column if area else 1,
            sheets=sheets,
            selected_sheet=selected,
            reading_note="Stored cell values; Excel formatting is not reproduced. "
            "Formulas show saved results when available, otherwise formula text. No recalculation.",
        )
