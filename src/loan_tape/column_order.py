"""Read-only ordering for every row and column in one saved data set."""

from __future__ import annotations

import sqlite3
import tempfile
import warnings
from collections.abc import Callable, Iterable
from contextlib import ExitStack
from dataclasses import dataclass
from itertools import zip_longest
from pathlib import Path
from threading import Event

from openpyxl import load_workbook  # type: ignore[import-untyped]

from loan_tape.column_profile import (
    MISSING_KINDS,
    Cell,
    ColumnExample,
    ColumnTable,
    _check_cancelled,
    _classify,
    _verify_source,
    limit_xml_storage,
)
from loan_tape.inspection import (
    CancellableReader,
    InspectionError,
    PreviewCell,
    _text,
    preflight_ooxml,
)
from loan_tape.selection import SavedTable
from loan_tape.session import AnalysisSession
from loan_tape.xml_selection import SavedXmlTable

ORDER_MODES = ("A–Z", "Z–A", "Ascending", "Descending")
ORDER_PAGE_SIZE = 50
ORDER_COLUMN_PAGE_SIZE = 10


def _text_compare(left: str, right: str) -> int:
    """Case-insensitive Unicode ordering with exact text as a stable tie-breaker."""
    folded_left, folded_right = left.casefold(), right.casefold()
    if folded_left != folded_right:
        return (folded_left > folded_right) - (folded_left < folded_right)
    return (left > right) - (left < right)


@dataclass(frozen=True)
class OrderedRow:
    """One source row with its ordering evidence and requested display cells."""

    source_row: int
    sort_cell: ColumnExample
    cells: tuple[PreviewCell, ...]


@dataclass(frozen=True)
class OrderedPage:
    """A bounded window over ordered rows and source columns."""

    rows: tuple[OrderedRow, ...]
    start_column: int
    column_labels: tuple[str, ...]


class OrderedDataSet:
    """A complete saved data set stored on disk and exposed through bounded pages."""

    def __init__(self, table: ColumnTable, sort_column: int) -> None:
        work = Path.cwd() / ".artifacts" / "column-order"
        work.mkdir(parents=True, exist_ok=True)
        self._temporary = tempfile.TemporaryDirectory(prefix="order-", dir=work)
        self._db = sqlite3.connect(
            Path(self._temporary.name) / "values.sqlite", check_same_thread=False
        )
        limit_xml_storage(self._db)
        self._db.execute("PRAGMA cache_size = -4096")
        self._db.execute("PRAGMA temp_store = FILE")
        self._db.create_collation("LOAN_TEXT", _text_compare)
        self._db.execute(
            """CREATE TABLE ordered_rows (
                source_row INTEGER PRIMARY KEY,
                kind TEXT NOT NULL,
                text_value TEXT NOT NULL,
                number_format TEXT NOT NULL,
                source_path TEXT,
                presence TEXT,
                missing INTEGER NOT NULL,
                numeric_value REAL,
                row_source_path TEXT
            )"""
        )
        self._db.execute(
            """CREATE TABLE ordered_cells (
                source_row INTEGER NOT NULL,
                source_column INTEGER NOT NULL,
                text_value TEXT NOT NULL,
                formula TEXT,
                number_format TEXT,
                source_path TEXT,
                presence TEXT,
                PRIMARY KEY (source_row, source_column)
            )"""
        )
        self.table = table
        self.sort_column = sort_column
        self.first_row = 1
        self.first_column = 1
        self.column_labels: tuple[str, ...] = ()
        self.has_headers = False
        self._total_rows = 0
        self._numeric_rows = 0
        self._filled_rows = 0
        self.closed = False

    def add_row(
        self,
        source_row: int,
        sort_cell: ColumnExample,
        cells: Iterable[tuple[int, PreviewCell]],
        *,
        row_source_path: str | None = None,
    ) -> None:
        if self.closed:
            raise InspectionError("The ordered data-set result is already closed.")
        try:
            self._db.execute(
                "INSERT INTO ordered_rows VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
                (
                    source_row,
                    sort_cell.kind,
                    sort_cell.text,
                    sort_cell.number_format,
                    sort_cell.source_path,
                    sort_cell.presence,
                    int(sort_cell.kind in MISSING_KINDS),
                    float(sort_cell.text) if sort_cell.kind == "Number" else None,
                    row_source_path,
                ),
            )
            self._db.executemany(
                "INSERT INTO ordered_cells VALUES (?, ?, ?, ?, ?, ?, ?)",
                (
                    (
                        source_row,
                        source_column,
                        cell.text,
                        cell.formula,
                        cell.number_format,
                        cell.source_path,
                        cell.presence,
                    )
                    for source_column, cell in cells
                    if _retain_cell(cell)
                ),
            )
        except (OverflowError, sqlite3.Error) as error:
            raise InspectionError(
                "Ordering storage failed; no complete ordered view was produced "
                "(512 MiB temporary-storage limit)."
            ) from error

    def finish(
        self,
        *,
        first_row: int,
        first_column: int,
        total_rows: int,
        column_labels: tuple[str, ...],
        has_headers: bool,
    ) -> OrderedDataSet:
        if self.closed:
            raise InspectionError("The ordered data-set result is already closed.")
        try:
            rows, numeric, missing = self._db.execute(
                """SELECT COUNT(*), COALESCE(SUM(kind = 'Number'), 0),
                COALESCE(SUM(missing), 0) FROM ordered_rows"""
            ).fetchone()
            if rows != total_rows:
                raise InspectionError(
                    "Ordered data-set coverage does not match the saved data set."
                )
            self._db.execute(
                "CREATE INDEX ordered_text ON ordered_rows(missing, text_value COLLATE LOAN_TEXT)"
            )
            self._db.execute("CREATE INDEX ordered_number ON ordered_rows(missing, numeric_value)")
            self._db.execute(
                "CREATE INDEX ordered_cell_window ON ordered_cells(source_column, source_row)"
            )
            self._db.commit()
        except sqlite3.Error as error:
            raise InspectionError(
                "Ordering storage failed; no complete ordered view was produced "
                "(512 MiB temporary-storage limit)."
            ) from error
        self.first_row = first_row
        self.first_column = first_column
        self.column_labels = column_labels
        self.has_headers = has_headers
        self._total_rows = int(rows)
        self._numeric_rows = int(numeric)
        self._filled_rows = int(rows - missing)
        return self

    @property
    def total_rows(self) -> int:
        return self._total_rows

    @property
    def total_columns(self) -> int:
        return len(self.column_labels)

    @property
    def numeric_rows(self) -> int:
        return 0 if self.closed else self._numeric_rows

    @property
    def filled_rows(self) -> int:
        return 0 if self.closed else self._filled_rows

    def page(
        self,
        mode: str,
        offset: int = 0,
        limit: int = ORDER_PAGE_SIZE,
        *,
        column_offset: int = 0,
        column_limit: int = ORDER_COLUMN_PAGE_SIZE,
    ) -> OrderedPage:
        if self.closed:
            raise InspectionError("The ordered data-set result is unavailable.")
        if mode not in ORDER_MODES:
            raise InspectionError("Choose one of the available ordering modes.")
        if (
            type(offset) is not int
            or type(limit) is not int
            or type(column_offset) is not int
            or type(column_limit) is not int
            or offset < 0
            or limit < 1
            or column_offset < 0
            or column_limit < 1
            or column_offset >= max(1, self.total_columns)
        ):
            raise InspectionError("The ordered page position is invalid.")
        if mode in {"A–Z", "Z–A"}:
            direction = "ASC" if mode == "A–Z" else "DESC"
            order = (
                f"missing ASC, text_value COLLATE LOAN_TEXT {direction}, "
                f"text_value {direction}, source_row ASC"
            )
        else:
            direction = "ASC" if mode == "Ascending" else "DESC"
            order = (
                "CASE WHEN kind = 'Number' THEN 0 WHEN missing = 1 THEN 2 ELSE 1 END ASC, "
                f"numeric_value {direction}, text_value COLLATE LOAN_TEXT {direction}, "
                "source_row ASC"
            )
        selected = tuple(
            self._db.execute(
                f"""SELECT source_row, kind, text_value, number_format, source_path,
                presence, row_source_path FROM ordered_rows
                ORDER BY {order} LIMIT ? OFFSET ?""",
                (limit, offset),
            )
        )
        start_column = self.first_column + column_offset
        end_column = min(
            self.first_column + self.total_columns - 1,
            start_column + column_limit - 1,
        )
        row_ids = [int(row[0]) for row in selected]
        by_location: dict[tuple[int, int], PreviewCell] = {}
        if row_ids:
            placeholders = ",".join("?" for _ in row_ids)
            try:
                cells = self._db.execute(
                    f"""SELECT source_row, source_column, text_value, formula,
                    number_format, source_path, presence FROM ordered_cells
                    WHERE source_row IN ({placeholders}) AND source_column BETWEEN ? AND ?""",
                    (*row_ids, start_column, end_column),
                )
                by_location = {
                    (row, source_column): PreviewCell(
                        text, formula, number_format, source_path, presence
                    )
                    for row, source_column, text, formula, number_format, source_path, presence in cells
                }
            except sqlite3.Error as error:
                raise InspectionError("The ordered page could not be read.") from error
        labels = self.column_labels[column_offset : column_offset + column_limit]
        rows: list[OrderedRow] = []
        for row, kind, text, number_format, source_path, presence, row_path in selected:
            values = tuple(
                by_location.get(
                    (row, source_column),
                    self._missing_cell(row_path, labels[source_column - start_column]),
                )
                for source_column in range(start_column, end_column + 1)
            )
            rows.append(
                OrderedRow(
                    row,
                    ColumnExample(row, kind, text, number_format, 1, source_path, presence),
                    values,
                )
            )
        return OrderedPage(tuple(rows), start_column, labels)

    def _missing_cell(self, row_path: str | None, label: str) -> PreviewCell:
        if isinstance(self.table, SavedXmlTable):
            return PreviewCell(
                "",
                source_path=f"{row_path} (absent field: {label})",
                presence="Absent",
            )
        return PreviewCell("", number_format="General")

    def close(self) -> None:
        if self.closed:
            return
        self.closed = True
        self._db.close()
        self._temporary.cleanup()


OrderedColumn = OrderedDataSet


def _retain_cell(cell: PreviewCell) -> bool:
    """Implicit plain blanks need no per-cell storage; all other evidence does."""
    return bool(
        cell.text
        or cell.formula is not None
        or cell.number_format not in {None, "General"}
        or cell.source_path is not None
        or cell.presence is not None
    )


def _preview_cell(source: Cell, saved: Cell) -> PreviewCell:
    value: object = source.value
    formula: str | None = None
    if source.data_type == "f":
        formula = (
            str(value)
            if isinstance(value, str)
            else str(getattr(value, "text", None) or "[Formula]")
        )
        saved_value = saved.value
        value = saved_value if saved_value is not None else formula
    number_format = source.number_format
    return PreviewCell(_text(value), formula, str(number_format) if number_format else None)


def _excel_order(
    path: Path,
    table: SavedTable,
    column: int,
    result: OrderedDataSet,
    cancelled: Event | None,
    progress: Callable[[str], None],
) -> OrderedDataSet:
    selected = table.selection
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise InspectionError("An Excel data set requires its .xlsx or .xlsm source file.")
    if (
        type(column) is not int
        or not selected.area.min_column <= column <= selected.area.max_column
    ):
        raise InspectionError("Choose a column inside the saved data set.")
    _check_cancelled(cancelled)
    progress("Checking the saved source…")
    _verify_source(path, selected.workbook_sha256, cancelled)
    first = selected.header_row + 1 if selected.header_row else selected.area.min_row
    last = selected.area.max_row
    total = max(0, last - first + 1)
    header_values = ["" for _ in range(selected.area.column_count)]
    rows_read = 0
    preflight_ooxml(path, cancelled)
    with ExitStack() as stack:
        source = stack.enter_context(CancellableReader(path, cancelled))
        book = load_workbook(source, read_only=True, data_only=False, keep_links=False)
        stack.callback(book.close)
        cached_source = stack.enter_context(CancellableReader(path, cancelled))
        cached = load_workbook(cached_source, read_only=True, data_only=True, keep_links=False)
        stack.callback(cached.close)
        if selected.sheet not in [sheet.title for sheet in book.worksheets]:
            raise InspectionError("The selected worksheet is unavailable. Reopen Inspect file.")
        worksheet, cached_sheet = book[selected.sheet], cached[selected.sheet]
        worksheet.reset_dimensions()
        cached_sheet.reset_dimensions()
        scan_first = selected.header_row or first
        progress(
            f"Reading all {total:,} rows × {selected.area.column_count:,} columns "
            "from the saved data set…"
        )
        pairs = zip_longest(
            worksheet.iter_rows(
                min_row=scan_first,
                max_row=last,
                min_col=selected.area.min_column,
                max_col=selected.area.max_column,
            ),
            cached_sheet.iter_rows(
                min_row=scan_first,
                max_row=last,
                min_col=selected.area.min_column,
                max_col=selected.area.max_column,
            ),
        )
        with warnings.catch_warnings():
            warnings.filterwarnings(
                "error", message="Cell .* marked as a date.*", category=UserWarning
            )
            try:
                for row_number, pair in enumerate(pairs, scan_first):
                    _check_cancelled(cancelled)
                    source_row, cached_row = pair
                    if source_row is None or cached_row is None:
                        raise InspectionError("The workbook changed while reading. Add it again.")
                    if len(source_row) != selected.area.column_count or len(cached_row) != len(
                        source_row
                    ):
                        raise InspectionError("The workbook data-set width changed while reading.")
                    preview = tuple(
                        _preview_cell(cell, saved)
                        for cell, saved in zip(source_row, cached_row, strict=True)
                    )
                    if row_number == selected.header_row:
                        header_values = [cell.text for cell in preview]
                        continue
                    sort_index = column - selected.area.min_column
                    sort_cell, _ = _classify(source_row[sort_index], row_number)
                    result.add_row(
                        row_number,
                        sort_cell,
                        zip(
                            range(selected.area.min_column, selected.area.max_column + 1),
                            preview,
                            strict=True,
                        ),
                    )
                    rows_read += 1
                    if rows_read % 10000 == 0:
                        progress(f"Read {rows_read:,} of {total:,} data rows…")
            except UserWarning as error:
                raise InspectionError(
                    "An Excel date could not be decoded without changing its value. "
                    "No complete ordered view was produced."
                ) from error
    if rows_read > total:
        raise InspectionError("The data-set reader returned too many rows.")
    for row in range(first + rows_read, last + 1):
        _check_cancelled(cancelled)
        result.add_row(row, ColumnExample(row, "Blank", "", "General"), ())
    progress("Verifying the source is unchanged…")
    _verify_source(path, selected.workbook_sha256, cancelled)
    labels = (
        tuple(header_values)
        if selected.header_row is not None
        else tuple("" for _ in range(selected.area.column_count))
    )
    return result.finish(
        first_row=first,
        first_column=selected.area.min_column,
        total_rows=total,
        column_labels=labels,
        has_headers=selected.header_row is not None,
    )


def _xml_order(
    path: Path,
    table: SavedXmlTable,
    column: int,
    result: OrderedDataSet,
    cancelled: Event | None,
    progress: Callable[[str], None],
) -> OrderedDataSet:
    from loan_tape.xml_column import XML_KINDS
    from loan_tape.xml_preview import _parse, _Target

    if path.suffix.lower() != ".xml":
        raise InspectionError("An XML data set requires its XML source file.")
    if type(column) is not int or not 1 <= column <= table.field_count:
        raise InspectionError("Choose a field inside the saved XML data set.")
    _check_cancelled(cancelled)
    progress("Checking the complete XML group and source identity…")
    discovery = _Target(cancelled, table.group_path, collect_preview=False)
    identity = _parse(path, discovery)
    if identity != table.source_sha256:
        raise InspectionError(
            "The saved XML source changed. Add it again before ordering the data set."
        )
    if discovery.problem:
        raise InspectionError(discovery.problem)
    candidates = [key for key, value in discovery.groups.items() if value.repeated]
    if not candidates:
        candidates = [next(iter(discovery.groups))]
    fields = tuple(discovery.fields)
    if (
        table.group_path not in candidates
        or discovery.records != table.record_count
        or len(fields) != table.field_count
    ):
        raise InspectionError(
            "The saved XML group no longer matches the analysis scope. Reopen its preview."
        )
    sort_field = fields[column - 1]
    rows_read = 0

    def visit(row: int, location: str, values: dict[str, PreviewCell]) -> None:
        nonlocal rows_read
        _check_cancelled(cancelled)
        if row != rows_read + 1:
            raise InspectionError("XML records were not delivered in source order.")
        source = values.get(
            sort_field,
            PreviewCell(
                "",
                source_path=f"{location} (absent field: {sort_field})",
                presence="Absent",
            ),
        )
        kind = (
            source.presence
            if source.presence != "Present"
            else ("Text" if source.text else "Empty text")
        )
        if kind not in XML_KINDS or source.source_path is None:
            raise InspectionError("XML field evidence is incomplete.")
        assert kind is not None
        result.add_row(
            row,
            ColumnExample(
                row,
                kind,
                source.text,
                "General",
                source_path=source.source_path,
                presence=source.presence,
            ),
            ((target.fields[key] + 1, cell) for key, cell in values.items()),
            row_source_path=location,
        )
        rows_read += 1
        if row % 10000 == 0:
            progress(f"Read {row:,} of {table.record_count:,} XML records…")

    progress(
        f"Reading all {table.record_count:,} records × {table.field_count:,} fields "
        "from the saved XML data set…"
    )
    target = _Target(
        cancelled,
        table.group_path,
        collect_preview=False,
        on_full_record=visit,
    )
    if _parse(path, target) != identity:
        raise InspectionError("The XML source changed during ordering. Add it again.")
    if target.problem:
        raise InspectionError(target.problem)
    if rows_read != table.record_count or tuple(target.fields) != fields:
        raise InspectionError("XML data-set coverage does not match the saved group.")
    _check_cancelled(cancelled)
    return result.finish(
        first_row=1,
        first_column=1,
        total_rows=table.record_count,
        column_labels=fields,
        has_headers=True,
    )


def order_data_set(
    path: Path,
    table: ColumnTable,
    column: int,
    session: AnalysisSession,
    *,
    cancelled: Event | None = None,
    progress: Callable[[str], None] = lambda message: None,
) -> OrderedDataSet:
    """Retain every source row and column while ordering by one explicit column."""
    del session
    result = OrderedDataSet(table, column)
    try:
        if isinstance(table, SavedXmlTable):
            return _xml_order(path, table, column, result, cancelled, progress)
        return _excel_order(path, table, column, result, cancelled, progress)
    except Exception:
        result.close()
        raise


def order_column(
    path: Path,
    table: ColumnTable,
    column: int,
    session: AnalysisSession,
    *,
    cancelled: Event | None = None,
    progress: Callable[[str], None] = lambda message: None,
) -> OrderedDataSet:
    """Compatibility wrapper for the original public function name."""
    return order_data_set(
        path,
        table,
        column,
        session,
        cancelled=cancelled,
        progress=progress,
    )
