"""Versioned, exact-header mapping profiles for explicitly reviewed Excel layouts."""

from __future__ import annotations

import hashlib
import re
from contextlib import ExitStack
from dataclasses import dataclass
from datetime import date, datetime, time
from functools import lru_cache
from importlib.resources import files
from pathlib import Path
from threading import Event

from openpyxl import load_workbook  # type: ignore[import-untyped]

from loan_tape.inspection import (
    CancellableReader,
    InspectionError,
    check_read_cancelled,
    preflight_ooxml,
)
from loan_tape.pack_format import DATA_TYPES, ID_PATTERN, VERSION_PATTERN, object_keys, parse_json
from loan_tape.selection import SavedTable

MAX_PROFILE_BYTES = 512 * 1024
MAX_PROFILE_COLUMNS = 16_384
PROFILE_ROLES = frozenset({"source", "calculated", "mixed"})
PROFILE_PRESENCE = frozenset({"required", "optional", "conditional", "unconfirmed"})
BUILTIN_PROFILE_RESOURCES = ("warehouse_model.json",)


class MappingProfileError(ValueError):
    """A mapping profile or source layout needs explicit review."""


def _text(value: object, name: str, *, maximum: int = 2_000) -> str:
    if not isinstance(value, str) or not value.strip() or len(value) > maximum:
        raise MappingProfileError(
            f"{name} must be non-empty text of at most {maximum:,} characters."
        )
    try:
        value.encode("utf-8")
    except UnicodeError as error:
        raise MappingProfileError(f"{name} contains an invalid Unicode character.") from error
    return value


def _optional_text(value: object, name: str, *, maximum: int = 2_000) -> str | None:
    return None if value is None else _text(value, name, maximum=maximum)


@dataclass(frozen=True)
class ProfileColumn:
    position: int
    key: str
    source_header: str
    occurrence: int
    expected_type: str
    role: str
    presence: str
    unit: str | None = None
    notes: str | None = None

    def __post_init__(self) -> None:
        if type(self.position) is not int or not 1 <= self.position <= MAX_PROFILE_COLUMNS:
            raise MappingProfileError("Profile column position is outside the Excel column limit.")
        if not isinstance(self.key, str) or not ID_PATTERN.fullmatch(self.key):
            raise MappingProfileError("Profile column key must be a lowercase stable identifier.")
        if not isinstance(self.source_header, str) or not self.source_header:
            raise MappingProfileError("Profile source header must retain non-empty source text.")
        if len(self.source_header) > 2_000:
            raise MappingProfileError("Profile source header exceeds 2,000 characters.")
        try:
            self.source_header.encode("utf-8")
        except UnicodeError as error:
            raise MappingProfileError(
                "Profile source header contains an invalid Unicode character."
            ) from error
        if type(self.occurrence) is not int or self.occurrence < 1:
            raise MappingProfileError("Header occurrence must be a positive whole number.")
        if self.expected_type not in DATA_TYPES:
            raise MappingProfileError("Profile column has an unsupported expected type.")
        if self.role not in PROFILE_ROLES:
            raise MappingProfileError("Profile column has an unsupported source/calculation role.")
        if self.presence not in PROFILE_PRESENCE:
            raise MappingProfileError("Profile column has an unsupported presence setting.")
        _optional_text(self.unit, "Profile column unit", maximum=200)
        _optional_text(self.notes, "Profile column notes", maximum=4_000)


@dataclass(frozen=True)
class MappingProfile:
    id: str
    version: str
    name: str
    description: str
    columns: tuple[ProfileColumn, ...]
    fingerprint: str

    def __post_init__(self) -> None:
        if not isinstance(self.id, str) or not ID_PATTERN.fullmatch(self.id):
            raise MappingProfileError("Mapping profile ID must be a lowercase stable identifier.")
        if not isinstance(self.version, str) or not VERSION_PATTERN.fullmatch(self.version):
            raise MappingProfileError("Mapping profile version must use major.minor.patch.")
        _text(self.name, "Mapping profile name", maximum=200)
        _text(self.description, "Mapping profile description", maximum=4_000)
        if not isinstance(self.columns, tuple) or not self.columns:
            raise MappingProfileError("Mapping profile must contain at least one column.")
        if len(self.columns) > MAX_PROFILE_COLUMNS:
            raise MappingProfileError("Mapping profile exceeds the Excel column limit.")
        if not re.fullmatch(r"[0-9a-f]{64}", self.fingerprint):
            raise MappingProfileError("Mapping profile fingerprint is invalid.")
        if [item.position for item in self.columns] != list(range(1, len(self.columns) + 1)):
            raise MappingProfileError("Mapping profile positions must be contiguous and ordered.")
        keys = [item.key for item in self.columns]
        if len(keys) != len(set(keys)):
            raise MappingProfileError("Mapping profile column keys must be unique.")
        occurrences: dict[str, int] = {}
        for item in self.columns:
            occurrences[item.source_header] = occurrences.get(item.source_header, 0) + 1
            if item.occurrence != occurrences[item.source_header]:
                raise MappingProfileError(
                    f"Header {item.source_header!r} has an invalid occurrence number."
                )


@dataclass(frozen=True)
class ProfileBinding:
    source_column: int
    definition: ProfileColumn

    def __post_init__(self) -> None:
        if type(self.source_column) is not int or not 1 <= self.source_column <= 16_384:
            raise MappingProfileError("Profile binding falls outside the Excel column limit.")


@dataclass(frozen=True)
class MappingProfileMatch:
    profile: MappingProfile
    bindings: tuple[ProfileBinding, ...]

    @property
    def description(self) -> str:
        return (
            f"Mapping profile: {self.profile.name} {self.profile.version} · "
            f"{len(self.bindings):,} columns matched exactly."
        )


def load_mapping_profile(data: bytes, location: str = "mapping profile") -> MappingProfile:
    """Load one strict data-only profile and fingerprint its exact reviewed bytes."""
    if len(data) > MAX_PROFILE_BYTES:
        raise MappingProfileError(f"{location}: exceeds the {MAX_PROFILE_BYTES:,}-byte limit.")
    try:
        document = parse_json(data.decode("utf-8"), location)
        row = object_keys(
            document,
            {"format_version", "id", "version", "name", "description", "columns"},
            set(),
            location,
        )
        if type(row["format_version"]) is not int or row["format_version"] != 1:
            raise MappingProfileError(f"{location}: unsupported format_version; expected 1.")
        if not isinstance(row["columns"], list):
            raise MappingProfileError(f"{location}: columns must be a list.")
        if len(row["columns"]) > MAX_PROFILE_COLUMNS:
            raise MappingProfileError(f"{location}: too many columns.")
        columns = []
        for index, value in enumerate(row["columns"]):
            item = object_keys(
                value,
                {
                    "position",
                    "key",
                    "source_header",
                    "occurrence",
                    "expected_type",
                    "role",
                    "presence",
                },
                {"unit", "notes"},
                f"{location}.columns[{index}]",
            )
            columns.append(
                ProfileColumn(
                    item["position"],
                    item["key"],
                    item["source_header"],
                    item["occurrence"],
                    item["expected_type"],
                    item["role"],
                    item["presence"],
                    item.get("unit"),
                    item.get("notes"),
                )
            )
        return MappingProfile(
            row["id"],
            row["version"],
            row["name"],
            row["description"],
            tuple(columns),
            hashlib.sha256(data).hexdigest(),
        )
    except MappingProfileError:
        raise
    except (UnicodeError, ValueError, KeyError, TypeError, AttributeError) as error:
        raise MappingProfileError(f"{location}: {error}") from error


@lru_cache(maxsize=1)
def load_builtin_profiles() -> tuple[MappingProfile, ...]:
    root = files("loan_tape").joinpath("profile_data")
    profiles = tuple(
        load_mapping_profile(root.joinpath(name).read_bytes(), f"built-in profile {name}")
        for name in BUILTIN_PROFILE_RESOURCES
    )
    ids = [item.id for item in profiles]
    if len(ids) != len(set(ids)):
        raise MappingProfileError("Built-in mapping profile IDs must be unique.")
    return profiles


def match_headers(
    profile: MappingProfile, headers: tuple[str, ...], *, start_column: int = 1
) -> MappingProfileMatch | None:
    """Require the complete exact raw-header sequence; never normalize or guess."""
    if type(start_column) is not int or start_column < 1:
        raise MappingProfileError("Profile matching needs a valid first source column.")
    if start_column + len(headers) - 1 > 16_384:
        raise MappingProfileError("Profile matching range exceeds the Excel column limit.")
    expected = tuple(item.source_header for item in profile.columns)
    if headers != expected:
        return None
    return MappingProfileMatch(
        profile,
        tuple(ProfileBinding(start_column + item.position - 1, item) for item in profile.columns),
    )


def match_builtin_headers(
    headers: tuple[str, ...], *, start_column: int = 1
) -> MappingProfileMatch | None:
    matches = tuple(
        match
        for profile in load_builtin_profiles()
        if (match := match_headers(profile, headers, start_column=start_column)) is not None
    )
    if len(matches) > 1:
        raise MappingProfileError("More than one built-in profile matches this exact layout.")
    return matches[0] if matches else None


def _digest(path: Path, cancelled: Event | None) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        while chunk := source.read(1024 * 1024):
            check_read_cancelled(cancelled)
            digest.update(chunk)
    return digest.hexdigest()


def _header_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    expression = getattr(value, "text", None)
    return expression if isinstance(expression, str) else str(value)


def read_table_headers(
    path: Path, table: SavedTable, *, cancelled: Event | None = None
) -> tuple[str, ...]:
    """Read the complete designated header row without evaluating or saving the workbook."""
    selection = table.selection
    if selection.header_row is None:
        raise InspectionError("Save a header row before matching a mapping profile.")
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise InspectionError("Mapping profiles currently support saved .xlsx and .xlsm data sets.")
    check_read_cancelled(cancelled)
    before = _digest(path, cancelled)
    if before != selection.workbook_sha256:
        raise InspectionError(
            "The saved workbook has changed. Add it again before matching a mapping profile."
        )
    preflight_ooxml(path, cancelled)
    with ExitStack() as stack:
        source = stack.enter_context(CancellableReader(path, cancelled))
        book = load_workbook(source, read_only=True, data_only=False, keep_links=False)
        stack.callback(book.close)
        if selection.sheet not in book.sheetnames:
            raise InspectionError("The saved data set worksheet is no longer available.")
        worksheet = book[selection.sheet]
        worksheet.reset_dimensions()
        rows = worksheet.iter_rows(
            min_row=selection.header_row,
            max_row=selection.header_row,
            min_col=selection.area.min_column,
            max_col=selection.area.max_column,
        )
        row = next(rows, ())
        check_read_cancelled(cancelled)
        headers = tuple(_header_text(cell.value) for cell in row)
    if len(headers) != selection.area.column_count:
        raise InspectionError("The complete saved data set header could not be read.")
    if _digest(path, cancelled) != before:
        raise InspectionError("The saved workbook changed while matching its mapping profile.")
    return headers


def match_builtin_table_profile(
    path: Path, table: SavedTable, *, cancelled: Event | None = None
) -> MappingProfileMatch | None:
    if table.selection.header_row is None:
        return None
    headers = read_table_headers(path, table, cancelled=cancelled)
    return match_builtin_headers(headers, start_column=table.selection.area.min_column)
