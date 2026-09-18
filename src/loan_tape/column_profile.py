"""Read a complete saved Excel or XML column without changing source values."""

from __future__ import annotations

import hashlib
import json
import sqlite3
import tempfile
import warnings
from collections.abc import Callable
from contextlib import ExitStack
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from pathlib import Path
from threading import Event
from typing import Protocol

from openpyxl import load_workbook  # type: ignore[import-untyped]

from loan_tape.inspection import CancellableReader, InspectionError, preflight_ooxml
from loan_tape.ranges import column_label
from loan_tape.selection import SavedTable
from loan_tape.session import AnalysisSession
from loan_tape.xml_selection import SavedXmlTable

ColumnTable = SavedTable | SavedXmlTable
MISSING_KINDS = {"Blank", "Empty text", "Absent", "Empty element", "Explicit nil"}

KINDS = ("Blank", "Empty text", "Text", "Number", "Date/time", "Boolean", "Excel error", "Formula")
EXAMPLES_PER_TYPE = 3
REPEATED_LIMIT = 10
XML_RESULT_MAX_BYTES = 512 * 1024 * 1024


class ColumnCancelled(InspectionError):
    """Cancellation never returns a partial summary."""


class Cell(Protocol):
    value: object
    data_type: str
    number_format: str | None


@dataclass(frozen=True)
class ColumnExample:
    row: int
    kind: str
    text: str
    number_format: str
    count: int = 1
    source_path: str | None = None
    presence: str | None = None


@dataclass(frozen=True)
class ColumnProfile:
    table: ColumnTable
    column: int
    session: AnalysisSession
    header: str | None
    first_row: int
    last_row: int
    counts: tuple[tuple[str, int], ...]
    numeric_zeroes: int
    whitespace_text: int
    distinct_values: int
    repeated_values: int
    extra_occurrences: int
    examples: tuple[ColumnExample, ...]
    repeated: tuple[ColumnExample, ...]
    text_zeroes: int = 0

    @property
    def total_rows(self) -> int:
        return max(0, self.last_row - self.first_row + 1)

    @property
    def filled_cells(self) -> int:
        return self.total_rows - sum(n for kind, n in self.counts if kind in MISSING_KINDS)

    @property
    def data_address(self) -> str:
        if isinstance(self.table, SavedXmlTable):
            return f"field {self.column}, records {self.first_row:,}–{self.last_row:,}"
        label = column_label(self.column)
        return (
            f"{label}{self.first_row}:{label}{self.last_row}" if self.total_rows else "No data rows"
        )


def column_scope(table: ColumnTable, session: AnalysisSession) -> str:
    if isinstance(table, SavedXmlTable):
        return f"{table.name} · XML {table.group_label}\nAll {table.record_count:,} records · {session.date_label}"
    selected = table.selection
    first = selected.header_row + 1 if selected.header_row else selected.area.min_row
    scope = (
        f"Data rows {first:,}–{selected.area.max_row:,}"
        if first <= selected.area.max_row
        else "No data rows below the header"
    )
    return (
        f"{table.name} · {selected.sheet}!{selected.area.address}\n{scope} · {session.date_label}"
    )


def limit_xml_storage(db: sqlite3.Connection) -> None:
    """Bound each XML analysis database; source paths can be much longer than values."""
    db.execute("PRAGMA page_size = 4096")
    db.execute(f"PRAGMA max_page_count = {max(1, XML_RESULT_MAX_BYTES // 4096)}")


class ColumnAccumulator:
    """Shared exact counts; distinct source values stay in temporary SQLite."""

    def __init__(self, db: sqlite3.Connection, kinds: tuple[str, ...]) -> None:
        self.db = db
        self.counts = dict.fromkeys(kinds, 0)
        self.examples: list[ColumnExample] = []
        self.numeric_zeroes = self.text_zeroes = self.whitespace = 0
        db.execute("PRAGMA cache_size = -2048")
        db.execute("PRAGMA temp_store = FILE")
        db.execute("""CREATE TABLE frequency (
            kind TEXT, value_key TEXT, text TEXT, number_format TEXT,
            first_row INTEGER, count INTEGER, source_path TEXT, presence TEXT,
            PRIMARY KEY (kind, value_key))""")

    def add(self, cell: ColumnExample, key: str, *, numeric_zero: bool = False) -> None:
        self.counts[cell.kind] += 1
        if self.counts[cell.kind] <= EXAMPLES_PER_TYPE:
            self.examples.append(cell)
        self.numeric_zeroes += numeric_zero
        self.text_zeroes += cell.kind == "Text" and cell.text == "0"
        self.whitespace += cell.kind == "Text" and cell.text.isspace()
        if cell.kind not in MISSING_KINDS:
            self.db.execute(
                """INSERT INTO frequency VALUES (?, ?, ?, ?, ?, 1, ?, ?)
                ON CONFLICT(kind, value_key) DO UPDATE SET count = count + 1""",
                (
                    cell.kind,
                    key,
                    cell.text,
                    cell.number_format,
                    cell.row,
                    cell.source_path,
                    cell.presence,
                ),
            )

    def finish(self) -> tuple[int, int, int, tuple[ColumnExample, ...]]:
        distinct, repeated, extra = self.db.execute(
            """SELECT COUNT(*), COALESCE(SUM(count > 1), 0),
            COALESCE(SUM(count - 1), 0) FROM frequency"""
        ).fetchone()
        examples = tuple(
            ColumnExample(row, kind, text, fmt, count, source, presence)
            for kind, text, fmt, row, count, source, presence in self.db.execute(
                """SELECT kind, text, number_format, first_row, count, source_path, presence
                FROM frequency WHERE count > 1 ORDER BY count DESC, first_row LIMIT ?""",
                (REPEATED_LIMIT,),
            )
        )
        return distinct, repeated, extra, examples


def _check_cancelled(cancelled: Event | None) -> None:
    if cancelled is not None and cancelled.is_set():
        raise ColumnCancelled("Column inspection cancelled. No complete summary was produced.")


def _verify_source(path: Path, expected: str, cancelled: Event | None) -> None:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        while chunk := source.read(1024 * 1024):
            _check_cancelled(cancelled)
            digest.update(chunk)
    if digest.hexdigest() != expected:
        raise InspectionError(
            "The saved workbook has changed. Add it again before inspecting columns."
        )


def _classify(cell: Cell, row: int) -> tuple[ColumnExample, str]:
    """Group decoded source types; never trim text, parse dates, or coerce identifiers."""
    value = cell.value
    number_format = cell.number_format or "General"
    if cell.data_type == "f":
        expression = value if isinstance(value, str) else getattr(value, "text", None)
        if not isinstance(expression, str):
            raise InspectionError(f"The formula at row {row} has an unsupported representation.")
        kind, text = "Formula", expression
    elif cell.data_type == "e":
        kind, text = "Excel error", str(value)
    elif value is None:
        kind = "Empty text" if cell.data_type in {"s", "str", "inlineStr"} else "Blank"
        text = ""
    elif isinstance(value, str):
        kind, text = ("Empty text" if value == "" else "Text"), value
    elif isinstance(value, bool):
        kind, text = "Boolean", str(value)
    elif isinstance(value, (datetime, date, time)):
        kind, text = "Date/time", value.isoformat()
    elif isinstance(value, timedelta):
        kind, text = "Date/time", str(value)
    elif isinstance(value, (int, float)):
        kind, text = "Number", str(value)
    else:
        raise InspectionError(f"The stored value at row {row} has an unsupported type.")
    # Keep the reader's numeric representation and date/time subtype distinct.
    key = json.dumps([type(value).__name__, text], ensure_ascii=False)
    return ColumnExample(row, kind, text, number_format), key


def inspect_column(
    path: Path,
    table: ColumnTable,
    column: int,
    session: AnalysisSession,
    *,
    cancelled: Event | None = None,
    progress: Callable[[str], None] = lambda message: None,
    on_cell: Callable[[ColumnExample], None] | None = None,
) -> ColumnProfile:
    """Scan the full saved column, including hidden/blank rows and the final row."""
    if isinstance(table, SavedXmlTable):
        from loan_tape.xml_column import inspect_xml_column

        return inspect_xml_column(
            path, table, column, session, cancelled=cancelled, progress=progress, on_cell=on_cell
        )
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise InspectionError(
            "Column inspection currently supports saved .xlsx and .xlsm data sets."
        )
    selected = table.selection
    if (
        type(column) is not int
        or not selected.area.min_column <= column <= selected.area.max_column
    ):
        raise InspectionError("Choose a column inside the saved data set.")
    _check_cancelled(cancelled)
    progress("Checking the saved source…")
    _verify_source(path, selected.workbook_sha256, cancelled)
    first = selected.header_row + 1 if selected.header_row else selected.area.min_row
    last = selected.area.max_row
    total = max(0, last - first + 1)
    rows_read = 0
    header: str | None = "" if selected.header_row is not None else None
    work = Path.cwd() / ".artifacts" / "column-inspection"
    work.mkdir(parents=True, exist_ok=True)
    preflight_ooxml(path, cancelled)
    # Exact frequency counts live on disk so a million unique values do not fill RAM.
    with tempfile.TemporaryDirectory(prefix="column-", dir=work) as temporary, ExitStack() as stack:
        db = sqlite3.connect(Path(temporary) / "counts.sqlite")
        stack.callback(db.close)
        accumulator = ColumnAccumulator(db, KINDS)
        counts, examples = accumulator.counts, accumulator.examples
        source = stack.enter_context(CancellableReader(path, cancelled))
        book = load_workbook(source, read_only=True, data_only=False, keep_links=False)
        stack.callback(book.close)
        _check_cancelled(cancelled)
        if selected.sheet not in [sheet.title for sheet in book.worksheets]:
            raise InspectionError("The selected worksheet is unavailable. Reopen Inspect file.")
        sheet = book[selected.sheet]
        sheet.reset_dimensions()
        progress(f"Reading all {total:,} data rows in column {column_label(column)}…")
        # openpyxl otherwise changes an unrepresentable Excel date into an error value.
        # Treat that conversion warning as a read failure, not a source-cell error.
        with warnings.catch_warnings():
            warnings.filterwarnings(
                "error", message="Cell .* marked as a date.*", category=UserWarning
            )
            try:
                for row, cells in enumerate(
                    sheet.iter_rows(
                        min_row=selected.header_row or first,
                        max_row=last,
                        min_col=column,
                        max_col=column,
                    ),
                    selected.header_row or first,
                ):
                    _check_cancelled(cancelled)
                    example, key = _classify(cells[0], row)
                    if row == selected.header_row:
                        header = example.text
                        continue
                    if row < first or row > last:
                        raise InspectionError(
                            "Column reader returned rows outside the saved data set."
                        )
                    rows_read += 1
                    accumulator.add(
                        example, key, numeric_zero=example.kind == "Number" and cells[0].value == 0
                    )
                    if on_cell is not None:
                        on_cell(example)
                    if rows_read % 10000 == 0:
                        progress(f"Read {rows_read:,} of {total:,} data rows…")
            except UserWarning as error:
                raise InspectionError(
                    "An Excel date could not be decoded without changing its value. "
                    "No complete column summary was produced."
                ) from error
        _check_cancelled(cancelled)
        # A successful XML stream can end before the explicit range does. Cells
        # absent after EOF are blank; reset_dimensions prevents openpyxl padding them.
        # Parse errors above still fail instead of reaching this accounting step.
        remaining = total - rows_read
        if remaining < 0:
            raise InspectionError("Column reader returned too many rows.")
        for row in range(
            first + rows_read,
            min(last + 1, first + rows_read + max(0, EXAMPLES_PER_TYPE - counts["Blank"])),
        ):
            examples.append(ColumnExample(row, "Blank", "", "General"))
        counts["Blank"] += remaining
        if on_cell is not None:
            for row in range(first + rows_read, last + 1):
                _check_cancelled(cancelled)
                on_cell(ColumnExample(row, "Blank", "", "General"))
                if (row - first + 1) % 10000 == 0:
                    progress(f"Read {row - first + 1:,} of {total:,} data rows…")
        distinct, repeated, extra, repeated_examples = accumulator.finish()
    progress("Verifying the source is unchanged…")
    _verify_source(path, selected.workbook_sha256, cancelled)
    _check_cancelled(cancelled)
    return ColumnProfile(
        table,
        column,
        session,
        header,
        first,
        last,
        tuple(counts.items()),
        accumulator.numeric_zeroes,
        accumulator.whitespace,
        distinct,
        repeated,
        extra,
        tuple(examples),
        repeated_examples,
        accumulator.text_zeroes,
    )
