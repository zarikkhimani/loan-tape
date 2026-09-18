"""Read original OOXML date evidence without converting serials or running Excel."""

from __future__ import annotations

import posixpath
from dataclasses import dataclass
from pathlib import Path
from threading import Event
from xml.etree.ElementTree import Element
from zipfile import ZipFile

from defusedxml.ElementTree import fromstring, iterparse  # type: ignore[import-untyped]
from openpyxl.styles.numbers import BUILTIN_FORMATS  # type: ignore[import-untyped]

from loan_tape.date_parser import DateEvidence, DateKind, DateSource, DateSystem
from loan_tape.inspection import InspectionError
from loan_tape.ooxml_safety import validate_ooxml_archive
from loan_tape.ranges import CellRange, parse_range
from loan_tape.selection import SourceSelection
from loan_tape.workbook_index import NS, REL, ScanCancelled, _digest, _part, check_cancelled


@dataclass(frozen=True)
class DateColumnEvidence:
    selection: SourceSelection
    column: int
    date_system: DateSystem
    cells: tuple[DateEvidence, ...]


def _xml(archive: ZipFile, part: str, root_name: str | None = None) -> Element:
    root: Element = fromstring(archive.read(part), forbid_dtd=True)
    if root_name and root.tag != f"{{{NS}}}{root_name}":
        raise InspectionError(f"Unsupported {root_name} XML format.")
    return root


def _relationships(archive: ZipFile, part: str) -> dict[str, Element]:
    root = _xml(archive, part)
    result: dict[str, Element] = {}
    for item in root:
        key = item.attrib["Id"]
        if key in result:
            raise InspectionError("Duplicate package relationship IDs.")
        result[key] = item
    return result


def _target(base: str, item: Element) -> str:
    if item.get("TargetMode", "").lower() == "external":
        raise InspectionError("External workbook data is unsupported; no links were followed.")
    return _part(base, item.attrib["Target"])


def _optional_part(base: str, rels: dict[str, Element], suffix: str) -> str | None:
    items = [item for item in rels.values() if item.get("Type", "").endswith(suffix)]
    if len(items) > 1:
        raise InspectionError(f"Duplicate {suffix} relationships.")
    return _target(base, items[0]) if items else None


def _formats(archive: ZipFile, part: str | None) -> tuple[str | None, ...]:
    if part is None:
        return ("General",)
    root = _xml(archive, part, "styleSheet")
    formats = dict(BUILTIN_FORMATS)
    custom_ids: set[int] = set()
    for item in root.findall(f"{{{NS}}}numFmts/{{{NS}}}numFmt"):
        key = int(item.attrib["numFmtId"])
        if key in custom_ids:
            raise InspectionError("Duplicate custom number formats.")
        custom_ids.add(key)
        formats[key] = item.attrib["formatCode"]
    xfs = root.findall(f"{{{NS}}}cellXfs/{{{NS}}}xf")
    if not xfs:
        raise InspectionError("Workbook has no cell style definitions.")
    # Unknown built-in format IDs remain unavailable, never guessed from locale.
    return tuple(formats.get(int(item.get("numFmtId", "0"))) for item in xfs)


def _string(element: Element) -> str:
    """Read plain/rich text in order, excluding phonetic annotations."""
    return "".join(
        (child.text or "")
        if child.tag == f"{{{NS}}}t"
        else "".join(part.text or "" for part in child.findall(f"{{{NS}}}t"))
        for child in element
        if child.tag in {f"{{{NS}}}t", f"{{{NS}}}r"}
    )


def _shared_strings(archive: ZipFile, part: str | None, cancelled: Event | None) -> tuple[str, ...]:
    if part is None:
        return ()
    root = _xml(archive, part, "sst")
    result = []
    for item in root.findall(f"{{{NS}}}si"):
        check_cancelled(cancelled)
        result.append(_string(item))
    return tuple(result)


def _cell(
    element: Element,
    source: DateSource,
    system: DateSystem,
    formats: tuple[str | None, ...],
    strings: tuple[str, ...],
    inherited_style: str,
) -> DateEvidence:
    storage = element.get("t", "n")
    style = int(element.get("s", inherited_style))
    if not 0 <= style < len(formats):
        raise InspectionError("Cell references an unavailable style; evidence is incomplete.")
    values = element.findall(f"{{{NS}}}v")
    formulas = element.findall(f"{{{NS}}}f")
    inline = element.findall(f"{{{NS}}}is")
    if len(values) > 1 or len(formulas) > 1 or len(inline) > 1:
        raise InspectionError("Cell contains duplicated value/formula elements.")
    if inline and storage != "inlineStr":
        raise InspectionError("Cell has inline text inconsistent with its stored type.")
    token = values[0].text if values else None
    kind: DateKind
    raw = token
    formula = None
    attributes: tuple[tuple[str, str], ...] = ()
    if formulas:
        kind = "formula"
        formula = formulas[0].text
        attributes = tuple(sorted(formulas[0].attrib.items()))
    elif storage == "inlineStr":
        if values:
            raise InspectionError("Inline string also contains a stored numeric value.")
        kind, raw = "text", _string(inline[0]) if inline else ""
    elif storage == "s":
        index = int(token) if token is not None else -1
        if not 0 <= index < len(strings):
            raise InspectionError("Cell references an unavailable shared string.")
        kind, raw = "text", strings[index]
    elif storage == "str":
        kind, raw = "text", token or ""
    elif token is None:
        if storage not in {"n"}:
            raise InspectionError("Nonblank cell type has no stored value.")
        kind = "blank"
    else:
        kinds: dict[str, DateKind] = {"n": "number", "d": "iso_date", "e": "error", "b": "boolean"}
        kind = kinds.get(storage, "unsupported")
    return DateEvidence(
        source, kind, raw, storage, token, formats[style], system, True, formula, attributes
    )


def _columns(
    archive: ZipFile,
    part: str,
    selection: SourceSelection,
    columns: tuple[int, ...],
    system: DateSystem,
    formats: tuple[str | None, ...],
    strings: tuple[str, ...],
    cancelled: Event | None,
) -> tuple[tuple[DateEvidence, ...], ...]:
    first = selection.header_row + 1 if selection.header_row is not None else selection.area.min_row
    requested = set(columns)
    cells: dict[int, dict[int, DateEvidence]] = {column: {} for column in columns}
    stack: list[Element] = []
    previous_row = 0
    data_blocks = 0
    column_styles: dict[int, list[str]] = {column: [] for column in columns}
    with archive.open(part) as stream:
        for event, element in iterparse(stream, events=("start", "end"), forbid_dtd=True):
            check_cancelled(cancelled)
            if event == "start":
                if not stack and element.tag != f"{{{NS}}}worksheet":
                    raise InspectionError("Unsupported worksheet format.")
                if element.tag == f"{{{NS}}}sheetData":
                    if len(stack) != 1:
                        raise InspectionError("Worksheet data block is nested unexpectedly.")
                    data_blocks += 1
                    if data_blocks > 1:
                        raise InspectionError("Duplicate worksheet data blocks.")
                stack.append(element)
                continue
            if element.tag == f"{{{NS}}}col":
                low, high = int(element.attrib["min"]), int(element.attrib["max"])
                if not 1 <= low <= high <= 16384:
                    raise InspectionError("Invalid column metadata.")
                if "style" in element.attrib:
                    for column in columns:
                        if low <= column <= high:
                            column_styles[column].append(element.attrib["style"])
                            if len(column_styles[column]) > 1:
                                raise InspectionError("Overlapping column styles are unsupported.")
            elif element.tag == f"{{{NS}}}mergeCell":
                merged = parse_range(element.attrib["ref"])
                if (
                    any(merged.min_column <= column <= merged.max_column for column in columns)
                    and merged.min_row <= selection.area.max_row
                    and merged.max_row >= first
                ):
                    raise InspectionError(
                        "Selected date cells intersect a merged range; choose unmerged data cells."
                    )
            elif element.tag == f"{{{NS}}}row":
                if len(stack) < 2 or stack[-2].tag != f"{{{NS}}}sheetData":
                    raise InspectionError("Worksheet row is outside the sheet data block.")
                row = int(element.get("r", str(previous_row + 1)))
                if not previous_row < row <= 1048576:
                    raise InspectionError("Unordered or duplicated worksheet rows.")
                previous_row, previous_column = row, 0
                for cell in element.findall(f"{{{NS}}}c"):
                    address = cell.get("r")
                    position = (
                        parse_range(address)
                        if address
                        else CellRange(row, previous_column + 1, row, previous_column + 1)
                    )
                    if (
                        position.min_row != row
                        or position.row_count != 1
                        or position.column_count != 1
                        or position.min_column <= previous_column
                    ):
                        raise InspectionError("Invalid or duplicated cell coordinates.")
                    previous_column = position.min_column
                    column = position.min_column
                    if first <= row <= selection.area.max_row and column in requested:
                        styles = column_styles[column]
                        inherited = element.get("s", styles[0] if styles else "0")
                        cells[column][row] = _cell(
                            cell,
                            DateSource(selection.workbook_sha256, row, column, selection.sheet),
                            system,
                            formats,
                            strings,
                            inherited,
                        )
                element.clear()
                if len(stack) > 1:
                    stack[-2].remove(element)
            stack.pop()
    if data_blocks != 1:
        raise InspectionError("Worksheet data block is missing or unsupported.")
    result: list[tuple[DateEvidence, ...]] = []
    for column in columns:
        column_result = []
        for row in range(first, selection.area.max_row + 1):
            check_cancelled(cancelled)
            column_result.append(
                cells[column].get(row)
                or DateEvidence(
                    DateSource(selection.workbook_sha256, row, column, selection.sheet),
                    "blank",
                    None,
                    date_system=system,
                    present=False,
                )
            )
        result.append(tuple(column_result))
    return tuple(result)


def read_excel_date_columns(
    path: Path,
    selection: SourceSelection,
    columns: tuple[int, ...],
    *,
    cancelled: Event | None = None,
    max_rows: int = 100_000,
) -> tuple[DateColumnEvidence, ...]:
    """Read mapped date columns in one atomic worksheet scan.

    The source package, styles, shared strings and worksheet are each opened once for
    the complete requested set. Malformed or unreadable evidence in any requested
    column prevents every result from being published.
    """
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise InspectionError("Date evidence reader supports .xlsx and .xlsm only.")
    if not isinstance(columns, tuple) or not columns:
        raise InspectionError("Choose at least one date column inside the saved selection.")
    if any(
        type(column) is not int
        or not selection.area.min_column <= column <= selection.area.max_column
        for column in columns
    ):
        raise InspectionError("Choose date columns inside the saved selection.")
    if len(set(columns)) != len(columns):
        raise InspectionError("Choose each date column only once.")
    first = selection.header_row + 1 if selection.header_row is not None else selection.area.min_row
    if type(max_rows) is not int or not 1 <= max_rows <= 1048576:
        raise ValueError("Date reader row limit must be between 1 and 1048576.")
    if selection.area.max_row - first + 1 > max_rows:
        raise InspectionError(
            "Selected date columns exceed the configured row limit; no rows were read."
        )
    check_cancelled(cancelled)
    try:
        before = _digest(path, cancelled)
        if before != selection.workbook_sha256:
            raise InspectionError("Source identity changed; select the preserved file again.")
        with ZipFile(path) as archive:
            validate_ooxml_archive(
                archive,
                error_type=InspectionError,
                cancelled=lambda: check_cancelled(cancelled),
            )
            root_rels = _relationships(archive, "_rels/.rels")
            part = _optional_part("", root_rels, "/officeDocument")
            if part is None:
                raise InspectionError("Workbook part is missing.")
            book = _xml(archive, part, "workbook")
            props = book.findall(f"{{{NS}}}workbookPr")
            if len(props) > 1:
                raise InspectionError("Duplicate workbook date-system properties.")
            prop = props[0] if props else None
            mode = prop.get("date1904", "0") if prop is not None else "0"
            if mode not in {"0", "1", "false", "true"}:
                raise InspectionError("Invalid workbook date system.")
            system: DateSystem = "1904" if mode in {"1", "true"} else "1900"
            folder, filename = posixpath.split(part)
            rels = _relationships(archive, posixpath.join(folder, "_rels", filename + ".rels"))
            sheets = [
                s
                for s in book.findall(f"{{{NS}}}sheets/{{{NS}}}sheet")
                if s.get("name") == selection.sheet
            ]
            if len(sheets) != 1:
                raise InspectionError("Selected worksheet is missing or ambiguous.")
            sheet_rel = rels[sheets[0].attrib[f"{{{REL}}}id"]]
            if not sheet_rel.get("Type", "").endswith("/worksheet"):
                raise InspectionError("Selected tab is not a worksheet.")
            formats = _formats(archive, _optional_part(part, rels, "/styles"))
            strings = _shared_strings(
                archive, _optional_part(part, rels, "/sharedStrings"), cancelled
            )
            cells = _columns(
                archive,
                _target(part, sheet_rel),
                selection,
                columns,
                system,
                formats,
                strings,
                cancelled,
            )
        if _digest(path, cancelled) != before:
            raise InspectionError(
                "Source changed while reading; no complete evidence was returned."
            )
        check_cancelled(cancelled)
        return tuple(
            DateColumnEvidence(selection, column, system, column_cells)
            for column, column_cells in zip(columns, cells, strict=True)
        )
    except (ScanCancelled, InspectionError):
        raise
    except Exception as error:
        raise InspectionError(f"Could not read complete date evidence: {error}") from error


def read_excel_date_column(
    path: Path,
    selection: SourceSelection,
    column: int,
    *,
    cancelled: Event | None = None,
    max_rows: int = 100_000,
) -> DateColumnEvidence:
    """Compatibility wrapper for callers that need one complete date column."""
    return read_excel_date_columns(
        path,
        selection,
        (column,),
        cancelled=cancelled,
        max_rows=max_rows,
    )[0]
