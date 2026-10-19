"""Opt-in real-Excel export smoke, using fictional mixed-cell text/date evidence."""

from __future__ import annotations

import argparse
import os
import tempfile
from datetime import UTC, date, datetime
from pathlib import Path
from threading import Event
from typing import Any
from unittest.mock import patch

from openpyxl import Workbook, load_workbook  # type: ignore[import-untyped]

from loan_tape.column_definitions import ColumnDefinition, save_column_definition, snapshot_field
from loan_tape.date_analysis import DateBinding, analyze_dates
from loan_tape.date_parser import DateProfile
from loan_tape.date_reader import read_excel_date_columns
from loan_tape.date_runs import save_date_run
from loan_tape.excel_process import start_excel
from loan_tape.inspection import InspectionError
from loan_tape.intake import IntakeStore
from loan_tape.packs import PackStore
from loan_tape.ranges import parse_range
from loan_tape.recipe_execution import (
    approve_excel_recipe_step,
    bind_excel_recipe,
    review_excel_recipe,
)
from loan_tape.recipes import recipe_from_revision, save_recipe
from loan_tape.selection import SavedTable, SourceSelection, TableSelections, save_tables
from loan_tape.text_cleanup import TextCleanupSettings, review_excel_text
from loan_tape.text_export import export_text_revision
from loan_tape.text_rule_types import ExactRule, ExactRulesSettings
from loan_tape.transformations import TransformationApproval, apply_transformation
from loan_tape.workbook_index import ScanCancelled
from loan_tape.workspace_reset import reset_file_session


def run(output: Path) -> None:
    output = output.resolve()
    if output.exists():
        raise FileExistsError("Choose a new smoke output name.")
    previous = Path.cwd()
    packs = Path(__file__).resolve().parents[1] / "packs"
    (previous / ".artifacts").mkdir(exist_ok=True)
    with tempfile.TemporaryDirectory(
        prefix="text-export-smoke-", dir=previous / ".artifacts"
    ) as directory:
        work = Path(directory)
        os.chdir(work)
        try:
            original = work / "fictional.xlsx"
            book = Workbook()
            sheet = book.active
            assert sheet is not None
            sheet.title = "Loans"
            sheet.append(["Borrower", "ID", "Closing", "Numeric", "Formula", "Literal", "Error"])
            sheet.append(
                ["CafÃ© Holdings", "00000000000000001", "10/18/2026", 0, "=1+1", "=CMD()", "#N/A"]
            )
            sheet.append(["CafÃ© Holdings", "0002", "ambiguous", 1.125, None, "+literal", None])
            sheet.append([" Acme ", "0003", None, -12, None, "'quote", None])
            sheet["F2"].data_type = "s"
            sheet.row_dimensions[3].hidden = True
            book.save(original)
            book.close()
            before = original.read_bytes()
            intake = IntakeStore(work / "inputs")
            record = intake.save_path(original)
            source = Path(str(intake.describe(record)["saved_path"]))
            table = SavedTable(
                "a" * 32,
                "Fictional export",
                SourceSelection(str(record["sha256"]), "Loans", parse_range("A1:G5"), 1),
            )
            save_tables(TableSelections(table.selection.workbook_sha256, (table,), table.id))
            store = PackStore(packs, work / ".artifacts/packs")
            for pack in ("loan.text", "loan.dates", "example.loans"):
                store.activate(pack)
            catalog = store.catalog()
            for col, header, field in (
                (1, "Borrower", "loan.text:borrower_name"),
                (2, "ID", "example.loans:loan_id"),
                (3, "Closing", "loan.dates:closing_date"),
            ):
                save_column_definition(
                    table,
                    ColumnDefinition(
                        col,
                        header,
                        snapshot_field(catalog, catalog.get_field(field)),
                        date_formats=("M/D/YYYY",) if col == 3 else (),
                    ),
                )
            first = review_excel_text(source, table, 1, TextCleanupSettings(), store)
            approval = TransformationApproval(
                first.preview.id,
                ((2, 1),),
                "Smoke reviewer",
                "Fictional approved repair",
                datetime.now(UTC).isoformat(),
            )
            parent = apply_transformation(first.snapshot, first.preview, approval)
            second = review_excel_text(
                source,
                table,
                1,
                ExactRulesSettings((ExactRule("Café Holdings", "=Reviewed Name"),)),
                store,
                parent=parent,
            )
            approval = TransformationApproval(
                second.preview.id,
                ((2, 1),),
                "Smoke reviewer",
                "Fictional literal replacement",
                datetime.now(UTC).isoformat(),
            )
            revision = apply_transformation(
                second.snapshot, second.preview, approval, parent=parent
            )
            recipe = recipe_from_revision(
                revision,
                name="Fictional export recipe",
                profile="Fictional template",
                applicability="Review the fictional reporting population.",
            )
            recipe_path = save_recipe(recipe)
            binding = bind_excel_recipe(
                source,
                table,
                store,
                recipe,
                profile="Fictional template",
                context="Fictional October population",
                reviewer="Smoke reviewer",
                confirmed=True,
            )
            staged = None
            for _ in recipe.steps:
                review = review_excel_recipe(source, table, store, recipe, binding, parent=staged)
                approval = TransformationApproval(
                    review.preview.id,
                    ((2, 1),),
                    "Smoke reviewer",
                    "Fresh cell choice for this recipe step",
                    datetime.now(UTC).isoformat(),
                )
                staged, _ = approve_excel_recipe_step(
                    source, table, store, recipe, binding, review, approval, parent=staged
                )
            assert staged is not None
            evidence = read_excel_date_columns(source, table.selection, (3,))
            analysis = analyze_dates(
                table,
                evidence,
                (
                    DateBinding(
                        3, "loan.dates:closing_date", DateProfile(text_formats=("M/D/YYYY",))
                    ),
                ),
                catalog,
                reference_date=date(2026, 10, 18),
            )
            date_run = save_date_run(analysis, work / ".artifacts/date-runs")
            result = export_text_revision(source, table, store, staged, output, date_run=date_run)
            book = load_workbook(result.path, read_only=True, data_only=False, keep_links=False)
            try:
                sheet = book["Loan Tape"]
                assert [c.value for c in sheet[2]] == [
                    "=Reviewed Name",
                    "00000000000000001",
                    "2026-10-18",
                    "0",
                    "=1+1",
                    "=CMD()",
                    "#N/A",
                ]
                assert sheet["A3"].value == "CafÃ© Holdings"
                assert sheet["D3"].value == "1.125" and sheet["F4"].value == "'quote"
                assert sheet.max_row == 5 and sheet.max_column == 7
                assert book["Steps"]["L2"].value == recipe.to_dict()["id"]
            finally:
                book.close()
            assert original.read_bytes() == before == source.read_bytes()
            # Exercise cancellation after a real owned Excel process has started.
            cancelled = Event()

            def cancel_after_start(job: Any) -> tuple[Any, int]:
                owned = start_excel(job)
                cancelled.set()
                return owned

            rejected = work / "cancelled.xlsx"
            with patch("loan_tape.text_export_excel.start_excel", side_effect=cancel_after_start):
                try:
                    export_text_revision(
                        source, table, store, staged, rejected, cancelled=cancelled
                    )
                except ScanCancelled:
                    pass
                else:
                    raise AssertionError("Cancellation did not stop the owned export.")
            assert not rejected.exists()
            try:
                export_text_revision(source, table, store, staged, rejected, timeout=0.001)
            except InspectionError as error:
                assert "timed out" in str(error)
            else:
                raise AssertionError("Timeout did not stop the owned export.")
            assert not rejected.exists() and not list(work.glob(".*.tmp.xlsx"))
            reset_file_session(intake, "b" * 32)
            assert not source.exists() and output.exists() and recipe_path.exists()
            assert original.read_bytes() == before
            print(
                f"Real xlwings {result.writer_version} text/recipe/date export validated: {result.path}"
            )
            print(
                f"SHA-256 {result.sha256}; complete cells, literal text, audit, retention, real cancellation/timeout and owned process cleanup passed."
            )
        finally:
            os.chdir(previous)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", required=True, type=Path)
    run(parser.parse_args().output)
