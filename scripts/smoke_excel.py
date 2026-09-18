"""Opt-in integration check using installed desktop Excel and synthetic workbooks."""

from __future__ import annotations

import atexit
import importlib
import sys
import tempfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook  # type: ignore[import-untyped]

from loan_tape.excel_compare import ComparedPreview, read_parallel_preview
from loan_tape.ranges import parse_range
from loan_tape.workbook_index import scan_workbook


def main() -> None:
    if sys.platform != "win32":
        raise RuntimeError("This integration check requires Windows and desktop Excel.")
    xw = importlib.import_module("xlwings")
    atexit.unregister(importlib.import_module("xlwings._xlwindows").cleanup)
    existing = set(xw.apps.keys())
    artifacts = Path.cwd() / ".artifacts"
    artifacts.mkdir(exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="excel-smoke-", dir=artifacts) as directory:
        path = Path(directory) / "synthetic.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.title = "Loans"
        sheet.append(["ID", "Amount", "Date", "Blank", "Error", "Boolean", "Literal"])
        sheet.append(["000123", 125.25, datetime(2030, 3, 4), None, "#N/A", True, "=1+1"])
        sheet["G2"].data_type = "s"
        sheet.append(["00124", 0, None, "NA"])
        hidden = book.create_sheet("Hidden")
        hidden.sheet_state = "hidden"
        hidden["C4"] = "Résumé"
        book.create_sheet("Empty")
        offset = book.create_sheet("Far away")
        offset["DZ5000"] = "000007"
        offset["EA5000"] = 0
        offset["DZ5001"] = "Résumé"
        offset["EA5250"] = "Last row"
        offset.row_dimensions[5001].hidden = True
        offset.column_dimensions["EA"].hidden = True
        book.save(path)
        book.close()
        before = path.read_bytes()
        for name in ("Loans", "Hidden", "Empty"):
            result = read_parallel_preview(path, sheet=name)
            assert isinstance(result, ComparedPreview)
            assert result.secondary is not None, result.message
            assert not result.differences, result.message
            assert result.primary.rows == () if name == "Empty" else result.primary.rows
            assert "agree" in result.message, result.message
            assert path.read_bytes() == before, "Source bytes changed"
            print(f"{name}: both readers agree; source unchanged.")
        index = scan_workbook(path)
        assert index.complete and index.sheets[-1].bounds == parse_range("DZ5000:EA5250")
        for address in ("DZ5000:EA5001", "EA5250"):
            page = parse_range(address)
            compared = read_parallel_preview(path, sheet="Far away", area=page)
            assert isinstance(compared, ComparedPreview) and compared.secondary is not None
            assert not compared.differences, compared.message
            assert compared.primary.start_row == page.min_row
            assert compared.secondary.start_column == page.min_column
            assert path.read_bytes() == before
            print(f"{address}: both readers agree at original source coordinates.")
    after = set(xw.apps.keys())
    assert after == existing, "Excel process list changed; check for a remaining reader instance."
    print(
        "Desktop Excel check passed; existing Excel instances remain and no reader instance remains."
    )


if __name__ == "__main__":
    main()
