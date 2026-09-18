"""Opt-in real-Excel smoke test for the full refined loan-tape export."""

from __future__ import annotations

import argparse
import tempfile
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook  # type: ignore[import-untyped]

from loan_tape.date_analysis import (
    CalendarBinding,
    DateBinding,
    DatePair,
    analyze_standardized_dates,
)
from loan_tape.date_calendar import build_calendar
from loan_tape.date_export import EXPECTED_SHEETS, export_date_run
from loan_tape.date_parser import DateProfile
from loan_tape.date_runs import save_date_run
from loan_tape.date_standardization import DateMapping, standardize_excel_dates
from loan_tape.pack_format import load_pack
from loan_tape.packs import Catalog
from loan_tape.ranges import parse_range
from loan_tape.selection import SavedTable, SourceSelection
from loan_tape.workbook_index import _digest

ROOT = Path(__file__).resolve().parents[1]


def create_source(path: Path) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "Dates"
    sheet.append(["Loan ID", "Closing Date", "Maturity Date"])
    for row in (
        ("0001", "1/8/21", "9/19/26"),
        ("0002", "1/14/23", "1/13/27"),
        ("0003", "10/9/19", "6/18/28"),
        ("0004", "=TODAY()", "3/17/27"),
        ("0005", None, "2/20/27"),
    ):
        sheet.append(row)
    sheet["B5"].data_type = "s"
    book.save(path)


def run(output: Path) -> None:
    if output.exists():
        raise FileExistsError(f"Choose a new smoke-test output: {output}")
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="loan-tape-date-export-") as directory:
        work = Path(directory)
        source = work / "synthetic-dates.xlsx"
        create_source(source)
        before = source.read_bytes()
        selection = SourceSelection(_digest(source, None), "Dates", parse_range("A1:C6"), 1)
        table = SavedTable("a" * 32, "Synthetic date review", selection)
        profile = DateProfile(
            text_formats=("M/D/YY",),
            two_digit_year_start=2000,
            numeric_dates="excel_serial",
        )
        catalog = Catalog("smoke", (load_pack(ROOT / "packs/loan-dates"),))
        calendar = build_calendar("us.sifma_fixed_income")
        standardized = standardize_excel_dates(
            source,
            table,
            (
                DateMapping(2, "loan.dates:closing_date", profile),
                DateMapping(3, "loan.dates:current_maturity_date", profile),
            ),
            catalog,
        )
        analysis = analyze_standardized_dates(
            standardized,
            (
                DateBinding(2, "loan.dates:closing_date", profile),
                DateBinding(
                    3,
                    "loan.dates:current_maturity_date",
                    profile,
                    calendar=CalendarBinding(calendar.id, "maturity"),
                ),
            ),
            reference_date=date(2026, 9, 17),
            pairs=(DatePair(2, 3),),
            calendars=(calendar,),
        )
        saved = save_date_run(analysis, work / "runs")
        exported = export_date_run(saved, output, source_path=source)
        assert source.read_bytes() == before
        assert exported.writer == "xlwings"
        assert exported.analysis_fingerprint == analysis.fingerprint

    book = load_workbook(output, read_only=True, data_only=False, keep_links=False)
    try:
        assert tuple(book.sheetnames) == EXPECTED_SHEETS
        assert book["Loan Tape"].max_row == 6
        assert book["Loan Tape"].max_column == 3
        assert book["Loan Tape"]["A2"].value == "0001"
        assert book["Loan Tape"]["B2"].value == datetime(2021, 1, 8)
        assert book["Loan Tape"]["C2"].value == datetime(2026, 9, 19)
        assert book["Loan Tape"]["B5"].value == "=TODAY()"
        assert book["Loan Tape"]["B5"].data_type != "f"
        assert book["Audit"]["B3"].value == "xlwings"
        assert book["Summary"]["A1"].value == "Refined loan tape and date review"
        assert all(
            cell.data_type != "f" for sheet in book.worksheets for row in sheet for cell in row
        )
        raw_values = [row[5].value for row in book["Date Results"].iter_rows(min_row=2)]
        assert "=TODAY()" in raw_values
    finally:
        book.close()
    print(f"date export smoke passed: {output}")


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    run(args.output.resolve())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
