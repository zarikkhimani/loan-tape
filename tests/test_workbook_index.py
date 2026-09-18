"""Whole-workbook coverage and source-coordinate page regression cases."""

from dataclasses import replace
from threading import Event
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook
from openpyxl.styles import PatternFill

from loan_tape.excel_compare import compare_previews
from loan_tape.inspection import read_preview
from loan_tape.ranges import CellRange, parse_range
from loan_tape.selection import SourceSelection, load_selection, save_selection, selection_path
from loan_tape.workbook_index import ScanCancelled, scan_workbook


def rewrite(path, member, change):
    with ZipFile(path) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    entries[member] = change(entries[member])
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        for name, value in entries.items():
            archive.writestr(name, value)


def test_discovers_far_cells_all_hidden_content_and_separate_blocks(tmp_path):
    path = tmp_path / "anywhere.xlsx"
    book = Workbook()
    book.active.title = "Empty first tab"
    sheet = book.create_sheet("Portfolio")
    sheet["G450"] = "ID"
    sheet["H450"] = "Amount"
    sheet["G451"] = "000123"
    sheet["H451"] = 0
    sheet["L450"] = "Separate table"
    sheet["L451"] = "NA"
    sheet["G700"] = "After a blank gap"
    sheet["XFD1048576"] = "At the Excel boundary"
    sheet.row_dimensions[451].hidden = True
    sheet.column_dimensions["H"].hidden = True
    hidden = book.create_sheet("Hidden data")
    hidden.sheet_state = "veryHidden"
    hidden["ZZ9000"] = "=1+1"
    hidden.merge_cells("ZZ9000:AAA9000")
    style = book.create_sheet("Only formatting")
    style["XFD1048576"].fill = PatternFill("solid", fgColor="FF0000")
    book.save(path)
    book.close()
    before = path.read_bytes()
    result = scan_workbook(path)
    assert result.complete
    assert len(result.sheets) == 4
    assert result.sheets[0].bounds is None
    main = result.sheets[1]
    assert main.bounds.address == "G450:XFD1048576"
    assert main.stored_cells == 8
    assert [area.address for area in main.areas] == ["G450:H451", "L450:L451", "G700", "XFD1048576"]
    assert main.hidden_rows == 1 and main.hidden_columns == 1
    assert result.sheets[2].visibility == "veryHidden"
    assert result.sheets[2].bounds.address == "ZZ9000"
    assert result.sheets[2].merged_ranges == 1
    assert result.sheets[3].bounds is None
    preview = read_preview(path, sheet="Portfolio", area=parse_range("XFD1048576"))
    assert preview.rows[0][0].text == "At the Excel boundary"
    assert (preview.start_row, preview.start_column) == (1048576, 16384)
    assert path.read_bytes() == before


def test_wrong_dimensions_do_not_hide_data_or_shift_preview(tmp_path):
    path = tmp_path / "wrong-dimensions.xlsx"
    book = Workbook()
    book.active["DZ5000"] = "000007"
    book.active["EA5000"] = 0
    book.active["DZ5001"] = "NULL"
    book.save(path)
    book.close()
    rewrite(path, "xl/worksheets/sheet1.xml", lambda xml: xml.replace(b"DZ5000:EA5001", b"A1:A1"))
    result = scan_workbook(path)
    assert result.complete
    assert result.sheets[0].bounds.address == "DZ5000:EA5001"
    preview = read_preview(path, area=parse_range("DZ5000:EA5001"))
    assert [[cell.text for cell in row] for row in preview.rows] == [["000007", "0"], ["NULL", ""]]
    assert (preview.start_row, preview.start_column) == (5000, 130)


def test_paging_preserves_source_positions_and_entire_range(tmp_path):
    path = tmp_path / "pages.xlsx"
    book = Workbook()
    for row in range(450, 951):
        book.active.cell(row, 7, f"ID-{row:06}")
    book.active.cell(950, 130, "Far right")
    book.save(path)
    book.close()
    full = parse_range("G450:EA950")
    first = read_preview(path, area=full.page(450, 7))
    last = read_preview(path, area=full.page(940, 127))
    assert len(first.rows) == 20 and first.column_count == 10
    assert first.rows[0][0].text == "ID-000450"
    assert len(last.rows) == 11 and last.start_row == 940 and last.start_column == 127
    assert last.rows[10][3].text == "Far right"
    assert full.row_count == 501 and full.column_count == 125
    assert not compare_previews(first, replace(first, start_row=850)).message.endswith(
        "preview only."
    )


def test_malformed_sheet_is_reported_incomplete_without_partial_counts(tmp_path):
    path = tmp_path / "broken-tail.xlsx"
    book = Workbook()
    book.active["A1"] = "Good early data"
    book.create_sheet("Readable")["X1000"] = "Still discoverable"
    book.save(path)
    book.close()
    rewrite(path, "xl/worksheets/sheet1.xml", lambda xml: xml[:-20])
    result = scan_workbook(path)
    assert not result.complete
    assert result.sheets[0].error and result.sheets[0].bounds is None
    assert result.sheets[1].bounds.address == "X1000"


def test_nonworksheet_tabs_are_explicit_in_coverage(tmp_path):
    path = tmp_path / "chart.xlsx"
    book = Workbook()
    book.active["A1"] = "Data"
    book.save(path)
    book.close()
    rewrite(
        path, "xl/_rels/workbook.xml.rels", lambda xml: xml.replace(b'/worksheet"', b'/chartsheet"')
    )
    result = scan_workbook(path)
    assert not result.complete and "Non-worksheet" in result.sheets[0].error


def test_scan_cancellation_is_not_a_successful_partial_scan(tmp_path):
    path = tmp_path / "cancel.xlsx"
    book = Workbook()
    book.active["A1"] = "Data"
    book.save(path)
    book.close()
    cancel = Event()
    with pytest.raises(ScanCancelled, match="incomplete"):
        scan_workbook(path, cancelled=cancel, progress=lambda message: cancel.set())


def test_many_areas_are_combined_explicitly_without_losing_coverage(tmp_path):
    path = tmp_path / "many.xlsx"
    book = Workbook()
    for row in range(1, 1000, 2):
        book.active.cell(row, 1, row)
    book.save(path)
    book.close()
    sheet = scan_workbook(path).sheets[0]
    assert sheet.condensed and sheet.stored_cells == 500
    assert sheet.areas == (parse_range("A1:A999"),)


def test_saved_scope_contains_all_rows_not_current_page(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    selection = SourceSelection("a" * 64, "Portfolio", parse_range("G450:AZ18400"), 450)
    save_selection(selection)
    assert load_selection("a" * 64) == selection
    assert load_selection("b" * 64) is None
    assert load_selection("a" * 64).area.row_count == 17951
    page = selection.area.page(15000, 7)
    assert page.row_count == 20
    assert load_selection("a" * 64) == selection
    selection_path("a" * 64).write_text('{"broken":true}')
    with pytest.raises(ValueError, match="Saved data set"):
        load_selection("a" * 64)


@pytest.mark.parametrize("text", ["A0", "XFE1", "A1048577", "B5:A1", "Sheet1!A1", "A1:B", "A1;B2"])
def test_invalid_range_is_rejected(text):
    with pytest.raises(ValueError):
        parse_range(text)


def test_absolute_address_and_header_validation():
    assert parse_range("$dz$5000:$EA$5001") == CellRange(5000, 130, 5001, 131)
    with pytest.raises(ValueError, match="Header row"):
        SourceSelection("a" * 64, "Loans", parse_range("G450:H600"), 1)
    with pytest.raises(ValueError, match="inside"):
        parse_range("G450:H600").page(1, 1)
