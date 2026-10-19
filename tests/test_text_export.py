import json
from dataclasses import replace
from threading import Event

import pytest
from openpyxl import load_workbook
from text_export_support import prepare_export, write_fixture

from loan_tape.column_definitions import ColumnRules, definition_for_column, save_column_definition
from loan_tape.inspection import InspectionError
from loan_tape.recipe_execution import recipe_scope
from loan_tape.text_export import export_text_revision
from loan_tape.text_export_excel import validate_output, validate_payload
from loan_tape.text_export_payload import build_text_export
from loan_tape.workbook_index import ScanCancelled


@pytest.fixture
def prepared(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    return prepare_export(tmp_path)


def payload(prepared, **options):
    _, source, _, table, store, revision, _ = prepared
    return build_text_export(
        source, table, recipe_scope(table), store.catalog(), revision, **options
    )


def test_complete_source_overlay_and_exact_audit(prepared):
    before = prepared[0].read_bytes()
    result = payload(prepared)
    sheets = {sheet["name"]: sheet for sheet in result["sheets"]}
    tape = sheets["Loan Tape"]
    assert len(tape["rows"]) == 7 and len(tape["rows"][0]) == 7
    assert tape["header_row"] == 2
    assert tape["rows"][2] == [
        "=Reviewed Name",
        "00000000000000000001",
        "10/18/2026",
        "0",
        "=1+1",
        "=1+1",
        "#N/A",
    ]
    assert tape["rows"][3][0] == "CafÃ© Holdings"  # unapproved hidden row retained
    assert tape["rows"][-1] == [None] * 7
    raw = [json.loads(row[4]) for row in sheets["Source Evidence"]["rows"][1:]]
    assert len(raw) == 49
    assert {r["source"]["kind"] for r in raw} >= {"text", "formula", "blank", "error", "number"}
    audit = sheets["Audit"]["rows"][1:]
    assert json.loads("".join(row[4] for row in audit)) == prepared[5].to_dict()
    assert prepared[0].read_bytes() == prepared[1].read_bytes() == before


def test_compatible_dates_compose_only_valid_cells(prepared):
    result = payload(prepared, date_run=prepared[6])
    tape = result["sheets"][0]["rows"]
    assert tape[2][2] == "2026-10-18" and tape[4][2] == "2026-10-19"
    assert tape[3][2] == "ambiguous" and tape[5][2] is None
    assert tape[2][0] == "=Reviewed Name"
    assert [sheet["name"] for sheet in result["sheets"]][-2:] == ["Date Results", "Date Findings"]


@pytest.mark.parametrize(
    "change", ["source", "mapping", "dictionary", "date_source", "date_evidence"]
)
def test_stale_or_incompatible_inputs_fail(prepared, change):
    from loan_tape._date_json import canonical
    from loan_tape.date_analysis import DateAnalysis, replay_analysis

    original, source, _, table, store, _, run = prepared
    if change == "source":
        source.write_bytes(original.read_bytes() + b"changed")
    elif change == "mapping":
        definition = definition_for_column(table, 3)
        save_column_definition(
            table, replace(definition, rules=ColumnRules(missing_tokens=("N/A",)))
        )
    elif change == "dictionary":
        store.deactivate("loan.dates")
    else:
        data = run.analysis.to_dict()
        if change == "date_source":
            data["inputs"]["dataset"]["name"] = "Different selection"
        else:
            data["inputs"]["evidence"][0]["cells"][0]["raw"] = "10/17/2026"
        changed = replay_analysis(data)
        run = replace(run, analysis=DateAnalysis(canonical(changed.to_dict())))
    with pytest.raises((ValueError, InspectionError)):
        payload(prepared, date_run=run)


@pytest.mark.parametrize("damage", ["formula", "identifier", "row", "sheet", "audit", "link"])
def test_validation_rejects_changed_output(prepared, tmp_path, damage):
    result = payload(prepared)
    target = tmp_path / "fixture.xlsx"
    write_fixture(target, result)
    assert dict(validate_output(target, result))["Loan Tape"] == 7
    book = load_workbook(target)
    if damage == "formula":
        book["Loan Tape"]["A3"] = "=1+1"
    elif damage == "identifier":
        book["Loan Tape"]["B3"] = 1
    elif damage == "row":
        book["Loan Tape"].delete_rows(7)
    elif damage == "sheet":
        book.create_sheet("Unexpected")
    elif damage == "link":
        book["Loan Tape"]["A3"].hyperlink = "https://example.invalid/fictional"
    else:
        book["Audit"]["E2"] = "changed"
    book.save(target)
    book.close()
    with pytest.raises(InspectionError):
        validate_output(target, result)


@pytest.mark.parametrize(
    "failure", ["write", "invalid", "mapping", "source", "cancel", "collision", "after_link"]
)
def test_publication_failures_preserve_sources_and_competing_files(
    prepared, tmp_path, monkeypatch, failure
):
    import loan_tape.text_export as exporter

    original, source, _, table, store, revision, _ = prepared
    before = original.read_bytes()
    target = tmp_path / "out.xlsx"
    event = Event()
    monkeypatch.setattr(exporter.sys, "platform", "win32")
    monkeypatch.setattr(exporter, "version", lambda _: "fixture")

    def writer(data, staging, **kwargs):
        if failure == "write":
            staging.write_bytes(b"partial")
            raise OSError("writer failed")
        write_fixture(staging, data)
        if failure == "invalid":
            staging.write_bytes(b"invalid")
        if failure == "mapping":
            definition = definition_for_column(table, 3)
            save_column_definition(
                table, replace(definition, rules=ColumnRules(missing_tokens=("N/A",)))
            )
        if failure == "source":
            source.write_bytes(before + b"changed")
        if failure == "cancel":
            event.set()
        if failure == "collision":
            target.write_bytes(b"competing owner")
        return "fixture"

    monkeypatch.setattr(exporter, "write_excel", writer)
    real_link = exporter.os.link
    if failure == "after_link":

        def link(a, b):
            real_link(a, b)
            event.set()

        monkeypatch.setattr(exporter.os, "link", link)
    with pytest.raises((OSError, ValueError, InspectionError, ScanCancelled)):
        export_text_revision(source, table, store, revision, target, cancelled=event)
    assert original.read_bytes() == before
    if failure == "collision":
        assert target.read_bytes() == b"competing owner"
    else:
        assert not target.exists()
    assert not list(tmp_path.glob(".*.tmp.xlsx"))


def test_output_limits_and_cancel_are_explicit(prepared):
    result = payload(prepared)
    result["sheets"][0]["rows"][0][0] = "x" * 32767
    with pytest.raises(ValueError, match="literal-cell"):
        validate_payload(result)
    event = Event()
    event.set()
    with pytest.raises(ScanCancelled):
        payload(prepared, cancelled=event)


def test_preexisting_output_and_platform_guard(prepared, tmp_path, monkeypatch):
    import loan_tape.text_export as exporter

    _, source, _, table, store, revision, _ = prepared
    target = tmp_path / "keep.xlsx"
    target.write_bytes(b"keep")
    monkeypatch.setattr(exporter.sys, "platform", "win32")
    with pytest.raises(FileExistsError):
        export_text_revision(source, table, store, revision, target)
    assert target.read_bytes() == b"keep"
    monkeypatch.setattr(exporter.sys, "platform", "linux")
    with pytest.raises(InspectionError, match="requires Windows"):
        export_text_revision(source, table, store, revision, tmp_path / "new.xlsx")


def test_incomplete_recipe_cannot_escape_through_export(tmp_path, monkeypatch):
    from recipe_support import approval, prepare_recipe

    from loan_tape.recipe_execution import (
        approve_excel_recipe_step,
        bind_excel_recipe,
        review_excel_recipe,
    )

    monkeypatch.chdir(tmp_path)
    context, recipe, _, _ = prepare_recipe(tmp_path)
    _, source, _, table, store = context
    binding = bind_excel_recipe(
        source,
        table,
        store,
        recipe,
        profile=recipe.to_dict()["profile"],
        context="Fictional current population",
        reviewer="Reviewer",
        confirmed=True,
    )
    review = review_excel_recipe(source, table, store, recipe, binding)
    partial, path = approve_excel_recipe_step(
        source, table, store, recipe, binding, review, approval(review)
    )
    assert path is None
    with pytest.raises(ValueError, match="incomplete staged recipe"):
        build_text_export(source, table, recipe_scope(table), store.catalog(), partial)


def test_bounded_export_evidence_and_complete_population(prepared, monkeypatch):
    import loan_tape.text_export_payload as builder

    monkeypatch.setattr(builder, "MAX_EXPORT_BYTES", 500)
    with pytest.raises(ValueError, match="byte or worksheet row limit"):
        payload(prepared)
    monkeypatch.setattr(builder, "MAX_COMPLETE_CELLS", 48)
    with pytest.raises(ValueError, match="complete selected-cell limit"):
        payload(prepared)


def test_mapping_conflict_between_valid_date_run_and_text_history(prepared):
    from datetime import date

    from loan_tape.date_analysis import DateBinding, analyze_dates
    from loan_tape.date_parser import DateProfile
    from loan_tape.date_reader import read_excel_date_columns

    _, source, _, table, store, _, run = prepared
    analysis = analyze_dates(
        table,
        read_excel_date_columns(source, table.selection, (3,)),
        (DateBinding(3, "loan.dates:closing_date", DateProfile(text_formats=("M/D/YYYY",))),),
        store.catalog(),
        reference_date=date(2026, 10, 18),
    )
    with pytest.raises(ValueError, match="mappings or interpretation settings conflict"):
        payload(prepared, date_run=replace(run, analysis=analysis))
