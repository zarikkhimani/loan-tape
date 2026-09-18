"""Exercise real readers against synthetic files, including source preservation."""

from datetime import datetime
from pathlib import Path
from zipfile import ZIP_DEFLATED, BadZipFile, ZipFile

import pytest
from openpyxl import Workbook

from loan_tape.inspection import MAX_COLUMNS, MAX_ROWS, InspectionError, read_preview


def values(preview):
    return [[cell.text for cell in row] for row in preview.rows]


def test_csv_preserves_text_duplicate_headers_blanks_and_embedded_newlines(tmp_path: Path):
    path = tmp_path / "data.csv"
    payload = (
        b'ID,Rate,Rate,Note\r\n000123,0,NA,"two, parts\r\nsecond line"\r\n\r\n000124,,NULL,=1+1\r\n'
    )
    path.write_bytes(payload)
    preview = read_preview(path, delimiter="Comma")
    assert values(preview) == [
        ["ID", "Rate", "Rate", "Note"],
        ["000123", "0", "NA", "two, parts\r\nsecond line"],
        [],
        ["000124", "", "NULL", "=1+1"],
    ]
    assert preview.column_count == 4
    assert not preview.more_rows
    assert path.read_bytes() == payload


@pytest.mark.parametrize(
    "delimiter,label", [(";", "semicolon"), ("\t", "tab"), ("|", "pipe"), (",", "comma")]
)
def test_csv_detects_common_separators(tmp_path: Path, delimiter, label):
    path = tmp_path / "data.csv"
    path.write_text(f"ID{delimiter}Balance\n0001{delimiter}125\n", encoding="utf-8-sig")
    preview = read_preview(path)
    assert values(preview)[1] == ["0001", "125"]
    assert label in preview.reading_note


def test_csv_encoding_is_explicit_and_recoverable(tmp_path: Path):
    path = tmp_path / "legacy.csv"
    path.write_bytes("Borrower,ID\nRésumé,0001\n".encode("cp1252"))
    with pytest.raises(InspectionError, match="encoding"):
        read_preview(path)
    assert values(read_preview(path, encoding="Windows-1252"))[1] == ["Résumé", "0001"]
    path.write_text("Borrower\tID\nRésumé\t0001\n", encoding="utf-16")
    assert values(read_preview(path))[1] == ["Résumé", "0001"]


def test_csv_limits_preview_without_converting_values(tmp_path: Path):
    path = tmp_path / "large.csv"
    path.write_text(
        (",".join(["0001"] * (MAX_COLUMNS + 1)) + "\n") * (MAX_ROWS + 1), encoding="utf-8"
    )
    preview = read_preview(path)
    assert len(preview.rows) == MAX_ROWS
    assert preview.column_count == MAX_COLUMNS
    assert preview.more_rows and preview.more_columns
    assert values(preview)[-1][-1] == "0001"


def test_csv_empty_single_column_and_malformed(tmp_path: Path):
    path = tmp_path / "data.csv"
    path.write_text("", encoding="utf-8")
    assert not read_preview(path).rows
    path.write_text("ID\n0001\n0002\n", encoding="utf-8")
    assert values(read_preview(path)) == [["ID"], ["0001"], ["0002"]]
    path.write_text('ID,Note\n001,"unclosed', encoding="utf-8")
    with pytest.raises(InspectionError):
        read_preview(path, delimiter="Comma")


def test_csv_irregular_rows_require_separator_without_dropping_fields(tmp_path: Path):
    path = tmp_path / "data.csv"
    path.write_text("A,B\nx,y,z\np\n", encoding="utf-8")
    with pytest.raises(InspectionError, match="separator"):
        read_preview(path)
    assert values(read_preview(path, delimiter="Comma")) == [["A", "B"], ["x", "y", "z"], ["p"]]


def write_workbook(path: Path):
    book = Workbook()
    sheet = book.active
    sheet.title = "Positions"
    sheet.append(["ID", "Balance", "Rate", "Maturity", "Total"])
    sheet.append(["000123", 0, 0.0625, datetime(2030, 3, 4), "=B2+1"])
    sheet["C2"].number_format = "0.00%"
    sheet["E3"] = "=B2+2"
    other = book.create_sheet("Another sheet")
    other.sheet_state = "hidden"
    other["C4"] = "source location"
    book.create_sheet("Empty")
    book.save(path)
    book.close()
    # Supply an independently controlled cached result; openpyxl does not calculate.
    with ZipFile(path) as archive:
        members = {name: archive.read(name) for name in archive.namelist()}
    members["xl/worksheets/sheet1.xml"] = members["xl/worksheets/sheet1.xml"].replace(
        b"<f>B2+1</f><v />", b"<f>B2+1</f><v>1</v>"
    )
    if path.suffix == ".xlsm":
        members["[Content_Types].xml"] = members["[Content_Types].xml"].replace(
            b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
            b"application/vnd.ms-excel.sheet.macroEnabled.main+xml",
        )
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        for name, content in members.items():
            archive.writestr(name, content)


@pytest.mark.parametrize("extension", [".xlsx", ".xlsm"])
def test_excel_sheets_formulas_raw_values_and_source_preservation(tmp_path: Path, extension):
    path = tmp_path / ("portfolio" + extension)
    write_workbook(path)
    before = path.read_bytes()
    preview = read_preview(path)
    assert preview.sheets == ("Positions", "Another sheet", "Empty")
    assert values(preview)[1] == ["000123", "0", "0.0625", "2030-03-04T00:00:00", "1"]
    assert preview.rows[1][4].formula == "=B2+1"
    assert preview.rows[1][2].number_format == "0.00%"
    assert preview.rows[2][4].text == "=B2+2"
    assert preview.rows[2][4].formula == "=B2+2"
    other = read_preview(path, sheet="Another sheet")
    assert values(other) == [["", "", ""], ["", "", ""], ["", "", ""], ["", "", "source location"]]
    assert not read_preview(path, sheet="Empty").rows
    with pytest.raises(InspectionError, match="worksheet"):
        read_preview(path, sheet="No such sheet")
    assert path.read_bytes() == before
    path.rename(tmp_path / ("closed-handles" + extension))


def test_excel_limits_and_incorrect_dimensions(tmp_path: Path):
    path = tmp_path / "large.xlsx"
    book = Workbook()
    sheet = book.active
    sheet["A1"] = "ID"
    sheet.cell(row=MAX_ROWS + 1, column=1, value="0002")
    sheet.cell(row=2, column=150, value="far column")
    book.save(path)
    book.close()
    with ZipFile(path) as archive:
        members = {name: archive.read(name) for name in archive.namelist()}
    members["xl/worksheets/sheet1.xml"] = members["xl/worksheets/sheet1.xml"].replace(
        f"A1:ET{MAX_ROWS + 1}".encode(), b"A1:A1"
    )
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        for name, content in members.items():
            archive.writestr(name, content)
    preview = read_preview(path)
    assert len(preview.rows) == MAX_ROWS
    assert preview.more_rows and preview.more_columns
    assert preview.column_count == MAX_COLUMNS
    assert values(preview)[0][0] == "ID"


@pytest.mark.parametrize("extension", [".xls", ".xlsb", ".pdf"])
def test_unsupported_formats_fail_explicitly(tmp_path: Path, extension):
    with pytest.raises(InspectionError, match="(?i)preview|inspect"):
        read_preview(tmp_path / ("unsupported" + extension))


def test_damaged_workbook_is_not_reported_as_empty(tmp_path: Path):
    path = tmp_path / "broken.xlsx"
    path.write_bytes(b"not a workbook")
    with pytest.raises(BadZipFile):
        read_preview(path)
