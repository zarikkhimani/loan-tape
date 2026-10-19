"""Full-population text evidence, explicit approvals, invalidation and atomic retention."""

import hashlib
import json
from dataclasses import FrozenInstanceError, replace
from datetime import UTC, datetime
from pathlib import Path
from threading import Event

import pytest
from openpyxl import Workbook

from loan_tape._date_json import canonical
from loan_tape.column_definitions import (
    ColumnDefinition,
    ColumnRules,
    DefinitionScope,
    snapshot_field,
)
from loan_tape.date_parser import DateEvidence, DateSource
from loan_tape.date_reader import DateColumnEvidence
from loan_tape.intake import IntakeStore
from loan_tape.pack_format import load_pack
from loan_tape.packs import Catalog
from loan_tape.ranges import parse_range
from loan_tape.selection import SavedTable, SourceSelection
from loan_tape.transformation_excel import (
    load_transformation,
    publish_transformation,
    read_transformation_snapshot,
)
from loan_tape.transformations import (
    CellProposal,
    TransformationApproval,
    TransformationOperation,
    TransformationPreview,
    TransformationRevision,
    apply_transformation,
    build_snapshot,
    derived_text_values,
    preview_transformation,
    replay_revision,
)
from loan_tape.workbook_index import ScanCancelled
from loan_tape.workspace_reset import reset_file_session

OPERATION = TransformationOperation("reviewed-example", "1.0.0", "synthetic-proposals", "1.0.0")


@pytest.fixture
def prepared(tmp_path):
    pack = load_pack(Path(__file__).parents[1] / "packs/example-loans")
    name = replace(
        pack.fields[0],
        id="borrower_name",
        label="Borrower name",
        entity="borrower",
        data_type="text",
        definition="Fictional source-supplied borrower name.",
        required=False,
        blank_allowed=True,
    )
    pack = replace(pack, fields=(*pack.fields, name), fingerprint="c" * 64)
    catalog = Catalog("synthetic", (pack,))
    path = tmp_path / "synthetic.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Loans"
    headers = ("Name", "Name", "Loan ID", "Balance", "Status", "Date", "Unmapped")
    for column, header in enumerate(headers, 2):
        sheet.cell(3, column, header)
    for column, value in enumerate(
        (" Acme ", "Beta", "00001234567890123456", 123.45, "CURRENT", "2026-10-18", "leave alone"),
        2,
    ):
        sheet.cell(4, column, value)
    sheet["B5"] = " Acme "
    sheet["B6"] = ""
    sheet["B7"] = " \t"
    sheet["B8"] = "N/A"
    sheet["B9"] = 0
    sheet["B10"] = "=1+1"
    sheet["B11"] = "#N/A"
    sheet["B12"] = False
    sheet["B13"].number_format = "@"  # Present blank, distinct from absent B14/B16.
    sheet["B15"] = "Café"
    sheet.row_dimensions[5].hidden = True
    sheet.column_dimensions["C"].hidden = True
    book.save(path)
    book.close()
    sha = hashlib.sha256(path.read_bytes()).hexdigest()
    table = SavedTable(
        "1" * 32, "Fictional loans", SourceSelection(sha, "Loans", parse_range("B3:H16"), 3)
    )
    definitions = []
    ids = (
        "borrower_name",
        "borrower_name",
        "loan_id",
        "principal_balance",
        "status",
        "maturity_date",
    )
    for column, (header, field_id) in enumerate(zip(headers[: len(ids)], ids, strict=True), 2):
        kwargs = {}
        if field_id == "principal_balance":
            kwargs = {"currency": "USD", "source_unit": "whole currency units"}
        if field_id == "maturity_date":
            kwargs = {"date_formats": ("YYYY-MM-DD",)}
        definitions.append(
            ColumnDefinition(
                column,
                header,
                snapshot_field(catalog, catalog.get_field("example.loans:" + field_id)),
                rules=ColumnRules(missing_tokens=("N/A",)),
                **kwargs,
            )
        )
    scope = DefinitionScope(table.id, table.name, table.selection, tuple(definitions))
    return path, table, scope, catalog


def snapshot(prepared):
    return read_transformation_snapshot(*prepared)


def proposal(row=4, column=2, before=" Acme ", after="Acme"):
    return CellProposal(
        row, column, before, after, "literal-review", "Reviewed synthetic spelling."
    )


def approval(preview, selected=((4, 2),)):
    return TransformationApproval(
        preview.id,
        selected,
        "Synthetic reviewer",
        "Reviewed against source.",
        datetime.now(UTC).isoformat(),
    )


def test_full_population_source_evidence_and_explicit_overlay(prepared):
    path, *_ = prepared
    before = path.read_bytes()
    source = snapshot(prepared)
    data = source.to_dict()
    assert data["population"] == {
        "first_row": 4,
        "row_count": 13,
        "mapped_cell_count": 78,
        "selected_column_count": 7,
        "unmapped_column_count": 1,
    }
    assert [item["header"]["occurrence"] for item in data["fields"][:2]] == [1, 2]
    cells = {(c["row"], c["column"]): c for c in data["cells"]}
    assert [cells[r, 2]["state"] for r in range(6, 15)] == [
        "empty_text",
        "whitespace_text",
        "missing_token",
        "number",
        "formula",
        "error",
        "boolean",
        "blank",
        "absent",
    ]
    assert cells[10, 2]["evidence"]["formula"] == "1+1"
    assert cells[4, 4]["evidence"]["raw"] == "00001234567890123456"
    assert cells[4, 5]["evidence"]["raw_token"] == "123.45"
    preview = preview_transformation(source, OPERATION, (proposal(), proposal(row=5)))
    counts = preview.to_dict()["reconciliation"]
    assert (
        sum(counts[k] for k in ("changed", "unchanged", "ineligible", "unresolved", "failed")) == 78
    )
    assert counts["changed"] == 2 and counts["eligible"] == 4
    revision = apply_transformation(source, preview, approval(preview))
    assert derived_text_values(revision) == ((4, 2, "Acme"),)
    assert revision.to_dict()["steps"][0]["rejected_count"] == 1
    assert revision.to_dict()["snapshot"] == data
    assert path.read_bytes() == before
    with pytest.raises(FrozenInstanceError):
        revision.document_json = "bad"
    data["cells"].clear()
    assert len(source.to_dict()["cells"]) == 78


def test_no_preselected_changes_and_literal_text_is_not_executed(prepared):
    source = snapshot(prepared)
    preview = preview_transformation(source, OPERATION, (proposal(after='=HYPERLINK("example")'),))
    rejected = apply_transformation(source, preview, approval(preview, ()))
    assert derived_text_values(rejected) == ()
    accepted = apply_transformation(source, preview, approval(preview))
    assert derived_text_values(accepted)[0][2] == '=HYPERLINK("example")'
    assert accepted.to_dict()["snapshot"]["cells"][0]["evidence"]["kind"] == "text"


@pytest.mark.parametrize(
    "row,column,before",
    [
        (4, 4, "00001234567890123456"),
        (4, 5, "123.45"),
        (4, 6, "CURRENT"),
        (4, 7, "2026-10-18"),
        (6, 2, ""),
        (7, 2, " \t"),
        (8, 2, "N/A"),
        (9, 2, "0"),
        (10, 2, ""),
        (11, 2, "#N/A"),
        (12, 2, "0"),
        (14, 2, ""),
    ],
)
def test_ineligible_cells_cannot_receive_proposals(prepared, row, column, before):
    with pytest.raises(ValueError, match="ineligible"):
        preview_transformation(snapshot(prepared), OPERATION, (proposal(row, column, before),))


@pytest.mark.parametrize("after", ["", "  ", "N/A"])
def test_missingness_cannot_be_changed_by_cleanup(prepared, after):
    source = snapshot(prepared)
    preview = preview_transformation(source, OPERATION, (proposal(after=after),))
    assert preview.to_dict()["reconciliation"]["unresolved"] == 1
    with pytest.raises(ValueError, match="only explicitly proposed"):
        apply_transformation(source, preview, approval(preview))


def test_exact_inputs_conflicts_out_of_scope_and_failed_operation(prepared):
    source = snapshot(prepared)
    with pytest.raises(ValueError, match="literal"):
        preview_transformation(source, OPERATION, (proposal(before="Acme"),))
    with pytest.raises(ValueError, match="conflicting"):
        preview_transformation(source, OPERATION, (proposal(), proposal(after="Different")))
    with pytest.raises(ValueError, match="outside"):
        preview_transformation(source, OPERATION, (proposal(column=8),))
    preview = preview_transformation(
        source,
        OPERATION,
        (
            proposal(),
            replace(proposal(row=5), status="failed", proposed_text=None, reason="Engine failed."),
        ),
    )
    with pytest.raises(ValueError, match="Failed transformation"):
        apply_transformation(source, preview, approval(preview))


def test_derived_parent_intermediate_values_and_replay(prepared):
    source = snapshot(prepared)
    first = preview_transformation(source, OPERATION, (proposal(),))
    parent = apply_transformation(source, first, approval(first))
    second = preview_transformation(
        source,
        replace(OPERATION, version="2.0.0"),
        (proposal(before="Acme", after="ACME"),),
        parent=parent,
    )
    result = apply_transformation(source, second, approval(second), parent=parent)
    assert derived_text_values(parent) == ((4, 2, "Acme"),)
    assert derived_text_values(result) == ((4, 2, "ACME"),)
    assert result.to_dict()["steps"][1]["preview"]["parent_revision_id"] == parent.id
    assert replay_revision(source, result.document_json) == result
    with pytest.raises(ValueError):
        apply_transformation(source, second, approval(second))
    damaged = result.to_dict()
    damaged["steps"][0]["applied_count"] = 99
    with pytest.raises(ValueError, match="replay"):
        preview_transformation(
            source, OPERATION, (), parent=TransformationRevision(canonical(damaged))
        )


def test_preview_parameter_change_requires_new_approval(prepared):
    source = snapshot(prepared)
    preview = preview_transformation(source, OPERATION, (proposal(),))
    changed = preview_transformation(
        source, replace(OPERATION, parameters_json='{"mode":"new"}'), (proposal(),)
    )
    with pytest.raises(ValueError, match="stale"):
        apply_transformation(source, changed, approval(preview))
    damaged = preview.to_dict()
    damaged["decisions"][0]["proposed_text"] = "Invented"
    with pytest.raises(ValueError, match="modified"):
        apply_transformation(source, TransformationPreview(canonical(damaged)), approval(preview))


def test_pages_cover_all_rows_and_reject_bad_bounds(prepared):
    preview = preview_transformation(snapshot(prepared), OPERATION, ())
    assert sum(len(preview.page(offset, 7)) for offset in range(0, 78, 7)) == 78
    assert preview.page(78) == ()
    with pytest.raises(ValueError, match="page size"):
        preview.page(0, 201)


def test_source_and_active_dictionary_changes_invalidate(prepared, tmp_path):
    path, table, scope, catalog = prepared
    preview = preview_transformation(snapshot(prepared), OPERATION, (proposal(),))
    changed_definition = replace(
        scope.definitions[0], rules=ColumnRules(missing_tokens=("MISSING",))
    )
    changed_scope = replace(scope, definitions=(changed_definition, *scope.definitions[1:]))
    with pytest.raises(ValueError, match="stale"):
        publish_transformation(
            path,
            table,
            changed_scope,
            catalog,
            preview,
            approval(preview),
            artifact_root=tmp_path / "artifacts",
        )
    changed_pack = replace(catalog.packs[0], fingerprint="e" * 64)
    with pytest.raises(ValueError, match="dictionary meaning"):
        read_transformation_snapshot(path, table, scope, replace(catalog, packs=(changed_pack,)))
    path.write_bytes(path.read_bytes() + b"changed")
    with pytest.raises(ValueError, match="identity changed"):
        publish_transformation(
            path,
            table,
            scope,
            catalog,
            preview,
            approval(preview),
            artifact_root=tmp_path / "artifacts",
        )
    assert not (tmp_path / "artifacts").exists()


def test_saved_revision_reloads_with_actual_source_and_detects_tampering(prepared, tmp_path):
    source = snapshot(prepared)
    preview = preview_transformation(source, OPERATION, (proposal(),))
    revision, path = publish_transformation(
        *prepared, preview, approval(preview), artifact_root=tmp_path / "artifacts"
    )
    assert load_transformation(path, *prepared) == revision
    assert path.name.startswith(prepared[1].selection.workbook_sha256)
    document = json.loads(path.read_text(encoding="utf-8"))
    document["revision"]["steps"][0]["rejected_count"] = 50
    path.write_text(canonical(document), encoding="utf-8")
    with pytest.raises(ValueError, match="replay"):
        load_transformation(path, *prepared)


@pytest.mark.parametrize("failure", ["cancel", "write", "source_change", "publish"])
def test_failure_after_staging_never_publishes_partial_result(
    prepared, tmp_path, monkeypatch, failure
):
    import loan_tape.transformation_excel as adapter

    preview = preview_transformation(snapshot(prepared), OPERATION, (proposal(),))
    cancelled = Event()
    original_fsync = adapter.os.fsync

    def stage(fd):
        original_fsync(fd)
        if failure == "cancel":
            cancelled.set()
        elif failure == "write":
            raise OSError("Synthetic disk failure")
        elif failure == "source_change":
            prepared[0].write_bytes(prepared[0].read_bytes() + b"changed")

    monkeypatch.setattr(adapter.os, "fsync", stage)
    if failure == "publish":

        def fail_link(*_):
            raise OSError("Synthetic publication failure")

        monkeypatch.setattr(adapter.os, "link", fail_link)
    with pytest.raises((OSError, ValueError, ScanCancelled)):
        publish_transformation(
            *prepared,
            preview,
            approval(preview),
            artifact_root=tmp_path / "artifacts",
            cancelled=cancelled,
        )
    assert list((tmp_path / "artifacts/transformations").iterdir()) == []


def test_publish_refuses_to_overwrite_existing_revision(prepared, tmp_path, monkeypatch):
    import loan_tape.transformation_excel as adapter

    class FixedId:
        hex = "a" * 32

    monkeypatch.setattr(adapter, "uuid4", FixedId)
    preview = preview_transformation(snapshot(prepared), OPERATION, (proposal(),))
    _, path = publish_transformation(
        *prepared, preview, approval(preview), artifact_root=tmp_path / "artifacts"
    )
    before = path.read_bytes()
    with pytest.raises(FileExistsError):
        publish_transformation(
            *prepared, preview, approval(preview), artifact_root=tmp_path / "artifacts"
        )
    assert path.read_bytes() == before
    assert not list(path.parent.glob("*.tmp"))


def test_session_cleanup_removes_revisions_and_staging_only_for_its_sources(prepared, tmp_path):
    store = IntakeStore(tmp_path / "inputs")
    store.save_path(prepared[0])
    artifacts = tmp_path / "artifacts"
    preview = preview_transformation(snapshot(prepared), OPERATION, (proposal(),))
    _, path = publish_transformation(*prepared, preview, approval(preview), artifact_root=artifacts)
    staging = path.with_suffix(".tmp")
    staging.write_text("interrupted")
    unrelated = path.parent / ("f" * 64 + "-" + "d" * 32 + ".json")
    unrelated.write_text("retain")
    before = prepared[0].read_bytes()
    assert reset_file_session(store, "2" * 32, artifacts) == (1, 2)
    assert not path.exists() and not staging.exists()
    assert unrelated.read_text() == "retain" and prepared[0].read_bytes() == before


def test_oversize_token_is_referenced_and_not_copied_into_result(prepared):
    path, table, scope, catalog = prepared
    raw = "é" * 9_000  # Under Excel's character limit, above the 16 KiB UTF-8 limit.
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Loans"
    worksheet.append(["Name"])
    worksheet.append([raw])
    workbook.save(path)
    workbook.close()
    sha = hashlib.sha256(path.read_bytes()).hexdigest()
    table = replace(table, selection=SourceSelection(sha, "Loans", parse_range("A1:A2"), 1))
    scope = DefinitionScope(
        table.id, table.name, table.selection, (replace(scope.definitions[0], column=1),)
    )
    source = read_transformation_snapshot(path, table, scope, catalog)
    cell = source.to_dict()["cells"][0]
    assert cell["evidence"]["raw"] == {
        "retained": False,
        "utf8_bytes": 18_000,
        "sha256": hashlib.sha256(raw.encode("utf-8")).hexdigest(),
    }
    assert cell["evidence"]["source"]["source_sha256"] == sha
    assert raw not in source.document_json
    preview = preview_transformation(source, OPERATION, ())
    assert preview.page(0, 1)[0]["reason"] == "source_token_limit"
    assert preview.page(0, 1)[0]["input_text"] is None


def test_limits_and_cancellation_reject_before_publication(prepared, monkeypatch):
    import loan_tape.transformation_excel as adapter
    import loan_tape.transformations as core

    source = snapshot(prepared)
    with monkeypatch.context() as local:
        local.setattr(adapter, "MAX_CELLS", 77)
        with pytest.raises(ValueError, match="mapped-cell limit"):
            snapshot(prepared)
    with monkeypatch.context() as local:
        local.setattr(core, "MAX_RESULT_BYTES", 100)
        with pytest.raises(ValueError, match="serialized result limit"):
            preview_transformation(source, OPERATION, ())
    with monkeypatch.context() as local:
        local.setattr(core, "MAX_STEPS", 1)
        preview = preview_transformation(source, OPERATION, (proposal(),))
        parent = apply_transformation(source, preview, approval(preview))
        with pytest.raises(ValueError, match="ordered-step limit"):
            preview_transformation(source, OPERATION, (), parent=parent)
    cancelled = Event()
    cancelled.set()
    with pytest.raises(ScanCancelled):
        preview_transformation(source, OPERATION, (), cancelled=cancelled)


def test_core_rejects_incomplete_reordered_or_foreign_evidence(prepared):
    _, old_table, old_scope, _ = prepared
    table = replace(
        old_table,
        selection=replace(old_table.selection, area=parse_range("B4:B5"), header_row=None),
    )
    definition = replace(old_scope.definitions[0], source_header="")
    scope = DefinitionScope(table.id, table.name, table.selection, (definition,))
    cells = tuple(
        DateEvidence(DateSource(table.selection.workbook_sha256, row, 2, "Loans"), "text", "Acme")
        for row in (4, 5)
    )
    column = DateColumnEvidence(table.selection, 2, "1900", cells)
    assert len(build_snapshot(table, scope, (column,), ()).to_dict()["cells"]) == 2
    with pytest.raises(ValueError, match="Incomplete"):
        build_snapshot(table, scope, (replace(column, cells=cells[:1]),), ())
    with pytest.raises(ValueError, match="coordinates"):
        build_snapshot(table, scope, (replace(column, cells=tuple(reversed(cells))),), ())


def test_blank_header_exact_mapping_and_dictionary_missing_codes(prepared):
    path, table, scope, catalog = prepared
    bad = replace(
        scope,
        definitions=(replace(scope.definitions[0], source_header="Wrong"), *scope.definitions[1:]),
    )
    with pytest.raises(ValueError, match="header"):
        read_transformation_snapshot(path, table, bad, catalog)
    pack = catalog.packs[0]
    fields = tuple(
        replace(field, missing_codes=("UNKNOWN",)) if field.id == "borrower_name" else field
        for field in pack.fields
    )
    changed = replace(catalog, packs=(replace(pack, fields=fields),))
    with pytest.raises(ValueError, match="missing codes"):
        read_transformation_snapshot(path, table, scope, changed)
    accepted = replace(
        scope,
        definitions=tuple(
            replace(item, rules=ColumnRules(missing_tokens=("N/A", "UNKNOWN")))
            for item in scope.definitions
        ),
    )
    assert read_transformation_snapshot(path, table, accepted, changed)


def test_page_boundary_and_headerless_blank_column(prepared):
    _, old_table, old_scope, _ = prepared
    table = replace(
        old_table,
        selection=replace(old_table.selection, area=parse_range("B4:B254"), header_row=None),
    )
    definition = replace(old_scope.definitions[0], source_header="")
    scope = DefinitionScope(table.id, table.name, table.selection, (definition,))
    cells = tuple(
        DateEvidence(
            DateSource(table.selection.workbook_sha256, row, 2, "Loans"), "text", "Fictional"
        )
        for row in range(4, 255)
    )
    source = build_snapshot(
        table, scope, (DateColumnEvidence(table.selection, 2, "1900", cells),), ()
    )
    preview = preview_transformation(source, OPERATION, ())
    assert len(preview.page()) == 200 and len(preview.page(200)) == 51
    assert preview.page(200)[-1]["row"] == 254
    assert preview.to_dict()["reconciliation"]["unchanged"] == 251


@pytest.mark.parametrize("mutation", ["version", "parent", "snapshot", "counts", "unknown"])
def test_revision_replay_rejects_changed_contract_or_evidence(prepared, mutation):
    source = snapshot(prepared)
    preview = preview_transformation(source, OPERATION, (proposal(),))
    revision = apply_transformation(source, preview, approval(preview))
    changed = revision.to_dict()
    if mutation == "version":
        changed["format_version"] = 99
    elif mutation == "parent":
        changed["steps"][0]["preview"]["parent_revision_id"] = "f" * 64
    elif mutation == "snapshot":
        changed["snapshot"]["cells"][0]["evidence"]["raw"] = "Forged original"
    elif mutation == "counts":
        changed["steps"][0]["preview"]["reconciliation"]["changed"] = 100
    else:
        changed["extra"] = "unrecognized"
    with pytest.raises(ValueError):
        replay_revision(source, canonical(changed))


def test_approval_requires_reviewer_time_and_unique_explicit_choices(prepared):
    preview = preview_transformation(snapshot(prepared), OPERATION, (proposal(),))
    with pytest.raises(ValueError, match="non-empty"):
        replace(approval(preview), reviewer="")
    with pytest.raises(ValueError, match="UTC"):
        replace(approval(preview), decided_at_utc="2026-10-18T12:00:00")
    with pytest.raises(ValueError, match="duplicate"):
        approval(preview, ((4, 2), (4, 2)))


def test_staging_collision_preserves_existing_file(prepared, tmp_path, monkeypatch):
    import loan_tape.transformation_excel as adapter

    class FixedId:
        hex = "a" * 32

    monkeypatch.setattr(adapter, "uuid4", FixedId)
    artifacts = tmp_path / "artifacts"
    directory = artifacts / "transformations"
    directory.mkdir(parents=True)
    staging = directory / (prepared[1].selection.workbook_sha256 + "-" + FixedId.hex + ".tmp")
    staging.write_bytes(b"existing staging evidence")
    preview = preview_transformation(snapshot(prepared), OPERATION, (proposal(),))
    with pytest.raises(FileExistsError):
        publish_transformation(*prepared, preview, approval(preview), artifact_root=artifacts)
    assert staging.read_bytes() == b"existing staging evidence"
    assert list(directory.glob("*.json")) == []
