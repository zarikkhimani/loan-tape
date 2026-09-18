"""Saved XML tokens survive extraction; malformed data cannot publish partial evidence."""

import hashlib
import re
from dataclasses import replace
from threading import Event
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904

from loan_tape.date_parser import DateProfile, parse_date
from loan_tape.date_reader import read_excel_date_column, read_excel_date_columns
from loan_tape.inspection import InspectionError
from loan_tape.ranges import parse_range
from loan_tape.selection import SourceSelection
from loan_tape.workbook_index import ScanCancelled


def selection(path, area="C9:C19", header=9):
    return SourceSelection(
        hashlib.sha256(path.read_bytes()).hexdigest(), "Dates", parse_range(area), header
    )


def rewrite(path, member, change):
    with ZipFile(path) as archive:
        content = {name: archive.read(name) for name in archive.namelist()}
    updated = change(content[member])
    assert updated != content[member], "Fixture mutation did not change the target XML"
    content[member] = updated
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        for name, data in content.items():
            archive.writestr(name, data)


def make_book(path):
    book = Workbook()
    sheet = book.active
    sheet.title = "Dates"
    sheet["C9"] = "Closing Date"
    sheet["C10"] = "1/8/21"
    sheet["C11"] = 60
    sheet["C11"].number_format = "mm/dd/yyyy"
    sheet["C12"] = 61.5
    sheet["C12"].number_format = "yyyy-mm-dd hh:mm:ss"
    sheet["C13"] = ""
    sheet["C14"].number_format = "yyyy-mm-dd"
    sheet["C15"] = "=TODAY()"
    sheet["C16"] = "#N/A"
    sheet["C17"] = False
    sheet["C18"] = " \t"
    sheet.row_dimensions[11].hidden = True
    sheet.column_dimensions["C"].hidden = True
    book.save(path)
    book.close()


def test_raw_serials_formats_gaps_tails_and_formula_cache_survive(tmp_path):
    path = tmp_path / "dates.xlsx"
    make_book(path)
    rewrite(
        path,
        "xl/worksheets/sheet1.xml",
        lambda s: re.sub(
            rb"<f>TODAY\(\)</f><v(?: />|></v>)", b"<f>TODAY()</f><v>45000</v>", s
        ).replace(b'<dimension ref="C9:C18"', b'<dimension ref="A1:A1"'),
    )
    before = path.read_bytes()
    report = read_excel_date_column(path, selection(path), 3)
    assert report.date_system == "1900"
    assert len(report.cells) == 10
    assert [c.source.row for c in report.cells] == list(range(10, 20))
    serial = report.cells[1]
    assert serial.raw == serial.raw_token == "60"
    assert serial.storage_type == "n" and serial.number_format == "mm/dd/yyyy"
    profile = DateProfile(
        text_formats=("M/D/YY",), two_digit_year_start=2000, numeric_dates="excel_serial"
    )
    assert parse_date(serial, profile).code == "excel_fictitious_1900_02_29"
    assert parse_date(report.cells[0], profile).parsed_date.isoformat() == "2021-01-08"
    assert parse_date(report.cells[2], profile).code == "timestamp_requires_policy"
    assert [parse_date(report.cells[i], profile).missing_kind for i in (3, 4, 8, 9)] == [
        "empty_text",
        "blank",
        "whitespace",
        "absent",
    ]
    formula = report.cells[5]
    assert formula.formula == "TODAY()" and formula.raw_token == "45000"
    assert parse_date(formula, profile).code == "formula_not_evaluated"
    assert parse_date(report.cells[6], profile).status == "source_error"
    assert parse_date(report.cells[7], profile).status == "unsupported"
    assert path.read_bytes() == before


def test_multiple_columns_are_returned_from_one_atomic_worksheet_scan(tmp_path, monkeypatch):
    import loan_tape.date_reader as reader

    path = tmp_path / "dates.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Dates"
    sheet["C9"] = "Closing Date"
    sheet["D9"] = "Maturity Date"
    sheet["C10"] = "2026-01-02"
    sheet["D10"] = "2027-02-03"
    sheet["C12"] = "2026-03-04"
    sheet["D12"] = "2027-04-05"
    book.save(path)
    book.close()
    chosen = selection(path, "C9:D12", 9)
    scans = 0
    scan = reader._columns

    def count_scans(*args, **kwargs):
        nonlocal scans
        scans += 1
        return scan(*args, **kwargs)

    monkeypatch.setattr(reader, "_columns", count_scans)
    reports = read_excel_date_columns(path, chosen, (4, 3))
    assert scans == 1
    assert [report.column for report in reports] == [4, 3]
    assert [cell.raw for cell in reports[0].cells] == ["2027-02-03", None, "2027-04-05"]
    assert [cell.raw for cell in reports[1].cells] == ["2026-01-02", None, "2026-03-04"]
    assert all(report.selection == chosen and report.date_system == "1900" for report in reports)


def test_1904_system_is_read_from_workbook_and_zero_is_a_value(tmp_path):
    path = tmp_path / "mac.xlsx"
    book = Workbook()
    book.epoch = CALENDAR_MAC_1904
    book.active.title = "Dates"
    book.active["C9"] = 0
    book.active["C9"].number_format = "yyyy-mm-dd"
    book.save(path)
    report = read_excel_date_column(path, selection(path, "C9", None), 3)
    assert report.date_system == "1904" and report.cells[0].raw == "0"
    result = parse_date(report.cells[0], DateProfile(numeric_dates="excel_serial"))
    assert result.parsed_date.isoformat() == "1904-01-01"


def test_explicit_iso_cell_retains_timestamp_before_policy(tmp_path):
    path = tmp_path / "iso.xlsx"
    make_book(path)
    rewrite(
        path,
        "xl/worksheets/sheet1.xml",
        lambda s: s.replace(
            b'<c r="C10" t="inlineStr"><is><t>1/8/21</t></is></c>',
            b'<c r="C10" t="d"><v>2026-01-01T00:30:00+14:00</v></c>',
        ),
    )
    evidence = read_excel_date_column(path, selection(path, "C10", None), 3).cells[0]
    assert evidence.kind == "iso_date" and evidence.raw.endswith("+14:00")
    assert parse_date(evidence, DateProfile()).status == "unsupported"
    assert (
        parse_date(evidence, DateProfile(timestamps="date_component")).parsed_date.isoformat()
        == "2026-01-01"
    )


def test_shared_rich_text_excludes_phonetic_runs_and_keeps_index(tmp_path):
    path = tmp_path / "strings.xlsx"
    make_book(path)
    rewrite(
        path,
        "xl/worksheets/sheet1.xml",
        lambda s: s.replace(
            b'<c r="C10" t="inlineStr"><is><t>1/8/21</t></is></c>', b'<c r="C10" t="s"><v>0</v></c>'
        ),
    )
    rewrite(
        path,
        "xl/_rels/workbook.xml.rels",
        lambda s: s.replace(
            b"</Relationships>",
            b'<Relationship Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings" Target="sharedStrings.xml" Id="rId99"/></Relationships>',
        ),
    )
    with ZipFile(path, "a", ZIP_DEFLATED) as archive:
        archive.writestr(
            "xl/sharedStrings.xml",
            '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><si><r><t>1/</t></r><rPh sb="0" eb="2"><t>not part of value</t></rPh><r><t>8/21</t></r></si></sst>',
        )
    value = read_excel_date_column(path, selection(path, "C10", None), 3).cells[0]
    assert value.raw == "1/8/21" and value.raw_token == "0" and value.storage_type == "s"


@pytest.mark.parametrize(
    "member,change,match",
    [
        (
            "xl/workbook.xml",
            lambda s: re.sub(rb"<workbookPr\s*/>", b'<workbookPr date1904="maybe"/>', s),
            "date system",
        ),
        (
            "xl/worksheets/sheet1.xml",
            lambda s: s.replace(b'<c r="C10"', b'<c s="9999" r="C10"'),
            "style",
        ),
        (
            "xl/worksheets/sheet1.xml",
            lambda s: s.replace(
                b'<c r="C10" t="inlineStr"><is><t>1/8/21</t></is></c>',
                b'<c r="C10" t="s"><v>99</v></c>',
            ),
            "shared string",
        ),
        (
            "xl/worksheets/sheet1.xml",
            lambda s: s.replace(b'<row r="11"', b'<row r="10"'),
            "duplicated worksheet rows",
        ),
        (
            "xl/worksheets/sheet1.xml",
            lambda s: s.replace(b'<c r="C10" t="inlineStr">', b'<c r="C9" t="inlineStr">'),
            "coordinates",
        ),
        (
            "xl/worksheets/sheet1.xml",
            lambda s: s.replace(b"</worksheet>", b"<bad></worksheet>"),
            "Could not read complete",
        ),
        (
            "xl/worksheets/sheet1.xml",
            lambda s: b'<!DOCTYPE worksheet [<!ENTITY x "x">]>' + s,
            "Could not read complete",
        ),
    ],
)
def test_unreadable_evidence_fails_instead_of_returning_partial_rows(
    tmp_path, member, change, match
):
    path = tmp_path / "broken.xlsx"
    make_book(path)
    rewrite(path, member, change)
    before = path.read_bytes()
    with pytest.raises(InspectionError, match=match):
        read_excel_date_column(path, selection(path), 3)
    assert path.read_bytes() == before


def test_source_identity_scope_limits_and_cancellation(tmp_path):
    path = tmp_path / "dates.xlsx"
    make_book(path)
    chosen = selection(path)
    with pytest.raises(InspectionError, match="identity changed"):
        read_excel_date_column(path, replace(chosen, workbook_sha256="a" * 64), 3)
    with pytest.raises(InspectionError, match="inside"):
        read_excel_date_column(path, chosen, 4)
    with pytest.raises(InspectionError, match="at least one"):
        read_excel_date_columns(path, chosen, ())
    with pytest.raises(InspectionError, match="only once"):
        read_excel_date_columns(path, chosen, (3, 3))
    with pytest.raises(InspectionError, match="row limit"):
        read_excel_date_column(path, chosen, 3, max_rows=5)
    cancel = Event()
    cancel.set()
    with pytest.raises(ScanCancelled):
        read_excel_date_column(path, chosen, 3, cancelled=cancel)
    empty = read_excel_date_column(path, selection(path, "C9", 9), 3)
    assert empty.cells == ()


def test_mid_read_source_change_is_not_published(tmp_path, monkeypatch):
    import loan_tape.date_reader as reader

    path = tmp_path / "dates.xlsx"
    make_book(path)
    read = reader._columns

    def change(*args, **kwargs):
        result = read(*args, **kwargs)
        with path.open("ab") as output:
            output.write(b"changed")
        return result

    monkeypatch.setattr(reader, "_columns", change)
    with pytest.raises(InspectionError, match="Source changed"):
        read_excel_date_column(path, selection(path), 3)


def test_merged_dates_are_not_silently_assigned_to_one_record(tmp_path):
    path = tmp_path / "merged.xlsx"
    book = Workbook()
    book.active.title = "Dates"
    book.active.merge_cells("C10:C11")
    book.active["C10"] = "2026-01-01"
    book.save(path)
    with pytest.raises(InspectionError, match="merged"):
        read_excel_date_column(path, selection(path, "C10:C11", None), 3)


def test_legacy_files_rejected_without_opening_excel(tmp_path):
    path = tmp_path / "old.xls"
    path.write_bytes(b"legacy")
    with pytest.raises(InspectionError, match="xlsx"):
        read_excel_date_column(path, selection(path), 3)


@pytest.mark.parametrize(
    "old,new",
    [
        (b"sheetData", b"unknownData"),
        (b"</sheetData>", b"</sheetData><sheetData/>"),
    ],
)
def test_missing_or_duplicate_data_blocks_cannot_be_reported_as_blank(tmp_path, old, new):
    path = tmp_path / "blocks.xlsx"
    make_book(path)
    rewrite(path, "xl/worksheets/sheet1.xml", lambda value: value.replace(old, new))
    with pytest.raises(InspectionError, match="data block"):
        read_excel_date_column(path, selection(path), 3)


def test_raw_serial_precision_survives_reader(tmp_path):
    path = tmp_path / "fraction.xlsx"
    make_book(path)
    raw = b"61.999999999999999999999999999999"
    rewrite(
        path,
        "xl/worksheets/sheet1.xml",
        lambda s: s.replace(b"<v>61.5</v>", b"<v>" + raw + b"</v>"),
    )
    cell = read_excel_date_column(path, selection(path, "C12", None), 3).cells[0]
    result = parse_date(
        cell, DateProfile(numeric_dates="excel_serial", timestamps="date_component")
    )
    assert cell.raw == raw.decode("ascii")
    assert result.parsed_date.isoformat() == "1900-03-01"


def test_cancellation_during_read_never_returns_partial_result(tmp_path, monkeypatch):
    import loan_tape.date_reader as reader

    path = tmp_path / "cancel.xlsx"
    make_book(path)
    cancel = Event()
    read = reader._cell

    def cancel_after_first(*args, **kwargs):
        result = read(*args, **kwargs)
        cancel.set()
        return result

    monkeypatch.setattr(reader, "_cell", cancel_after_first)
    with pytest.raises(ScanCancelled):
        read_excel_date_column(path, selection(path), 3, cancelled=cancel)
