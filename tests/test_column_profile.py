"""Full saved-range coverage and raw distinctions for column inspection."""

import hashlib
import re
from dataclasses import replace
from datetime import UTC, datetime
from threading import Event
from unittest.mock import patch
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook, load_workbook

from loan_tape.column_profile import ColumnCancelled, inspect_column
from loan_tape.inspection import InspectionError, read_preview
from loan_tape.ranges import parse_range
from loan_tape.selection import SavedTable, SourceSelection
from loan_tape.session import AnalysisSession

SESSION = AnalysisSession("session", datetime(2026, 9, 16, tzinfo=UTC))


def table_for(path, area, header=None, sheet="Loans", name="Loans"):
    sha = hashlib.sha256(path.read_bytes()).hexdigest()
    return SavedTable("1" * 32, name, SourceSelection(sha, sheet, parse_range(area), header))


def rewrite_sheet(path, change):
    with ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    parts["xl/worksheets/sheet1.xml"] = change(parts["xl/worksheets/sheet1.xml"])
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        for name, value in parts.items():
            archive.writestr(name, value)


@pytest.mark.parametrize("extension", [".xlsx", ".xlsm"])
def test_full_distant_column_keeps_blanks_zeroes_types_formulas_and_last_row(
    tmp_path, monkeypatch, extension
):
    monkeypatch.chdir(tmp_path)
    path = tmp_path / ("data" + extension)
    book = Workbook()
    sheet = book.active
    sheet.title = "Loans"
    sheet["DZ450"] = "ID"
    values = {
        451: "000123",
        452: "000123",
        453: 123,
        454: 0,
        455: "0",
        456: False,
        457: True,
        458: datetime(2035, 1, 2),
        459: "01/02/2030",
        460: "",
        461: "   ",
        462: "N/A",
        463: "#N/A",
        464: "=1+1",
        465: "=1+1",
        900: "Late",
        949: "Late",
        950: "TailOnly",
    }
    for row, value in values.items():
        sheet.cell(row, 130, value)
    sheet["EA950"] = "Other column"
    sheet["DZ951"] = "Outside range"
    sheet.row_dimensions[950].hidden = True
    sheet.column_dimensions["DZ"].hidden = True
    sheet.sheet_state = "hidden"
    book.create_sheet("Visible")
    book.save(path)
    book.close()
    rewrite_sheet(
        path, lambda xml: re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1:A1"', xml)
    )
    table = table_for(path, "DZ450:EA950", 450)
    before = path.read_bytes()
    with patch("loan_tape.column_profile.load_workbook", wraps=load_workbook) as reader:
        profile = inspect_column(path, table, 130, SESSION)
    assert reader.call_args.kwargs == {"read_only": True, "data_only": False, "keep_links": False}
    assert profile.total_rows == 500 and profile.header == "ID"
    assert profile.data_address == "DZ451:DZ950"
    assert profile.filled_cells == 17
    assert dict(profile.counts) == {
        "Blank": 482,
        "Empty text": 1,
        "Text": 9,
        "Number": 2,
        "Date/time": 1,
        "Boolean": 2,
        "Excel error": 1,
        "Formula": 2,
    }
    assert profile.numeric_zeroes == 1 and profile.whitespace_text == 1
    assert profile.distinct_values == 14
    assert profile.repeated_values == 3 and profile.extra_occurrences == 3
    assert {(x.text, x.count, x.row) for x in profile.repeated} == {
        ("000123", 2, 451),
        ("=1+1", 2, 464),
        ("Late", 2, 900),
    }
    assert sum(count for _, count in profile.counts) == profile.total_rows
    assert profile.table == table and profile.session is SESSION
    assert path.read_bytes() == before
    assert list((tmp_path / ".artifacts/column-inspection").iterdir()) == []
    path.rename(tmp_path / ("closed" + extension))


def test_header_excludes_preamble_and_can_leave_zero_data_rows(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    path = tmp_path / "headers.xlsx"
    book = Workbook()
    book.active.title = "Loans"
    book.active["A1"] = "Title"
    book.active["A3"] = "Header"
    book.active["A4"] = "Value"
    book.save(path)
    book.close()
    without = inspect_column(path, table_for(path, "A1:A5"), 1, SESSION)
    assert without.total_rows == 5 and without.header is None
    assert dict(without.counts)["Text"] == 3
    after = inspect_column(path, table_for(path, "A1:A5", 3), 1, SESSION)
    assert after.total_rows == 2 and dict(after.counts)["Blank"] == 1
    empty = inspect_column(path, table_for(path, "A1:A3", 3), 1, SESSION)
    assert empty.total_rows == 0 and empty.header == "Header"
    assert not empty.examples and not empty.repeated


def test_counting_reaches_full_range_with_many_distinct_values(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    path = tmp_path / "long.xlsx"
    book = Workbook()
    book.active.title = "Loans"
    book.active.append(["ID"])
    for row in range(2, 20002):
        book.active.append([f"{row:06}"])
    book.active.append(["000002"])
    book.save(path)
    book.close()
    profile = inspect_column(path, table_for(path, "A1:A20002", 1), 1, SESSION)
    assert profile.total_rows == profile.filled_cells == 20001
    assert profile.distinct_values == 20000
    assert profile.extra_occurrences == 1
    assert profile.repeated[0].text == "000002" and profile.repeated[0].count == 2
    assert len(profile.examples) == 3  # Example limit never limits population counts.


def test_blank_tail_and_maximum_excel_coordinates_are_not_truncated(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    path = tmp_path / "edge.xlsx"
    book = Workbook()
    book.active.title = "Loans"
    book.active["XFD1048574"] = "Header"
    book.active["XFD1048575"] = 0
    book.save(path)
    book.close()
    profile = inspect_column(
        path, table_for(path, "XFD1048574:XFD1048576", 1048574), 16384, SESSION
    )
    assert profile.total_rows == 2 and profile.numeric_zeroes == 1
    assert dict(profile.counts)["Blank"] == 1
    assert profile.examples[-1].row == 1048576


def test_cancellation_and_failed_reads_never_return_partial_counts(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    path = tmp_path / "cancel.xlsx"
    book = Workbook()
    book.active.title = "Loans"
    book.active["A1"] = "Header"
    book.active["A30000"] = "last"
    book.save(path)
    book.close()
    table = table_for(path, "A1:A30000", 1)
    cancelled = Event()

    def progress(message):
        if "Read 10,000" in message:
            cancelled.set()

    with pytest.raises(ColumnCancelled):
        inspect_column(path, table, 1, SESSION, cancelled=cancelled, progress=progress)
    assert list((tmp_path / ".artifacts/column-inspection").iterdir()) == []
    with pytest.raises(ColumnCancelled):
        inspect_column(path, table, 1, SESSION, cancelled=cancelled)
    rewrite_sheet(
        path,
        lambda xml: xml.replace(b"inlineStr", b"n", 1).replace(
            b"<is><t>Header</t></is>", b"<v>not-a-number</v>", 1
        ),
    )
    with pytest.raises(ValueError):
        inspect_column(path, table_for(path, "A1:A30000", 1), 1, SESSION)
    assert list((tmp_path / ".artifacts/column-inspection").iterdir()) == []
    path.rename(tmp_path / "no-open-handles.xlsx")


def test_changed_source_invalid_sheet_and_outside_columns_fail_clearly(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    path = tmp_path / "source.xlsx"
    book = Workbook()
    book.active.title = "Loans"
    book.active["B2"] = "Header"
    book.active["B3"] = "value"
    book.save(path)
    book.close()
    table = table_for(path, "B2:B3", 2)
    with pytest.raises(InspectionError, match="inside"):
        inspect_column(path, table, 1, SESSION)
    missing = replace(table, selection=replace(table.selection, sheet="Missing"))
    with pytest.raises(InspectionError, match="unavailable"):
        inspect_column(path, missing, 2, SESSION)
    wrong_hash = replace(table, selection=replace(table.selection, workbook_sha256="f" * 64))
    with pytest.raises(InspectionError, match="changed"):
        inspect_column(path, wrong_hash, 2, SESSION)

    def modify_at_end(message):
        if message.startswith("Verifying"):
            with path.open("ab") as handle:
                handle.write(b"changed")

    with pytest.raises(InspectionError, match="changed"):
        inspect_column(path, table, 2, SESSION, progress=modify_at_end)


def test_unreadable_date_is_not_counted_as_a_source_error(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    path = tmp_path / "date.xlsx"
    book = Workbook()
    book.active.title = "Loans"
    book.active["A1"] = datetime(2030, 1, 1)
    book.save(path)
    book.close()
    rewrite_sheet(path, lambda xml: re.sub(rb"<v>[0-9]+</v>", b"<v>999999999</v>", xml))
    with pytest.raises(InspectionError, match="date could not be decoded"):
        inspect_column(path, table_for(path, "A1:A1"), 1, SESSION)


def test_preview_can_show_a_blank_example_beyond_stored_rows(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    path = tmp_path / "blank-example.xlsx"
    book = Workbook()
    book.active.title = "Loans"
    book.active["DZ5"] = "Header"
    book.active["DZ6"] = "Value"
    book.save(path)
    book.close()
    profile = inspect_column(path, table_for(path, "DZ5:EA100", 5), 130, SESSION)
    blank = next(item for item in profile.examples if item.kind == "Blank")
    assert blank.row == 7
    page = read_preview(path, sheet="Loans", area=parse_range("DZ7:EA26"))
    assert page.start_row == 7 and page.start_column == 130
    assert page.column_count == 2 and len(page.rows) == 20
    assert all(cell.text == "" for row in page.rows for cell in row)
