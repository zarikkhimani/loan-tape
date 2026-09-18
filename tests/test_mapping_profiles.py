"""Exact, source-preserving coverage for reusable workbook mapping profiles."""

import hashlib
import json
from pathlib import Path
from uuid import uuid4

import pytest
from openpyxl import Workbook

from loan_tape.mapping_profiles import (
    MappingProfileError,
    load_builtin_profiles,
    load_mapping_profile,
    match_builtin_table_profile,
    match_headers,
)
from loan_tape.ranges import CellRange
from loan_tape.selection import SavedTable, SourceSelection


def _warehouse_headers() -> tuple[str, ...]:
    return tuple(item.source_header for item in load_builtin_profiles()[0].columns)


def _saved_table(path: Path, *, header_row: int | None = 5) -> SavedTable:
    sha256 = hashlib.sha256(path.read_bytes()).hexdigest()
    return SavedTable(
        uuid4().hex,
        "Portfolio",
        SourceSelection(sha256, "Inputs and Portfolio", CellRange(5, 2, 6, 95), header_row),
    )


def test_builtin_warehouse_profile_preserves_reviewed_layout() -> None:
    (profile,) = load_builtin_profiles()

    assert (profile.id, profile.version, len(profile.columns)) == (
        "warehouse-model",
        "1.0.0",
        94,
    )
    assert profile.columns[3].source_header == (
        "Indentifier (CUSIP, ISIN or Internal Indentifier Code)"
    )
    assert profile.columns[7].source_header == "Notional "
    assert [item.occurrence for item in profile.columns if item.source_header == "Aaa"] == [1, 2, 3]
    assert profile.columns[10].role == "calculated"
    assert profile.columns[18].role == "mixed"
    assert profile.columns[88].presence == "optional"


def test_header_match_is_complete_exact_and_position_bound() -> None:
    profile = load_builtin_profiles()[0]
    headers = _warehouse_headers()

    match = match_headers(profile, headers, start_column=2)

    assert match is not None
    assert match.description == (
        "Mapping profile: Warehouse Model 1.0.0 · 94 columns matched exactly."
    )
    assert match.bindings[0].source_column == 2
    assert match.bindings[-1].source_column == 95
    changed = list(headers)
    changed[7] = changed[7].strip()
    assert match_headers(profile, tuple(changed), start_column=2) is None
    assert match_headers(profile, headers[:-1], start_column=2) is None


@pytest.mark.parametrize(
    ("change", "message"),
    [
        ({"occurrence": 2}, "invalid occurrence"),
        ({"unexpected": True}, "unsupported keys"),
    ],
)
def test_profile_loader_rejects_ambiguous_or_unknown_content(
    change: dict[str, object], message: str
) -> None:
    document = {
        "format_version": 1,
        "id": "test-profile",
        "version": "1.0.0",
        "name": "Test",
        "description": "Test profile",
        "columns": [
            {
                "position": 1,
                "key": "loan_id",
                "source_header": "Loan ID",
                "occurrence": 1,
                "expected_type": "identifier",
                "role": "source",
                "presence": "required",
                **change,
            }
        ],
    }

    with pytest.raises(MappingProfileError, match=message):
        load_mapping_profile(json.dumps(document).encode(), "test profile")


def test_saved_table_match_reads_full_header_without_changing_workbook(tmp_path: Path) -> None:
    path = tmp_path / "warehouse.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Inputs and Portfolio"
    for column, header in enumerate(_warehouse_headers(), 2):
        sheet.cell(5, column, header)
        sheet.cell(6, column, 0)
    book.save(path)
    book.close()
    original = path.read_bytes()
    table = _saved_table(path)

    match = match_builtin_table_profile(path, table)

    assert match is not None
    assert match.bindings[0].source_column == 2
    assert match.bindings[-1].source_column == 95
    assert path.read_bytes() == original


def test_saved_table_match_requires_a_saved_header(tmp_path: Path) -> None:
    path = tmp_path / "warehouse.xlsx"
    book = Workbook()
    book.active.title = "Inputs and Portfolio"
    book.save(path)
    book.close()

    assert match_builtin_table_profile(path, _saved_table(path, header_row=None)) is None
