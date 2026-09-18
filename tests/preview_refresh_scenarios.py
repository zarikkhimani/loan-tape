"""Exercise persistent context, saved setup, narrow layouts, and retry in real Tk."""

import sys
import time
from pathlib import Path
from threading import Event
from unittest.mock import patch

from desktop_scenarios import wait_for_copy
from desktop_scenarios import wait_for_preview as wait_hidden
from openpyxl import Workbook

from loan_tape.inspection import read_preview
from loan_tape.intake import IntakeStore
from loan_tape.ui import LoanTapeApp, create_root


def wait_for_preview(app):
    wait_hidden(app)
    app.root.deiconify()
    app.root.update()


def visible_inside(root, widget):
    assert widget.winfo_ismapped(), str(widget)
    assert root.winfo_rootx() <= widget.winfo_rootx()
    assert widget.winfo_rootx() + widget.winfo_width() <= root.winfo_rootx() + root.winfo_width(), (
        str(widget),
        widget.winfo_geometry(),
    )
    assert (
        widget.winfo_rooty() + widget.winfo_height() <= root.winfo_rooty() + root.winfo_height()
    ), (str(widget), widget.winfo_geometry())


def run(scenario, scale, work):
    root = create_root()
    root.withdraw()
    root.tk.call("tk", "scaling", scale)
    errors = []
    root.report_callback_exception = lambda *args: errors.append(args)
    app = LoanTapeApp(root, IntakeStore(work / "inputs"))
    release = Event()
    try:
        if scenario == "workbook":
            source = work / ("A long but meaningful portfolio filename " * 3 + ".xlsx")
            book = Workbook()
            for name in ("Loans", "Facilities"):
                sheet = book.active if name == "Loans" else book.create_sheet()
                sheet.title = name
                sheet.append(["Identifier", "Balance"])
                for number in range(50):
                    sheet.append([f"{number:06}", number * 100])
            book.save(source)
            book.close()
        elif scenario == "xml":
            source = work / "Portfolio.xml"
            source.write_text(
                "<loans>"
                + "".join(
                    f'<loan id="{row:06}"><balance>{row}</balance></loan>' for row in range(60)
                )
                + "</loans>"
            )
        else:
            source = work / "Portfolio.csv"
            source.write_text("id,balance\n000123,0\n")
        original = source.read_bytes()
        app.add_file(source)
        wait_for_copy(app)
        with patch("loan_tape.preview_ui.read_parallel_preview", side_effect=read_preview):
            app.inspect_file()
            wait_for_preview(app)
            panel = app.inspector
            root.geometry("760x720+0+0")
            root.deiconify()
            root.update()
            if scenario == "workbook":
                nav = panel.navigator
                assert not panel.sidebar_collapsed
                nav.header_text.set("1")
                nav.use_button.invoke()
                wait_for_preview(app)
                root.update()
                assert panel.sidebar_collapsed and not panel.rail.winfo_ismapped()
                for widget in (
                    panel.sheet_selector,
                    panel.data_selector,
                    panel.setup_button,
                    panel.zoom_selector,
                    nav.end_button,
                ):
                    visible_inside(root, widget)
                assert "…" in panel.filename_label.cget("text")
                assert panel.column_actions.winfo_ismapped()
                assert panel.column_button.instate(["disabled"])
                assert panel.workflow_hint.get() == "Next: select a column heading."
                panel._select_column(0, entire=True)
                root.update()
                assert panel.workflow_hint.get() == "Ready: choose a review action."
                visible_inside(root, panel.definition_button)
                visible_inside(root, panel.column_button)
                visible_inside(root, panel.check_button)
                nav.down.invoke()
                wait_for_preview(app)
                assert panel.preview.start_row == 21
                assert panel.whole_column is not None
                panel.sheet_selector.set("Facilities")
                panel.sheet_selector.event_generate("<<ComboboxSelected>>")
                wait_for_preview(app)
                assert not panel.sidebar_collapsed
                nav.table_name.set("Facilities saved")
                nav.use_button.invoke()
                wait_for_preview(app)
                panel.data_selector.current(0)
                panel.data_selector.event_generate("<<ComboboxSelected>>")
                wait_for_preview(app)
                assert panel.sheet_name.get() == "Loans" and panel.sidebar_collapsed
                panel.setup_button.invoke()
                root.update()
                nav.edit_button.invoke()
                nav.table_name.set("Edited name")
                assert nav.has_unsaved_setup()
                panel._toggle_sidebar()
                root.update()
                assert panel.sidebar_collapsed
                assert panel.setup_button.cget("text") == "Unsaved setup"
                with patch("loan_tape.preview_ui.messagebox.askokcancel", return_value=False):
                    panel.sheet_selector.set("Facilities")
                    panel.sheet_selector.event_generate("<<ComboboxSelected>>")
                    root.update()
                    assert panel.sheet_name.get() == "Loans"
                    panel.data_selector.current(1)
                    panel.data_selector.event_generate("<<ComboboxSelected>>")
                    root.update()
                    assert nav.table_name.get() == "Edited name"
                # Closing still guards hidden drafts, and Cancel preserves the exact setup.
                with patch("loan_tape.ui.messagebox.askokcancel", return_value=False):
                    app.close_preview()
                assert app.inspector is panel and nav.table_name.get() == "Edited name"
                nav.use_range()
                wait_for_preview(app)
                app.close_preview()
                app.inspect_file()
                wait_for_preview(app)
                assert app.inspector.sidebar_collapsed
                assert app.inspector.navigator.table_name.get() == "Edited name"
            elif scenario == "xml":
                nav = panel.xml_navigation
                nav.name.set("Saved loans")
                nav.save_button.invoke()
                root.update()
                assert not nav.editor.winfo_ismapped()
                panel._select_column(0, entire=True)
                root.update()
                for widget in (
                    panel.xml_group_selector,
                    nav.saved_selector,
                    nav.edit_button,
                    nav.buttons["Go"],
                    panel.check_button,
                ):
                    visible_inside(root, widget)
                nav.buttons["End"].invoke()
                wait_for_preview(app)
                assert panel.preview.start_row == 41
                nav.edit_button.invoke()
                root.update()
                visible_inside(root, nav.name_entry)
                visible_inside(root, nav.save_button)
                nav.name.set("Renamed loans")
                nav.save_button.invoke()
                root.update()
                assert not nav.editor.winfo_ismapped()
                nav.new_button.invoke()
                root.update()
                assert nav.editor.winfo_ismapped()
            else:
                assert panel.preview.rows[1][0].text == "000123"
                for widget in (*panel.selectors, panel.zoom_selector):
                    visible_inside(root, widget)

                def paused(*args, **kwargs):
                    assert release.wait(5)
                    return read_preview(*args, **kwargs)

                with patch("loan_tape.preview_ui.read_parallel_preview", side_effect=paused):
                    panel.load()
                    root.update()
                    assert panel.busy and panel.grid_state.winfo_ismapped()
                    assert not panel.retry_button.winfo_ismapped()
                    app.tabs.select(app.files_page)
                    root.update()
                    release.set()
                    wait_for_preview(app)
                    assert str(app.tabs.select()) == str(app.files_page)
                    app.activate_preview()
                with patch(
                    "loan_tape.preview_ui.read_parallel_preview", side_effect=PermissionError
                ):
                    panel.load()
                    wait_for_preview(app)
                    root.update()
                    visible_inside(root, panel.retry_button)
                panel.retry_button.invoke()
                wait_for_preview(app)
                root.update()
                assert not panel.grid_state.winfo_ismapped()
                assert panel.preview.rows[1][0].text == "000123"
                empty = work / "Empty.xlsx"
                empty_book = Workbook()
                empty_book.save(empty)
                empty_book.close()
                app.add_file(empty)
                wait_for_copy(app)
                app.inspect_file()
                wait_for_preview(app)
                root.update()
                assert app.inspector.state_title.cget("text") == "No data to display"
                assert not app.inspector.retry_button.winfo_ismapped()
        assert source.read_bytes() == original
        assert not errors, errors
    finally:
        release.set()
        with patch("loan_tape.ui.messagebox.askokcancel", return_value=True):
            app.close()
        deadline = time.monotonic() + 10
        while not app.closed:
            root.update()
            assert time.monotonic() < deadline
            time.sleep(0.01)


if __name__ == "__main__":
    run(sys.argv[1], float(sys.argv[2]), Path(sys.argv[3]))
