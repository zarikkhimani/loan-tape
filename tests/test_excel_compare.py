"""Independent-reader comparisons and preflight of the complete workbook."""

import importlib
import sys
from dataclasses import replace
from pathlib import Path
from threading import Barrier, Event, Timer
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName

from loan_tape import excel_compare
from loan_tape.excel_compare import (
    ComparedPreview,
    ComparisonSkipped,
    check_excel_input,
    compare_previews,
    read_parallel_preview,
)
from loan_tape.inspection import FilePreview, InspectionError, PreviewCell, read_preview


def plain_book(path):
    book = Workbook()
    book.active.append(["ID", "Amount", "Blank", "Missing", "Literal"])
    book.active.append(["000123", 0, None, "NA", "=1+1"])
    book.active["E2"].data_type = "s"
    book.save(path)
    book.close()


def test_plain_book_preserves_formula_like_text_and_source_bytes(tmp_path):
    path = tmp_path / "data.xlsx"
    plain_book(path)
    before = path.read_bytes()
    check_excel_input(path)
    preview = read_preview(path)
    assert [cell.text for cell in preview.rows[1]] == ["000123", "0", "", "NA", "=1+1"]
    assert path.read_bytes() == before


@pytest.mark.parametrize("kind", ["distant_formula", "hidden_formula", "defined_name"])
def test_executable_content_outside_preview_is_rejected(tmp_path, kind):
    path = tmp_path / "data.xlsx"
    book = Workbook()
    book.active["A1"] = "Plain visible data"
    if kind == "distant_formula":
        book.active["A1000"] = "=1+1"
    elif kind == "hidden_formula":
        hidden = book.create_sheet("Hidden")
        hidden.sheet_state = "veryHidden"
        hidden["A1"] = "=1+1"
    else:
        book.defined_names.add(DefinedName("Calculation", attr_text="1+1"))
    book.save(path)
    book.close()
    with pytest.raises(ComparisonSkipped, match="formulas or named"):
        check_excel_input(path)


@pytest.mark.parametrize(
    "part", ["xl/vbaProject.bin", "xl/connections.xml", "xl/macrosheets/sheet1.xml"]
)
def test_features_that_could_execute_or_refresh_are_rejected(tmp_path, part):
    path = tmp_path / "data.xlsx"
    plain_book(path)
    with ZipFile(path, "a", ZIP_DEFLATED) as archive:
        archive.writestr(part, b"synthetic placeholder")
    with pytest.raises(ComparisonSkipped, match="beyond plain data"):
        check_excel_input(path)


def test_external_relationship_is_rejected(tmp_path):
    path = tmp_path / "data.xlsx"
    plain_book(path)
    with ZipFile(path) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    entries["_rels/.rels"] = entries["_rels/.rels"].replace(
        b"</Relationships>",
        b'<Relationship Id="synthetic" TargetMode="External" Target="https://example.invalid/" Type="external"/></Relationships>',
    )
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        for name, content in entries.items():
            archive.writestr(name, content)
    with pytest.raises(ComparisonSkipped, match="external links"):
        check_excel_input(path)


def test_excel_is_not_started_for_rejected_source(tmp_path, monkeypatch):
    path = tmp_path / "data.xlsx"
    book = Workbook()
    book.active["A300"] = "=1+1"
    book.save(path)
    book.close()
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(excel_compare.sys, "platform", "win32")

    def forbidden(*args, **kwargs):
        pytest.fail("Excel worker must not start for a rejected source")

    monkeypatch.setattr(excel_compare, "start_excel", forbidden)
    with pytest.raises(ComparisonSkipped):
        excel_compare.read_xlwings_preview(path)
    assert not list((tmp_path / ".artifacts/excel-comparison").iterdir())


def test_readers_run_concurrently_and_identify_cell_differences(monkeypatch):
    barrier = Barrier(2, timeout=3)
    primary = FilePreview(
        rows=((PreviewCell("000123"), PreviewCell("0")),),
        column_count=2,
        more_rows=False,
        more_columns=False,
    )
    secondary = replace(primary, rows=((PreviewCell("123"), PreviewCell("0")),))

    def first(*args, **kwargs):
        barrier.wait()
        return primary

    def second(*args, **kwargs):
        barrier.wait()
        return secondary

    monkeypatch.setattr(excel_compare, "read_preview", first)
    monkeypatch.setattr(excel_compare, "read_xlwings_preview", second)
    result = read_parallel_preview(Path("data.xlsx"))
    assert isinstance(result, ComparedPreview)
    assert result.differences == {(0, 0)}
    assert "1 cell value differences" in result.message
    assert result.primary.rows[0][0].text == "000123"


def test_unavailable_excel_keeps_primary_with_explicit_message(tmp_path, monkeypatch):
    path = tmp_path / "data.xlsx"
    plain_book(path)

    def unavailable(*args, **kwargs):
        raise InspectionError("Excel is unavailable for this test.")

    monkeypatch.setattr(excel_compare, "read_xlwings_preview", unavailable)
    result = read_parallel_preview(path)
    assert isinstance(result, ComparedPreview)
    assert result.primary.rows[1][0].text == "000123"
    assert result.secondary is None
    assert result.message == "Excel is unavailable for this test."


def test_csv_never_calls_excel(tmp_path, monkeypatch):
    path = tmp_path / "data.csv"
    path.write_text("ID,Value\n000123,0\n")

    def forbidden(*args, **kwargs):
        pytest.fail("CSV must keep the text-preserving reader")

    monkeypatch.setattr(excel_compare, "read_xlwings_preview", forbidden)
    result = read_parallel_preview(path)
    assert isinstance(result, FilePreview)
    assert result.rows[1][0].text == "000123"


def test_cancelled_request_starts_no_process(monkeypatch):
    cancelled = Event()
    cancelled.set()
    monkeypatch.setattr(excel_compare.sys, "platform", "win32")
    with pytest.raises(InspectionError, match="cancelled"):
        excel_compare.read_xlwings_preview(Path("absent.xlsx"), cancelled=cancelled)


def test_comparison_reports_different_sheet_lists():
    primary = FilePreview(
        rows=(), column_count=0, more_rows=False, more_columns=False, sheets=("One",)
    )
    other = replace(primary, sheets=("Two",))
    assert "different worksheet lists" in compare_previews(primary, other).message


@pytest.mark.skipif(sys.platform != "win32", reason="Windows worker lifecycle integration")
@pytest.mark.parametrize("action", ["timeout", "cancel"])
def test_stalled_worker_is_terminated_and_temporary_copy_removed(tmp_path, monkeypatch, action):
    path = tmp_path / "data.xlsx"
    plain_book(path)
    before = path.read_bytes()
    monkeypatch.chdir(tmp_path)
    real_start = excel_compare.start_in_job
    api = importlib.import_module("win32api")
    events = importlib.import_module("win32event")
    owned = []

    def stall(executable, arguments, job):
        process, pid = real_start(sys.executable, ["-c", "import time; time.sleep(60)"], job)
        owned.append(api.OpenProcess(0x00100000, False, pid))  # SYNCHRONIZE
        return process, pid

    # Both helper processes stay in the real owned job, without requiring Office.
    monkeypatch.setattr(excel_compare, "start_in_job", stall)
    monkeypatch.setattr(excel_compare, "start_excel", lambda job: stall("", [], job))
    cancelled = Event()
    timer = Timer(0.3, cancelled.set)
    timer.start()
    try:
        with pytest.raises(
            InspectionError, match="timed out" if action == "timeout" else "cancelled"
        ):
            excel_compare.read_xlwings_preview(
                path,
                cancelled=cancelled if action == "cancel" else None,
                timeout=0.1 if action == "timeout" else 5,
            )
    finally:
        timer.cancel()
        for handle in owned:
            try:
                assert events.WaitForSingleObject(handle, 5000) == 0
            finally:
                handle.Close()
    assert owned
    assert path.read_bytes() == before
    assert not list((tmp_path / ".artifacts/excel-comparison").iterdir())
