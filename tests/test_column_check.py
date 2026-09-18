"""Read-only checking policies, full scope, and complete finding lifecycle."""

import hashlib
import re
from dataclasses import replace
from datetime import UTC, date, datetime
from threading import Event

import pytest
from openpyxl import Workbook
from test_column_profile import rewrite_sheet, table_for

from loan_tape.column_check import ColumnRules, check_column, check_value
from loan_tape.column_profile import ColumnCancelled, ColumnExample
from loan_tape.inspection import InspectionError
from loan_tape.session import AnalysisSession

SESSION = AnalysisSession("check-test", datetime(2026, 9, 17, tzinfo=UTC))


def cell(kind, text, row=1000):
    return ColumnExample(row, kind, text, "General")


def test_identifiers_text_numbers_and_missing_values_are_not_coerced():
    ref = SESSION.reference_date
    identifier = ColumnRules("Identifier")
    assert not check_value(cell("Text", "000123"), identifier, ref)
    assert not check_value(cell("Text", "N/A"), identifier, ref)
    assert "non-text" in check_value(cell("Number", "123"), identifier, ref)[0]
    assert "whitespace" in check_value(cell("Text", " 000123 "), identifier, ref)[0]
    for kind, text in [("Blank", ""), ("Empty text", ""), ("Text", " \t")]:
        assert "required" in check_value(cell(kind, text), identifier, ref)[0]
        assert not check_value(cell(kind, text), ColumnRules("Text", required=False), ref)
    assert not check_value(cell("Number", "0"), ColumnRules("Number"), ref)
    assert not check_value(cell("Number", "-123.5"), ColumnRules("Number"), ref)
    for kind, text in [
        ("Text", "0"),
        ("Text", "1,234"),
        ("Boolean", "False"),
        ("Date/time", "2030-01-01"),
    ]:
        assert (
            "Expected a stored number"
            in check_value(cell(kind, text), ColumnRules("Number"), ref)[0]
        )
    assert "not finite" in check_value(cell("Number", "inf"), ColumnRules("Number"), ref)[0]
    assert "Expected text" in check_value(cell("Number", "0"), ColumnRules("Text"), ref)[0]
    for expected in ("Identifier", "Text", "Number", "Date"):
        assert (
            "not evaluated" in check_value(cell("Formula", "=1+1"), ColumnRules(expected), ref)[0]
        )
        assert (
            "Stored Excel error"
            in check_value(cell("Excel error", "#N/A"), ColumnRules(expected), ref)[0]
        )


def test_explicit_date_conventions_calendar_validity_and_reference_boundary():
    reference = date(2026, 2, 1)
    us = ColumnRules("Date", date_format="MM/DD/YYYY", flag_before_reference=True)
    uk = replace(us, date_format="DD/MM/YYYY")
    ambiguous = cell("Text", "01/02/2026")
    assert "2026-01-02" in check_value(ambiguous, us, reference)[0]
    assert not check_value(ambiguous, uk, reference)  # Equal reference day is not past.
    strict = ColumnRules("Date", date_format="YYYY-MM-DD", min_year=2000, max_year=2050)
    assert not check_value(cell("Text", "2024-02-29"), strict, reference)
    for value in ("2023-02-29", "2026-13-01", "2026-2-01", "01/02/2026", "20260101", " 2026-02-01"):
        assert check_value(cell("Text", value), strict, reference)
    assert "before the earliest" in check_value(cell("Text", "1999-12-31"), strict, reference)[0]
    assert (
        "after the latest"
        in check_value(cell("Date/time", "2051-01-01T00:00:00"), strict, reference)[0]
    )
    assert not check_value(cell("Date/time", "2050-12-31T23:59:00"), strict, reference)
    for value in ("12:00:00", "1 day, 0:00:00"):
        assert check_value(cell("Date/time", value), strict, reference)
    assert check_value(cell("Number", "45000"), strict, reference)  # No guessed serial dates.
    assert check_value(cell("Text", "2026-01-01"), ColumnRules("Date"), reference)
    assert not check_value(cell("Date/time", "2026-01-01"), ColumnRules("Date"), reference)


@pytest.mark.parametrize(
    "settings",
    [
        {"expected_type": "Auto"},
        {"expected_type": "Date", "min_year": 0},
        {"expected_type": "Date", "max_year": 10000},
        {"expected_type": "Date", "min_year": 2030, "max_year": 2020},
        {"expected_type": "Date", "min_year": True},
        {"expected_type": "Date", "date_format": "auto"},
        {"expected_type": "Number", "flag_before_reference": True},
    ],
)
def test_invalid_settings_are_rejected(settings):
    with pytest.raises(InspectionError):
        ColumnRules(**settings)


def test_every_finding_is_accessible_across_hidden_distant_rows_and_blank_tail(
    tmp_path, monkeypatch
):
    monkeypatch.chdir(tmp_path)
    path = tmp_path / "distant.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Loans"
    sheet["DZ4000"] = "Preamble"
    sheet["DZ5000"] = "Loan ID"
    for row in range(5001, 5045):
        sheet.cell(row, 130, row)
    sheet["DZ5001"] = "000123"
    sheet["DZ5002"] = "=1+1"
    sheet["DZ5003"] = "#N/A"
    sheet["DZ5004"] = ""
    sheet["DZ5005"] = " \t"
    sheet["DZ5006"] = "N/A"
    sheet.row_dimensions[5044].hidden = True
    sheet.column_dimensions["DZ"].hidden = True
    sheet.sheet_state = "hidden"
    book.create_sheet("Visible")
    sheet["EA5000"] = "Side table"
    sheet["EA5044"] = "Leave this alone"
    book.save(path)
    book.close()
    rewrite_sheet(
        path, lambda xml: re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1:A1"', xml)
    )
    original = path.read_bytes()
    table = table_for(path, "DZ4000:DZ5050", 5000)
    report = check_column(path, table, 130, SESSION, ColumnRules("Identifier"))
    try:
        assert report.profile.total_rows == 50 and report.finding_count == 48
        assert report.profile.first_row == 5001 and report.profile.header == "Loan ID"
        assert report.profile.session is SESSION
        findings = [
            item for offset in range(0, report.finding_count, 20) for item in report.page(offset)
        ]
        assert len(findings) == 48 and len({f.cell.row for f in findings}) == 48
        assert findings[0].cell.row == 5002 and findings[-1].cell.row == 5050
        assert {f.cell.row for f in findings} == set(range(5001, 5051)) - {5001, 5006}
        assert findings[-1].cell.kind == "Blank"
        assert any(f.cell.kind == "Empty text" for f in findings)
        assert any("Whitespace-only" in f.reason for f in findings)
        assert "=1+1" == findings[0].cell.text
        assert not report.page(48)
        with pytest.raises(InspectionError):
            report.page(-1)
    finally:
        report.close()
    assert not list((tmp_path / ".artifacts/column-checks").iterdir())
    with pytest.raises(InspectionError, match="closed"):
        report.page()
    optional = check_column(path, table, 130, SESSION, ColumnRules("Identifier", required=False))
    assert optional.allowed_blanks == 8 and optional.finding_count == 40
    optional.close()
    assert path.read_bytes() == original
    path.rename(tmp_path / "handles-closed.xlsx")


def test_zero_data_rows_and_last_excel_cell_are_accounted_for(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    path = tmp_path / "edge.xlsx"
    book = Workbook()
    book.active.title = "Loans"
    book.active["XFD1048574"] = "Header"
    book.active["XFD1048575"] = "0001"
    book.save(path)
    book.close()
    report = check_column(
        path,
        table_for(path, "XFD1048574:XFD1048576", 1048574),
        16384,
        SESSION,
        ColumnRules("Identifier"),
    )
    assert report.profile.total_rows == 2 and report.finding_count == 1
    assert report.page()[0].cell.row == 1048576
    report.close()
    empty = check_column(
        path,
        table_for(path, "XFD1048574:XFD1048574", 1048574),
        16384,
        SESSION,
        ColumnRules("Identifier"),
    )
    assert empty.profile.total_rows == 0 and not empty.finding_count and not empty.page()
    empty.close()


def test_many_blank_tail_findings_and_cancellation_do_not_truncate_or_leak(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    path = tmp_path / "blanks.xlsx"
    book = Workbook()
    book.active.title = "Loans"
    book.active["A1"] = "Header"
    book.save(path)
    book.close()
    table = table_for(path, "A1:A20002", 1)
    report = check_column(path, table, 1, SESSION, ColumnRules("Text"))
    assert report.finding_count == 20001
    assert len(report.page(19980)) == 20
    assert report.page(20000)[0].cell.row == 20002
    report.close()
    cancelled = Event()

    def progress(message):
        if "Read 10,000" in message:
            cancelled.set()

    with pytest.raises(ColumnCancelled):
        check_column(
            path, table, 1, SESSION, ColumnRules("Text"), cancelled=cancelled, progress=progress
        )
    assert not list((tmp_path / ".artifacts/column-checks").iterdir())
    assert not list((tmp_path / ".artifacts/column-inspection").iterdir())


def test_hash_read_and_callback_failures_discard_incomplete_results(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    path = tmp_path / "failures.xlsx"
    book = Workbook()
    book.active.title = "Loans"
    book.active["A1"] = "Text"
    book.save(path)
    book.close()
    table = table_for(path, "A1:A3")

    def change_source(message):
        if message.startswith("Verifying"):
            with path.open("ab") as handle:
                handle.write(b"changed")

    with pytest.raises(InspectionError, match="changed"):
        check_column(path, table, 1, SESSION, ColumnRules("Number"), progress=change_source)
    assert not list((tmp_path / ".artifacts/column-checks").iterdir())
    table = table_for(path, "A1:A3")

    def fail_value(*args):
        raise RuntimeError("Simulated rule failure")

    with monkeypatch.context() as context:
        context.setattr("loan_tape.column_check.check_value", fail_value)
        with pytest.raises(RuntimeError, match="rule failure"):
            check_column(path, table, 1, SESSION, ColumnRules("Number"))
    assert not list((tmp_path / ".artifacts/column-checks").iterdir())
    rewrite_sheet(
        path,
        lambda xml: xml.replace(b"inlineStr", b"n").replace(
            b"<is><t>Text</t></is>", b"<v>broken</v>"
        ),
    )
    with pytest.raises(ValueError):
        check_column(path, table_for(path, "A1:A3"), 1, SESSION, ColumnRules("Number"))
    assert not list((tmp_path / ".artifacts/column-checks").iterdir())
    assert not list((tmp_path / ".artifacts/column-inspection").iterdir())


def test_date_checks_use_captured_session_day_not_the_wall_clock(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    path = tmp_path / "dates.xlsx"
    book = Workbook()
    book.active.title = "Loans"
    for row in [(datetime(2026, 9, 16),), (datetime(2026, 9, 17),), ("09/18/2026",)]:
        book.active.append(row)
    book.save(path)
    book.close()
    before = hashlib.sha256(path.read_bytes()).hexdigest()
    report = check_column(
        path,
        table_for(path, "A1:A3"),
        1,
        SESSION,
        ColumnRules("Date", date_format="MM/DD/YYYY", flag_before_reference=True),
    )
    assert report.finding_count == 1 and report.page()[0].cell.row == 1
    assert "reference date 2026-09-17" in report.page()[0].reason
    report.close()
    assert hashlib.sha256(path.read_bytes()).hexdigest() == before
