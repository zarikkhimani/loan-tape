"""Full-sheet, editable structure suggestions; synthetic sources only."""

from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.worksheet.table import Table

from loan_tape.ranges import parse_range
from loan_tape.workbook_index import scan_workbook


def block(sheet, top, left, bottom, width):
    for column in range(left, left + width):
        sheet.cell(top, column, f"Header {column}")
        for row in range(top + 1, bottom + 1):
            sheet.cell(row, column, row * 100 + column)


def test_title_rows_and_formatting_do_not_expand_suggested_table(tmp_path):
    path = tmp_path / "titles.xlsx"
    book = Workbook()
    sheet = book.active
    sheet["A1"] = "Portfolio summary"
    sheet.merge_cells("A1:D1")
    sheet["A2"] = "Reporting context"
    block(sheet, 3, 1, 70, 4)
    sheet["B20"] = None  # A missing value inside a table is not an edge.
    sheet["A74"] = "Source note"
    sheet["XFD1048576"].fill = PatternFill("solid", fgColor="FF0000")
    book.save(path)
    before = path.read_bytes()
    index = scan_workbook(path)
    assert index.complete
    sheet = index.sheets[0]
    assert sheet.bounds == parse_range("A1:D74")
    assert [(table.area.address, table.header_row) for table in sheet.tables] == [("A3:D70", 3)]
    assert path.read_bytes() == before


def test_stacked_tables_have_separate_ends_beside_taller_table(tmp_path):
    path = tmp_path / "neighbours.xlsx"
    book = Workbook()
    sheet = book.active
    block(sheet, 10, 2, 20, 3)
    block(sheet, 25, 2, 40, 3)
    block(sheet, 10, 8, 65, 2)
    sheet.sheet_state = "hidden"
    book.create_sheet("Visible")
    book.save(path)
    result = scan_workbook(path)
    assert result.complete
    tables = result.sheets[0].tables
    assert [(t.area.address, t.header_row) for t in tables] == [
        ("B10:D20", 10),
        ("H10:I65", 10),
        ("B25:D40", 25),
    ]


def test_distant_table_and_numeric_first_row_are_not_limited_to_preview(tmp_path):
    path = tmp_path / "distant.xlsx"
    book = Workbook()
    block(book.active, 5000, 130, 5300, 3)
    for column in (150, 151):
        book.active.cell(6000, column, 12)
        book.active.cell(6001, column, 13)
    book.save(path)
    result = scan_workbook(path)
    assert result.complete
    assert [(t.area.address, t.header_row) for t in result.sheets[0].tables] == [
        ("DZ5000:EB5300", 5000),
        ("ET6000:EU6001", None),
    ]


@pytest.mark.parametrize("headers", [0, 1])
def test_declared_excel_table_preserves_blank_columns_and_tail(tmp_path, headers):
    path = tmp_path / "declared.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.append(["ID", "Optional", "Amount"])
    sheet.append(["000123", None, 0])
    sheet["A20"] = "After blank rows"
    sheet.add_table(Table(displayName="Portfolio", ref="A1:C50", headerRowCount=headers))
    book.save(path)
    before = path.read_bytes()
    result = scan_workbook(path)
    assert result.complete
    assert len(result.sheets[0].tables) == 1
    suggested = result.sheets[0].tables[0]
    assert suggested.area.address == "A1:C50"
    assert suggested.header_row == (1 if headers else None)
    assert suggested.name == "Portfolio" and suggested.basis == "excel_table"
    assert path.read_bytes() == before


def test_more_than_200_tables_are_all_grouped_separately(tmp_path):
    path = tmp_path / "many.xlsx"
    book = Workbook()
    for row in range(1, 604, 3):
        block(book.active, row, 1, row + 1, 2)
    book.save(path)
    result = scan_workbook(path)
    assert result.complete
    sheet = result.sheets[0]
    assert len(sheet.tables) == 201
    assert [t.area.min_row for t in sheet.tables] == list(range(1, 604, 3))
    assert all(t.header_row == t.area.min_row for t in sheet.tables)
    assert sheet.stored_cells == 804
    assert sheet.bounds.address == "A1:B602"


@pytest.mark.parametrize("top,left", [(1, 4), (11, 1)])
def test_plain_table_touching_declared_table_is_not_lost(tmp_path, top, left):
    path = tmp_path / "touching.xlsx"
    book = Workbook()
    sheet = book.active
    block(sheet, 1, 1, 10, 3)
    sheet.add_table(Table(displayName="Declared", ref="A1:C10"))
    block(sheet, top, left, top + 5, 3)
    book.save(path)
    before = path.read_bytes()
    result = scan_workbook(path)
    assert result.complete
    tables = result.sheets[0].tables
    assert len(tables) == 2
    assert tables[0].name == "Declared" and tables[0].area.address == "A1:C10"
    assert tables[1].area.min_row == top and tables[1].area.min_column == left
    assert tables[1].area.max_row == top + 5 and tables[1].header_row == top
    assert path.read_bytes() == before


def test_many_side_by_side_tables_and_last_worksheet_cell(tmp_path):
    path = tmp_path / "all-directions.xlsx"
    book = Workbook()
    sheet = book.active
    for column in range(1, 604, 3):
        block(sheet, 1, column, 3, 2)
    block(sheet, 1048575, 16383, 1048576, 2)
    sheet.row_dimensions[1048576].hidden = True
    sheet.column_dimensions["XFD"].hidden = True
    book.save(path)
    result = scan_workbook(path)
    assert result.complete
    indexed = result.sheets[0]
    assert len(indexed.tables) == 202
    assert indexed.tables[-1].area.address == "XFC1048575:XFD1048576"
    assert indexed.tables[-1].header_row == 1048575
    assert indexed.stored_cells == 201 * 6 + 4


@pytest.mark.parametrize("problem", ["missing_relationships", "wrong_type"])
def test_broken_table_metadata_never_silently_loses_groups(tmp_path, problem):
    path = tmp_path / "broken-tables.xlsx"
    book = Workbook()
    block(book.active, 1, 1, 5, 2)
    book.active.add_table(Table(displayName="Broken", ref="A1:B5"))
    book.save(path)
    with ZipFile(path) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    member = "xl/worksheets/_rels/sheet1.xml.rels"
    if problem == "missing_relationships":
        del entries[member]
    else:
        entries[member] = entries[member].replace(b'/table"', b'/unsupported"')
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        for name, content in entries.items():
            archive.writestr(name, content)
    before = path.read_bytes()
    result = scan_workbook(path)
    assert not result.complete and result.sheets[0].error
    assert not result.sheets[0].tables
    assert path.read_bytes() == before
