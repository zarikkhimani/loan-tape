"""Orderly desktop shutdown and cancellation at reader load boundaries."""

import os
import subprocess
import sys
from pathlib import Path
from threading import Event

import pytest
from openpyxl import Workbook, load_workbook

from loan_tape.inspection import CancellableReader, InspectionError, read_preview
from loan_tape.workers import WorkerGroup


def test_registry_keeps_retired_readers_and_allows_copies_to_finish():
    workers = WorkerGroup()
    cancelled, read_started, finish_read, finish_copy = Event(), Event(), Event(), Event()

    def read():
        read_started.set()
        assert cancelled.wait(5)
        assert finish_read.wait(5)

    reader = workers.start(read, "old-preview", cancelled)
    copier = workers.start(lambda: finish_copy.wait(5), "file-copy")
    try:
        assert read_started.wait(5)
        workers.cancel_reads()
        assert cancelled.is_set()
        assert set(workers.active_names()) == {"old-preview", "file-copy"}
        assert not reader.daemon and not copier.daemon
        with pytest.raises(RuntimeError, match="closing"):
            workers.start(lambda: None, "new-work")
    finally:
        finish_read.set()
        finish_copy.set()
        reader.join(5)
        copier.join(5)
    assert not workers.active_names()


def test_cancellation_interrupts_workbook_load_before_first_row(tmp_path, monkeypatch):
    path = tmp_path / "data.xlsx"
    book = Workbook()
    book.active["A1"] = "test"
    book.save(path)
    book.close()
    cancelled = Event()

    def interrupt_load(source, **kwargs):
        cancelled.set()
        return load_workbook(source, **kwargs)

    monkeypatch.setattr("loan_tape.inspection.load_workbook", interrupt_load)
    with pytest.raises(InspectionError, match="cancelled"):
        read_preview(path, cancelled=cancelled)
    path.rename(tmp_path / "handles-closed.xlsx")


def test_buffered_reads_check_cancellation_even_with_buffered_bytes(tmp_path):
    path = tmp_path / "read.csv"
    path.write_bytes(b"test\n" * 1000)
    cancelled = Event()
    with CancellableReader(path, cancelled) as source:
        assert source.read(1) == b"t"
        cancelled.set()
        with pytest.raises(InspectionError, match="cancelled"):
            source.read(1)
        with pytest.raises(InspectionError, match="cancelled"):
            source.read1(1)
        with pytest.raises(InspectionError, match="cancelled"):
            source.readinto(bytearray(1))
    path.unlink()


@pytest.mark.skipif(
    os.name != "nt" and not os.environ.get("DISPLAY"), reason="Requires a desktop display"
)
@pytest.mark.parametrize("scenario", ["discovery", "preview", "column", "check", "retired"])
def test_main_window_waits_for_worker_cleanup(scenario, tmp_path):
    result = subprocess.run(
        [
            sys.executable,
            str(Path(__file__).with_name("shutdown_scenarios.py")),
            scenario,
            str(tmp_path),
        ],
        cwd=tmp_path,
        capture_output=True,
        text=True,
        timeout=20,
    )
    assert result.returncode == 0, result.stdout + result.stderr
