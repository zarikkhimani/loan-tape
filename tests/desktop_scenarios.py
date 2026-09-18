"""Native integration cases, isolated to one Tcl/Tk interpreter per process."""

import io
import json
import shutil
import sys
import time
from dataclasses import replace
from datetime import datetime
from pathlib import Path
from threading import Event
from unittest.mock import patch

from openpyxl import Workbook
from test_inspection import write_workbook
from test_table_structure import block

from loan_tape.column_check import check_column
from loan_tape.column_profile import ColumnCancelled
from loan_tape.excel_compare import compare_previews, read_parallel_preview
from loan_tape.inspection import InspectionError, PreviewCell, read_preview
from loan_tape.intake import IntakeStore
from loan_tape.mapping_profiles import load_builtin_profiles
from loan_tape.selection import load_selection, load_tables
from loan_tape.ui import LoanTapeApp, create_root


def wait_for_copy(app: LoanTapeApp) -> None:
    deadline = time.monotonic() + 10
    while app.worker is not None and not app.closed:
        app.root.update()
        assert time.monotonic() <= deadline, "Desktop copy did not finish."
        time.sleep(0.01)
    if not app.closed:
        app.root.update()


def wait_for_preview(app: LoanTapeApp) -> None:
    assert app.inspector is not None
    app.root.withdraw()
    deadline = time.monotonic() + 10
    while app.inspector.busy and not app.inspector.closed:
        app.root.update()
        assert time.monotonic() <= deadline, "Preview did not finish."
        time.sleep(0.01)
    if not app.closed:
        app.root.update()


def wait_for_column(app: LoanTapeApp) -> None:
    window = app.inspector.column_inspector
    assert window is not None
    window.window.withdraw()
    deadline = time.monotonic() + 10
    while window.busy and not window.closed:
        app.root.update()
        assert time.monotonic() <= deadline, "Column inspection did not finish."
        time.sleep(0.01)
    app.root.update()


def dispatch_drop(app: LoanTapeApp, paths: list[Path]) -> str:
    """Use the DnD extension's registered callback, including its Tcl event conversion."""
    script = str(app.root.tk.call("bind", str(app.drop_zone), "<<Drop>>"))
    callback = script.split()[0]
    data = app.root.tk.call("format", "%s", tuple(str(path) for path in paths))
    return str(
        app.root.tk.call(
            callback,
            "copy",
            "copy",
            1,
            "",
            "",
            "DND_Files",
            "DND_Files",
            data,
            "<<Drop>>",
            "DND_Files",
            "",
            "DND_Files",
            "DND_Files",
            "DND_Files",
            "DND_Files",
            str(app.drop_zone),
            0,
            0,
        )
    )


def run(scenario: str, work: Path) -> None:
    store = IntakeStore(work / "inputs")
    if scenario == "inspect_date_setup":
        shutil.copytree(
            Path(__file__).resolve().parents[1] / "packs" / "loan-dates",
            work / "packs" / "loan-dates",
        )
    if scenario == "legacy":
        store.save("legacy.csv", io.BytesIO(b"legacy"), 6)
    root = create_root()
    root.withdraw()
    app = LoanTapeApp(root, store)
    callback_errors = []
    root.report_callback_exception = lambda *args: callback_errors.append(args)
    source = work / "Résumé {Q1} 100%.CSV"
    payload = b"loan_id,balance\r\n000123,0\r\n"
    source.write_bytes(payload)
    try:
        if scenario in {"drop", "create"}:
            assert dispatch_drop(app, [source]) == "copy"
            wait_for_copy(app)
            assert app.records, app.status.get()
            assert source.read_bytes() == payload
            assert Path(app.saved_path.get()).read_bytes() == payload
            assert app.original_path.get() == str(source.absolute())
            assert len(app.history.get_children()) == 1
            app.original_button.invoke()
            assert root.clipboard_get() == str(source.absolute())
            app.saved_button.invoke()
            assert root.clipboard_get() == app.saved_path.get()
        elif scenario == "session":
            session = app.session
            records = Path.cwd() / ".artifacts" / "sessions"
            assert len(list(records.glob("*.json"))) == 1
            record_path = records / (session.id + ".json")
            record_bytes = record_path.read_bytes()
            record = json.loads(record_bytes)
            assert record["reference_date"] == session.reference_date.isoformat()
            assert app.session_label.cget("text") == session.startup_label
            root.geometry("760x680")
            root.deiconify()
            root.update()
            assert (
                app.session_label.winfo_rootx() + app.session_label.winfo_width()
                <= root.winfo_rootx() + root.winfo_width()
            )
            root.withdraw()
            app.add_file(source)
            wait_for_copy(app)
            for _ in range(2):
                with patch("loan_tape.session.datetime") as clock:
                    app.inspect_button.invoke()
                    wait_for_preview(app)
                    assert app.inspector.session is session
                    assert app.inspector.session_label.cget("text") == session.date_label
                    app.inspector.close()
                    clock.now.assert_not_called()
            assert len(list(records.glob("*.json"))) == 1
            assert record_path.read_bytes() == record_bytes
            app.close()
            assert not record_path.exists()
            assert store.list_files() == []
            root = create_root()
            root.withdraw()
            app = LoanTapeApp(root, store)
            assert app.session.id != session.id
            assert len(list(records.glob("*.json"))) == 1
        elif scenario == "reopen":
            assert app.records == {}
            assert not store.directory.exists()
            assert app.saved_heading.cget("text") == "No session files"
        elif scenario == "browse":
            workbook = work / "synthetic.xlsx"
            workbook.write_bytes(b"synthetic opaque workbook bytes")
            with patch("loan_tape.ui.filedialog.askopenfilename", return_value=str(workbook)):
                app.browse_button.invoke()
            wait_for_copy(app)
            assert Path(app.saved_path.get()).read_bytes() == workbook.read_bytes()
            with patch("loan_tape.ui.filedialog.askopenfilename", return_value=""):
                app.browse_button.invoke()
            assert len(app.history.get_children()) == 1
        elif scenario == "multiple":
            assert dispatch_drop(app, [source, work / "other.csv"]) == "refuse_drop"
            assert "one file at a time" in app.status.get()
            assert not store.directory.exists()
        elif scenario in {"invalid", "empty"}:
            bad_file = work / ("wrong.pdf" if scenario == "invalid" else "empty.csv")
            bad_file.write_bytes(b"wrong" if scenario == "invalid" else b"")
            dispatch_drop(app, [bad_file])
            wait_for_copy(app)
            assert not app.records
            assert not app.browse_button.instate(["disabled"])
            assert str(app.status_label.cget("foreground")) == "#a1342d"
            dispatch_drop(app, [source])
            wait_for_copy(app)
            assert len(app.records) == 1
        elif scenario == "close_busy":
            release = Event()
            original_save = store.save_path

            def paused_copy(path: Path):
                assert release.wait(5), "Test did not release copy worker"
                return original_save(path)

            with patch.object(store, "save_path", side_effect=paused_copy):
                try:
                    assert app.add_file(source)
                    assert dispatch_drop(app, [source]) == "refuse_drop"
                    app.close()
                    assert app.closing and not app.closed
                finally:
                    release.set()
                wait_for_copy(app)
            assert app.closed
            assert store.list_files() == []
            assert source.read_bytes() == payload
        elif scenario.startswith("inspect_"):
            assert app.inspect_button.instate(["disabled"])
            if scenario in {"inspect_excel", "inspect_comparison"}:
                source = work / "preview.xlsx"
                write_workbook(source)
            elif scenario == "inspect_mapping_profile":
                source = work / "warehouse.xlsx"
                book = Workbook()
                sheet = book.active
                sheet.title = "Inputs and Portfolio"
                for column, item in enumerate(load_builtin_profiles()[0].columns, 2):
                    sheet.cell(5, column, item.source_header)
                    sheet.cell(6, column, 0)
                sheet["A20"] = "=1+1"  # Keep Excel out of this UI-only test.
                book.save(source)
                book.close()
            elif scenario == "inspect_date_setup":
                source = work / "dates.xlsx"
                book = Workbook()
                sheet = book.active
                sheet.title = "Loans"
                sheet.append(("Closing Date", "Current Maturity Date"))
                sheet.append(("2026-01-15", "2031-01-15"))
                sheet.append(("2026-02-28", "2032-02-28"))
                book.save(source)
                book.close()
            elif scenario == "inspect_navigation":
                source = work / "anywhere.xlsx"
                book = Workbook()
                book.active.title = "Empty"
                sheet = book.create_sheet("Loans")
                for number in range(5000, 5501):
                    sheet.cell(number, 130, f"ID-{number:06}")
                sheet["IZ5500"] = "Far right end"
                hidden = book.create_sheet("Hidden")
                hidden.sheet_state = "hidden"
                hidden["Z9999"] = "=1+1"  # Never start Excel during this UI-only test.
                book.save(source)
                book.close()
            elif scenario == "inspect_suggestions":
                source = work / "structure.xlsx"
                book = Workbook()
                sheet = book.active
                sheet["B4998"] = "Portfolio report"
                block(sheet, 5000, 2, 5060, 12)
                block(sheet, 5000, 20, 5010, 3)
                sheet["A6000"] = "=1+1"  # Keep Excel out of this UI-only test.
                book.save(source)
                book.close()
            elif scenario == "inspect_tables":
                source = work / "multiple-tables.xlsx"
                book = Workbook()
                sheet = book.active
                sheet.title = "Loans"
                for row in range(5, 51):
                    sheet.cell(row, 2, f"ID-{row:05}")
                for column, value in enumerate(["Loan ID", "Balance", "Balance", None], 2):
                    sheet.cell(5, column, value)
                sheet["K5"] = "Currency"
                sheet["L5"] = "Off page header"
                sheet["M5"] = "Last header"
                sheet["B90"] = "Second title"
                sheet["B92"] = "Issuer"
                sheet["C92"] = "Amount"
                sheet["B93"] = "Borrower"
                sheet["C93"] = "Outstanding"
                sheet["B110"] = "Second end"
                sheet["Z5"] = "Side ID"
                sheet["AA5"] = "Side amount"
                sheet["Z6"] = "000123"
                sheet["AA10"] = "Side end"
                sheet["A200"] = "=1+1"  # Do not launch Excel for this UI scenario.
                book.save(source)
                book.close()
            elif scenario.startswith("inspect_check_") or scenario in {
                "inspect_column",
                "inspect_column_cancel",
                "inspect_column_error",
            }:
                source = work / "column.xlsx"
                book = Workbook()
                sheet = book.active
                sheet.title = "Loans"
                sheet["DZ5000"] = "Loan ID"
                for row in range(5001, 5501):
                    sheet.cell(row, 130, "COMMON")
                sheet["DZ5005"] = None
                sheet["DZ5025"] = "000123"
                sheet["DZ5200"] = 0
                sheet["DZ5400"] = datetime(2035, 1, 1)
                sheet["DZ5500"] = 0
                sheet["DZ5600"] = "Other ID"
                sheet["DZ5601"] = "ONLYSECOND"
                sheet["EA5000"] = "Side ID"
                sheet["EA5001"] = "001"
                sheet["EA5002"] = "002"
                hidden = book.create_sheet("Hidden")
                hidden.sheet_state = "hidden"
                hidden["A1"] = "=1+1"
                book.save(source)
                book.close()
            elif scenario == "inspect_scan_close":
                source = work / "scanning.xlsx"
                write_workbook(source)
            elif scenario == "inspect_error":
                source = work / "damaged.xlsx"
                source.write_bytes(b"not a workbook")
            app.add_file(source)
            wait_for_copy(app)
            assert not app.inspect_button.instate(["disabled"])
            # Inspection must use the preserved copy, even if the original is gone.
            source.unlink()
            if scenario == "inspect_scan_close":
                started = Event()
                stopped = Event()

                def waiting_scan(path, *, cancelled, progress):
                    from loan_tape.workbook_index import ScanCancelled

                    started.set()
                    assert cancelled.wait(5)
                    stopped.set()
                    raise ScanCancelled("Scan cancelled; incomplete.")

                with patch("loan_tape.preview_ui.scan_workbook", side_effect=waiting_scan):
                    app.inspect_button.invoke()
                    assert started.wait(5)
                    app.inspector.close()
                    assert stopped.wait(5)
                assert app.inspector.closed
            elif scenario == "inspect_mapping_profile":
                preserved = Path(app.saved_path.get())
                original = preserved.read_bytes()
                app.inspect_button.invoke()
                wait_for_preview(app)
                inspector = app.inspector
                nav = inspector.navigator
                assert nav is not None and nav.profile_match is None
                assert nav.area is not None and nav.area.address == "B5:CQ6"
                assert nav.header_row == 5
                nav.table_name.set("Portfolio")
                nav.use_button.invoke()
                wait_for_preview(app)
                assert nav.profile_match is not None
                assert nav.profile_match.profile.id == "warehouse-model"
                assert len(nav.profile_match.bindings) == 94
                assert nav.profile_text.get() == (
                    "Mapping profile: Warehouse Model 1.0.0 · 94 columns matched exactly."
                )
                assert nav.profile_label.winfo_ismapped()
                assert preserved.read_bytes() == original
            elif scenario == "inspect_date_setup":
                app.inspect_button.invoke()
                wait_for_preview(app)
                inspector = app.inspector
                nav = inspector.navigator
                assert nav is not None and inspector.date_button is not None
                nav.header_text.set("1")
                nav.use_button.invoke()
                wait_for_preview(app)
                assert not inspector.date_button.instate(["disabled"])
                assert not app.pack_store.pins()
                inspector.date_button.invoke()
                workflow = inspector.date_workflow
                assert workflow is not None and not workflow.ready
                assert workflow.enable_dates_button is not None
                assert not workflow.enable_dates_button.instate(["disabled"])
                assert "does not change the workbook" in " ".join(
                    child.cget("text")
                    for child in workflow.window.winfo_children()[0].winfo_children()
                    if child.winfo_class() == "TLabel" and "text" in child.keys()
                )
                window_name = str(workflow.window)
                with patch.object(
                    app.pack_store, "activate", side_effect=OSError("synthetic write failure")
                ):
                    workflow.enable_dates_button.invoke()
                    deadline = time.monotonic() + 10
                    while workflow.busy and not workflow.closed:
                        root.update()
                        assert time.monotonic() <= deadline, "Date setup failure was not reported."
                        time.sleep(0.01)
                    root.update()
                assert not workflow.ready
                assert not workflow.enable_dates_button.instate(["disabled"])
                assert "Could not enable loan date definitions" in workflow.status.get()
                workflow.enable_dates_button.invoke()
                deadline = time.monotonic() + 10
                while workflow.busy and not workflow.closed:
                    root.update()
                    assert time.monotonic() <= deadline, "Date definitions did not activate."
                    time.sleep(0.01)
                root.update()
                assert str(workflow.window) == window_name and workflow.ready
                assert workflow.enable_dates_button is None
                assert workflow.tabs.winfo_exists()
                assert "Enabled loan date definitions" in workflow.status.get()
                pins = app.pack_store.pins()
                assert len(pins) == 1 and pins[0].id == "loan.dates"
            elif scenario == "inspect_navigation":
                app.inspect_button.invoke()
                wait_for_preview(app)
                inspector = app.inspector
                navigator = inspector.navigator
                assert navigator.index.complete
                app.root.geometry("980x740")
                app.root.deiconify()
                root.update()
                assert navigator.use_button.winfo_ismapped()
                right_edge = app.root.winfo_rootx() + app.root.winfo_width()
                assert (
                    inspector.session_label.winfo_rootx() + inspector.session_label.winfo_width()
                    <= right_edge
                )
                assert (
                    navigator.use_button.winfo_rootx() + navigator.use_button.winfo_width()
                    <= right_edge
                )
                assert inspector.table.winfo_height() >= 420, (
                    app.root.winfo_geometry(),
                    inspector.table.winfo_height(),
                    navigator.frame.winfo_height(),
                    [
                        (str(w), w.winfo_height(), w.winfo_reqheight())
                        for w in inspector.frame.winfo_children()[0].winfo_children()
                    ],
                )
                app.root.withdraw()
                assert not navigator.overview_window.winfo_viewable()
                app.root.deiconify()
                root.update()
                navigator.find_button.invoke()
                root.update()
                assert navigator.overview_window.winfo_viewable()
                assert str(navigator.details_tabs.select()) == str(navigator.areas)
                inspector.details_button.invoke()
                assert str(navigator.details_tabs.select()) == str(navigator.details)
                navigator.overview_window.withdraw()
                app.root.withdraw()
                assert len(navigator.tree.get_children()) == 3
                assert inspector.sheet_name.get() == "Loans"
                assert navigator.area.address == "DZ5000:IZ5500"
                assert inspector.table.heading("DZ", "text") == "DZ"
                assert inspector.table.item("5000", "values")[0] == "ID-005000"
                inspector.table.selection_set("5000")
                inspector._selection()
                assert "DZ5000" in inspector.cell_text.get("1.0", "end")
                navigator.range_text.set("DZ5000:IZ5500")
                navigator.header_text.set("5000")
                navigator.use_button.invoke()
                wait_for_preview(app)
                full = navigator.saved
                assert full.area.row_count == 501 and full.header_row == 5000
                navigator.down.invoke()
                wait_for_preview(app)
                assert inspector.preview.start_row == 5020
                assert navigator.saved == full
                navigator.right.invoke()
                wait_for_preview(app)
                assert inspector.preview.start_column == 140
                navigator.jump_text.set("IZ5500")
                navigator.jump()
                wait_for_preview(app)
                assert inspector.table.item("5500", "values") == ("Far right end",)
                assert "IZ5500" in inspector.status.get()
                assert load_selection(navigator.index.sha256) == full
                navigator.range_text.set("A0")
                navigator.apply_range()
                assert "range" in navigator.message.get().lower()
                assert navigator.saved == full
                navigator.tree.selection_set("sheet-2")
                root.update()
                assert inspector.sheet_name.get() == "Loans"
                navigator.open_area_button.invoke()
                wait_for_preview(app)
                assert inspector.sheet_name.get() == "Hidden"
                assert inspector.table.item("9999", "values")[0] == "=1+1"
                assert "hidden" in navigator.tree.item("sheet-2", "values")
                navigator.tree.selection_set("sheet-1")
                navigator.all_data_button.invoke()
                wait_for_preview(app)
                assert inspector.sheet_name.get() == "Loans"
                assert navigator.area == navigator.current.bounds
                inspector.close()
                app.inspect_button.invoke()
                wait_for_preview(app)
                assert app.inspector.navigator.saved == full
                assert app.inspector.navigator.area == full.area
                assert app.inspector.preview.start_row == 5000
                assert app.inspector.navigator.header_text.get() == "5000"
            elif scenario == "inspect_suggestions":
                app.inspect_button.invoke()
                wait_for_preview(app)
                inspector, nav = app.inspector, app.inspector.navigator
                original = Path(app.saved_path.get()).read_bytes()
                assert nav.area.address == "B5000:M5060"
                assert nav.header_row == 5000
                assert nav.start_text.get() == "B5000" and nav.end_text.get() == "M5060"
                assert not nav.tables  # Suggestions do not silently create saved definitions.
                assert nav.table_count.get() == "2 data sets found across this sheet."
                assert len(nav.table_list.get_children()) == 3  # Two groups + whole sheet.
                assert nav.table_list.item("0", "values") == ("B5000:M5060",)
                assert nav.table_list.item("1", "values") == ("T5000:V5010",)
                nav.table_list.selection_set("1")
                root.update()
                wait_for_preview(app)
                assert nav.area.address == "T5000:V5010" and not nav.tables
                nav.table_list.selection_set("0")
                root.update()
                wait_for_preview(app)
                assert nav.area.address == "B5000:M5060"
                assert inspector.zoom.get() == "80%"
                app.root.geometry("980x740")
                app.root.deiconify()
                root.update()
                assert inspector.sidebar.winfo_width() <= 240
                assert nav.table_list.heading("#0", "text") == "Data set"
                assert nav.use_button.cget("text") == "Save data set"
                area, page, preview = nav.area, nav.page_area, inspector.preview
                opened_width = inspector.table.winfo_width()
                assert inspector.sidebar_toggle.master is inspector.rail_edge
                assert (
                    inspector.sidebar_toggle.winfo_rootx() + inspector.sidebar_toggle.winfo_width()
                    <= inspector.workspace.winfo_rootx()
                )
                inspector.sidebar_toggle.invoke()
                root.update()
                assert inspector.sidebar_collapsed and not inspector.rail.winfo_ismapped()
                assert (
                    inspector.sidebar_toggle.winfo_ismapped()
                    and inspector.rail_edge.winfo_ismapped()
                )
                assert inspector.table.winfo_width() > opened_width + 200
                assert (nav.area, nav.page_area, inspector.preview) == (area, page, preview)
                inspector.sidebar_toggle.invoke()
                root.update()
                assert not inspector.sidebar_collapsed and inspector.rail.winfo_ismapped()
                assert inspector.table.winfo_width() == opened_width
                inspector.zoom.set("125%")
                inspector._apply_zoom()
                root.update()
                assert inspector.table.yview()[1] < 1
                assert inspector.vertical_scroll.winfo_width() <= 14
                assert inspector.horizontal_scroll.winfo_height() <= 14
                inspector.table.event_generate("<MouseWheel>", delta=-120)
                root.update()
                assert inspector.table.yview()[0] > 0
                assert inspector.row_numbers.yview() == inspector.table.yview()
                vertical_position = inspector.table.yview()
                inspector.table.event_generate("<Shift-MouseWheel>", delta=-120)
                root.update()
                assert inspector.table.xview()[0] > 0
                assert inspector.table.yview() == vertical_position
                # High-resolution wheel deltas accumulate instead of jumping a row per event.
                inspector._scroll_rows("moveto", "0")
                for _ in range(3):
                    inspector.table.event_generate("<MouseWheel>", delta=-10)
                root.update()
                assert inspector.table.yview()[0] == 0
                inspector.table.event_generate("<MouseWheel>", delta=-10)
                root.update()
                assert inspector.table.yview()[0] > 0
                # Clicking the far track jumps across this page, not across the data set.
                inspector.table.xview_moveto(0)
                root.update()
                bar = inspector.horizontal_scroll
                x, y = bar.winfo_width() - 18, bar.winfo_height() // 2
                assert "trough" in bar.identify(x, y)
                bar.event_generate("<Button-1>", x=x, y=y)
                bar.event_generate("<ButtonRelease-1>", x=x, y=y)
                root.update()
                assert inspector.table.xview()[0] > 0.3
                assert inspector.table.xview()[1] == 1
                assert (nav.area, nav.page_area, inspector.preview) == (area, page, preview)
                inspector._scroll_rows("moveto", "0")
                inspector.table.xview_moveto(0)
                root.update()
                # Use the native heading gesture, then inspect the rendered column layer.
                box = inspector.table.bbox("5000", "C")
                x, y = box[0] + 10, box[1] // 2
                inspector.table.event_generate("<ButtonPress-1>", x=x, y=y)
                inspector.table.event_generate("<ButtonRelease-1>", x=x, y=y)
                root.update()
                overlay = inspector.column_highlight.canvas
                assert inspector.column_highlight.column == "C"
                assert inspector.whole_column == (inspector.sheet_name.get(), nav.area, 3)
                assert not inspector.table.selection() and not inspector.row_numbers.selection()
                assert overlay.winfo_ismapped()
                expected_visible = sum(
                    bool(box) and box[1] + box[3] <= overlay.winfo_height()
                    for item in inspector.table.get_children()
                    for box in [inspector.table.bbox(item, "C")]
                )
                assert len(overlay.find_withtag("cell")) == expected_visible
                assert overlay.winfo_x() == inspector.table.bbox("5000", "C")[0]
                assert "Entire data set column: C5000:C5060" in inspector.cell_summary.get()
                # Wheel events over highlighted cells still reach the grid and pinned rows.
                overlay.event_generate("<MouseWheel>", delta=-120, x=10, y=50)
                root.update()
                assert inspector.table.yview()[0] > 0
                assert inspector.row_numbers.yview() == inspector.table.yview()
                overlay.event_generate("<Shift-MouseWheel>", delta=-120, x=10, y=50)
                root.update()
                assert inspector.table.xview()[0] > 0
                assert overlay.winfo_x() == max(1, inspector.table.bbox("5003", "C")[0])
                inspector.sidebar_toggle.invoke()
                root.update()
                assert overlay.winfo_ismapped() and inspector.column_highlight.column == "C"
                inspector.sidebar_toggle.invoke()
                root.update()
                inspector._scroll_rows("moveto", "0")
                inspector.table.xview_moveto(0)
                root.update()
                # A click through the highlight becomes an ordinary cell selection.
                cell_box = inspector.table.bbox("5001", "C")
                overlay.event_generate("<ButtonPress-1>", x=10, y=cell_box[1] + 4)
                overlay.event_generate("<ButtonRelease-1>", x=10, y=cell_box[1] + 4)
                root.update()
                assert inspector.table.selection() == ("5001",)
                assert inspector.whole_column is None and not overlay.winfo_ismapped()
                # A selected source column survives row paging, including off-page headers.
                inspector.table.tk.call(inspector.table.heading("C", "command"))
                nav.down.invoke()
                wait_for_preview(app)
                root.update()
                assert inspector.preview.start_row == 5020
                assert inspector.column_highlight.column == "C" and inspector.active_column == 1
                nav.up.invoke()
                wait_for_preview(app)
                inspector.zoom.set("80%")
                inspector._apply_zoom()
                small = inspector.table.column("B", "width")
                inspector.zoom.set("100%")
                inspector.zoom_selector.event_generate("<<ComboboxSelected>>")
                root.update()
                assert inspector.table.column("B", "width") > small
                assert (
                    inspector.row_numbers.column("#0", "width")
                    >= inspector.grid_font.measure("1048576") + 12
                )
                assert inspector.preview.column_count == 10 and len(inspector.preview.rows) == 20
                for percent in (65, 50, 35, 25, 20, 15):
                    previous_width = inspector.table.column("B", "width")
                    inspector.zoom.set(f"{percent}%")
                    inspector.zoom_selector.event_generate("<<ComboboxSelected>>")
                    root.update()
                    assert inspector.table.column("B", "width") < previous_width
                    assert inspector.preview.column_count == 10
                    assert len(inspector.preview.rows) == 20
                    assert nav.area.address == "B5000:M5060"
                    assert inspector.column_highlight.column == "C"
                    assert inspector.column_highlight.canvas.winfo_width() == inspector.column_width
                assert inspector.table.column("B", "width") == 22
                assert inspector.grid_font.cget("size") == 2
                nav.end_button.invoke()
                wait_for_preview(app)
                assert inspector.table.column("D", "width") == 22  # Paging retains 15%.
                nav.start_button.invoke()
                wait_for_preview(app)
                inspector.zoom.set("80%")
                inspector._apply_zoom()
                assert (
                    "Start B5000" in inspector.status.get()
                    and "End M5060" in inspector.status.get()
                )
                nav.end_button.invoke()
                wait_for_preview(app)
                assert inspector.preview.start_row == 5041 and inspector.preview.start_column == 4
                assert inspector.table.exists("5060")
                assert (
                    "Bottom reached" in inspector.status.get()
                    and "Right edge reached" in inspector.status.get()
                )
                assert nav.area.address == "B5000:M5060"
                nav.start_button.invoke()
                wait_for_preview(app)
                assert inspector.preview.start_row == 5000 and inspector.preview.start_column == 2
                nav.end_text.set("M5050")
                nav.preview_button.invoke()
                wait_for_preview(app)
                assert nav.area.address == "B5000:M5050"
                nav.use_button.invoke()
                wait_for_preview(app)
                first = nav.tables[0]
                nav.new_button.invoke()
                wait_for_preview(app)
                assert (
                    nav.area.address == "T5000:V5010"
                )  # Skip the already-saved overlapping suggestion.
                assert nav.header_row == 5000
                inspector.table.selection_set("5009")
                inspector._select_column(2)
                inspector._set_corner(True)
                wait_for_preview(app)
                assert nav.end_text.get() == "V5009" and nav.area.address == "T5000:V5009"
                nav.use_button.invoke()
                wait_for_preview(app)
                assert len(nav.tables) == 2 and nav.tables[0] == first
                inspector.close()
                app.inspect_button.invoke()
                wait_for_preview(app)
                nav = app.inspector.navigator
                assert (
                    nav.area.address == "T5000:V5009"
                )  # Restore edited boundaries, not fresh guesses.
                assert len(nav.tables) == 2
                assert nav.table_list.item("0", "values") == ("B5000:M5050",)
                nav.table_list.selection_set("0")
                root.update()
                wait_for_preview(app)
                assert nav.area.address == "B5000:M5050"
                assert nav.editing_id == first.id and len(nav.tables) == 2
                nav.table_list.selection_set("1")
                root.update()
                wait_for_preview(app)
                assert nav.area.address == "T5000:V5009" and len(nav.tables) == 2
                nav.table_list.selection_set(str(len(nav.current.tables)))
                root.update()
                wait_for_preview(app)
                assert nav.area == nav.current.bounds and nav.header_row is None
                assert Path(app.saved_path.get()).read_bytes() == original
            elif scenario == "inspect_tables":
                app.inspect_button.invoke()
                wait_for_preview(app)
                inspector = app.inspector
                nav = inspector.navigator
                assert inspector.date_button is not None
                assert inspector.date_button.instate(["disabled"])
                saved_copy = Path(app.saved_path.get())
                original_bytes = saved_copy.read_bytes()
                assert nav.editor.winfo_manager() == "grid"
                assert inspector.column_actions.winfo_manager() == "grid"
                assert inspector.column_button.instate(["disabled"])
                assert inspector.workflow_hint.get() == "Next: save the data set."
                assert not inspector.detail_frame.winfo_manager()
                inspector.detail_button.invoke()
                assert inspector.detail_frame.winfo_manager() == "grid"
                inspector.detail_button.invoke()
                assert not inspector.detail_frame.winfo_manager()
                nav.range_text.set("B5:M50")
                nav.apply_range()
                wait_for_preview(app)
                inspector.table.selection_set("5")
                inspector._selection()
                inspector.header_button.invoke()
                wait_for_preview(app)
                assert inspector.table.heading("B", "text") == "B · Loan ID"
                assert inspector.table.heading("C", "text") == "C · Balance"
                assert inspector.table.heading("D", "text") == "D · Balance"
                assert inspector.table.heading("E", "text") == "E · [blank header]"
                assert "header" in inspector.table.item("5", "tags")
                nav.table_name.set("Top loans")
                nav.use_button.invoke()
                wait_for_preview(app)
                top = nav.tables[0]
                assert not nav.editor.winfo_manager()
                assert nav.summary.winfo_manager() == "grid"
                assert inspector.column_actions.winfo_manager() == "grid"
                assert inspector.column_button.instate(["disabled"])
                assert inspector.workflow_hint.get() == "Next: select a column heading."
                inspector._select_column(0)
                assert not inspector.column_button.instate(["disabled"])
                assert not inspector.date_button.instate(["disabled"])
                assert inspector.workflow_hint.get() == "Ready: choose a review action."
                nav.edit_button.invoke()
                assert nav.editor.winfo_manager() == "grid"
                nav.table_name.set("Unsaved name")
                assert inspector.column_button.instate(["disabled"])
                assert inspector.date_button.instate(["disabled"])
                assert inspector.column_actions.winfo_manager() == "grid"
                assert inspector.workflow_hint.get() == "Next: save the data set."
                nav.table_name.set(top.name)
                nav.use_button.invoke()
                wait_for_preview(app)
                assert not nav.editor.winfo_manager()
                assert not inspector.date_button.instate(["disabled"])
                assert top.selection.area.address == "B5:M50"
                nav.down.invoke()
                wait_for_preview(app)
                assert "5" not in inspector.table.get_children()
                assert inspector.table.heading("B", "text") == "B · Loan ID"
                inspector.table.selection_set("25")
                inspector._selection()
                assert "Column header (B5): Loan ID" in inspector.cell_text.get("1.0", "end")
                nav.right.invoke()
                wait_for_preview(app)
                assert inspector.table.heading("L", "text") == "L · Off page header"
                assert inspector.table.heading("M", "text") == "M · Last header"
                for name, area, header in [
                    ("Below loans", "B90:C110", "92"),
                    ("Side loans", "Z5:AA10", "5"),
                ]:
                    nav.new_button.invoke()
                    wait_for_preview(app)
                    assert nav.editor.winfo_manager() == "grid"
                    assert nav.area is not None
                    nav.table_name.set(name)
                    nav.range_text.set(area)
                    nav.header_text.set(header)
                    nav.use_button.invoke()
                    wait_for_preview(app)
                assert len(nav.tables) == 3
                assert nav.tables[0] == top
                assert inspector.table.heading("Z", "text") == "Z · Side ID"
                assert inspector.table.item("6", "values")[0] == "000123"
                assert load_tables(nav.index.sha256).tables == nav.tables
                nav.table_selector.current(1)
                nav.choose_table()
                wait_for_preview(app)
                assert inspector.table.heading("B", "text") == "B · Issuer"
                inspector.table.selection_set("93")
                inspector._selection()
                inspector.header_button.invoke()
                wait_for_preview(app)
                nav.use_button.invoke()
                wait_for_preview(app)
                assert inspector.table.heading("B", "text") == "B · Borrower"
                assert len(nav.tables) == 3 and nav.tables[0] == top
                assert nav.tables[1].selection.header_row == 93
                before_invalid = load_tables(nav.index.sha256)
                nav.header_text.set("5")
                nav.use_button.invoke()
                assert "inside the selected range" in nav.message.get()
                assert load_tables(nav.index.sha256) == before_invalid
                nav.header_text.set("93")
                inspector.close()
                app.inspect_button.invoke()
                wait_for_preview(app)
                inspector = app.inspector
                nav = inspector.navigator
                assert len(nav.tables) == 3
                assert nav.table_name.get() == "Below loans"
                assert inspector.table.heading("B", "text") == "B · Borrower"
                nav.table_selector.current(0)
                nav.choose_table()
                wait_for_preview(app)
                assert inspector.table.heading("B", "text") == "B · Loan ID"
                nav.remove_button.invoke()
                wait_for_preview(app)
                assert len(nav.tables) == 2
                assert [t.name for t in load_tables(nav.index.sha256).tables] == [
                    "Below loans",
                    "Side loans",
                ]
                assert saved_copy.read_bytes() == original_bytes
            elif scenario.startswith("inspect_check_"):
                app.inspect_button.invoke()
                wait_for_preview(app)
                inspector = app.inspector
                nav = inspector.navigator
                saved_copy = Path(app.saved_path.get())
                source_bytes = saved_copy.read_bytes()
                assert inspector.check_button.instate(["disabled"])
                inspector._select_column(0)
                assert inspector.check_button.instate(["disabled"])
                inspector.check_column()
                assert inspector.column_checker is None
                assert "Save this data set" in nav.message.get()
                nav.range_text.set("DZ5000:DZ5500")
                nav.header_text.set("5000")
                nav.table_name.set("Main table")
                nav.use_button.invoke()
                wait_for_preview(app)
                full = nav.table_for_inspection()
                inspector.table.tk.call(inspector.table.heading("DZ", "command"))
                assert inspector.check_button.cget("text") == "Check column DZ"
                inspector.check_button.invoke()
                child = inspector.column_checker
                assert child and not child.busy
                child.window.withdraw()
                child.run_button.invoke()
                assert child.report is None and "Choose Identifier" in child.status.get()
                child.expected.set("Number")

                def wait_check():
                    deadline = time.monotonic() + 10
                    while child.busy and not child.closed:
                        root.update()
                        assert time.monotonic() < deadline, "Column check timed out"
                        time.sleep(0.01)
                    root.update()

                def clean_results():
                    assert not list((work / ".artifacts/column-checks").glob("check-*"))

                if scenario == "inspect_check_cancel":
                    started = Event()

                    def paused_check(*args, cancelled, **kwargs):
                        started.set()
                        assert cancelled.wait(5)
                        raise ColumnCancelled("Cancelled")

                    with patch("loan_tape.column_check_ui.check_column", side_effect=paused_check):
                        child.run_button.invoke()
                        assert started.wait(5)
                        inspector.close()
                        child.worker.join(5)
                        assert child.closed and child.cancelled.is_set()
                        assert not child.worker.is_alive() and child.report is None
                    clean_results()
                elif scenario == "inspect_check_delivery":
                    ready, release = Event(), Event()

                    def completed_after_close(*args, **kwargs):
                        report = check_column(*args, **kwargs)
                        ready.set()
                        assert release.wait(5)
                        return report

                    with patch(
                        "loan_tape.column_check_ui.check_column", side_effect=completed_after_close
                    ):
                        child.run_button.invoke()
                        assert ready.wait(5)
                        inspector.close()
                        release.set()
                        child.worker.join(5)
                        assert not child.worker.is_alive() and child.report is None
                    clean_results()
                    # Also close a completed result before the Tk polling callback collects it.
                    app.inspect_button.invoke()
                    wait_for_preview(app)
                    inspector = app.inspector
                    inspector._select_column(0)
                    inspector.check_button.invoke()
                    child = inspector.column_checker
                    child.window.withdraw()
                    child.expected.set("Number")
                    child.run_button.invoke()
                    child.worker.join(5)
                    assert not child.worker.is_alive() and child.report is None
                    child.close()
                    clean_results()
                elif scenario == "inspect_check_error":
                    with patch(
                        "loan_tape.column_check_ui.check_column",
                        side_effect=InspectionError("Cannot read column"),
                    ):
                        child.run_button.invoke()
                        wait_check()
                    assert child.report is None
                    assert "No complete column check" in child.status.get()
                    assert not child.grid.get_children()
                    assert child.jump_button.instate(["disabled"])
                    child.run_button.invoke()
                    wait_check()
                    assert child.report.finding_count == 498
                    child.close()
                    clean_results()
                else:
                    child.run_button.invoke()
                    wait_check()
                    assert child.report.profile.table == full
                    assert child.report.profile.session is app.session
                    assert child.report.profile.total_rows == 500
                    assert child.report.finding_count == 498
                    assert "checked all 500" in child.status.get()
                    assert len(child.grid.get_children()) == 20
                    assert child.previous.instate(["disabled"])
                    assert child.next.instate(["!disabled"])
                    child.next.invoke()
                    assert child.offset == 20
                    child.previous.invoke()
                    assert child.offset == 0
                    child.show_page(480)
                    assert len(child.grid.get_children()) == 18
                    assert child.grid.get_children()[-1] == "5499"
                    assert child.next.instate(["disabled"])
                    first_report = child.report
                    child.expected.set("Date")
                    assert child.report is None and first_report._closed
                    child.min_year.set("not a year")
                    child.run_button.invoke()
                    assert not child.busy and "whole number" in child.status.get()
                    child.min_year.set("2000")
                    child.max_year.set("2030")
                    child.run_button.invoke()
                    wait_check()
                    assert child.report.finding_count == 500
                    assert child.report.rules.max_year == 2030
                    app.root.geometry("980x740")
                    app.root.deiconify()
                    child.window.geometry("800x660")
                    child.window.deiconify()
                    root.update()
                    assert child.grid.winfo_height() >= 60
                    assert child.jump_button.winfo_ismapped()
                    assert (
                        inspector.check_button.winfo_rootx() + inspector.check_button.winfo_width()
                        < app.root.winfo_rootx() + app.root.winfo_width()
                    )
                    child.window.withdraw()
                    child.expected.set("Identifier")
                    assert child.report is None and not child.grid.get_children()
                    child.run_button.invoke()
                    wait_check()
                    assert child.report.finding_count == 4
                    assert child.report.rules.max_year is None
                    assert set(child.grid.get_children()) == {"5005", "5200", "5400", "5500"}
                    child.grid.selection_set("5400")
                    child._selected()
                    assert "non-text" in child.detail.get("1.0", "end")
                    child.jump_button.invoke()
                    wait_for_preview(app)
                    assert child.closed and child.report is None
                    assert inspector.table.selection() == ("5400",)
                    assert "DZ5400" in inspector.cell_text.get("1.0", "end")
                    clean_results()
                    nav.range_text.set("DZ5000:DZ5200")
                    assert inspector.check_button.instate(["disabled"])
                    inspector.check_column()
                    assert inspector.column_checker.closed
                    assert "Save your data set changes" in nav.message.get()
                    nav.range_text.set(full.selection.area.address)
                    for name, area, header, expected in [
                        ("Side", "EA5000:EA5002", "5000", 0),
                        ("Below", "DZ5600:DZ5602", "5600", 1),
                    ]:
                        nav.new_button.invoke()
                        wait_for_preview(app)
                        nav.range_text.set(area)
                        nav.table_name.set(name)
                        nav.header_text.set(header)
                        nav.use_button.invoke()
                        wait_for_preview(app)
                        inspector._select_column(0)
                        inspector.check_button.invoke()
                        child = inspector.column_checker
                        child.window.withdraw()
                        child.expected.set("Identifier")
                        child.run_button.invoke()
                        wait_check()
                        assert child.report.profile.table.name == name
                        assert child.report.profile.total_rows == 2
                        assert child.report.finding_count == expected
                        if name == "Below":
                            child.grid.selection_set("5602")
                            child._selected()
                            child.jump_button.invoke()
                            wait_for_preview(app)
                            assert inspector.table.selection() == ("5602",)
                        else:
                            assert "No findings under these settings" in child.page_label.get()
                            child.close()
                    clean_results()
                assert saved_copy.read_bytes() == source_bytes
                assert not source.exists()
            elif scenario in {"inspect_column", "inspect_column_cancel", "inspect_column_error"}:
                app.inspect_button.invoke()
                wait_for_preview(app)
                inspector = app.inspector
                nav = inspector.navigator
                original = Path(app.saved_path.get()).read_bytes()
                inspector._select_column(0)
                assert inspector.column_button.instate(["disabled"])
                inspector.inspect_column()
                assert inspector.column_inspector is None
                assert "Save this data set" in nav.message.get()
                nav.range_text.set("DZ5000:DZ5500")
                nav.header_text.set("5000")
                nav.table_name.set("Main table")
                nav.use_button.invoke()
                wait_for_preview(app)
                full = nav.table_for_inspection()
                # Exercise the registered column-heading command, not a guessed column.
                inspector.table.tk.call(inspector.table.heading("DZ", "command"))
                assert inspector.column_button.cget("text") == "Inspect column DZ"
                if scenario == "inspect_column_cancel":
                    started, stopped = Event(), Event()

                    def paused_column(*args, cancelled, **kwargs):
                        started.set()
                        assert cancelled.wait(5)
                        stopped.set()
                        raise ColumnCancelled("Cancelled")

                    with patch("loan_tape.column_ui.inspect_column", side_effect=paused_column):
                        inspector.column_button.invoke()
                        assert started.wait(5)
                        child = inspector.column_inspector
                        inspector.close()
                        assert stopped.wait(5)
                        assert child.closed and child.profile is None
                        assert child.cancelled.is_set()
                        child.worker.join(5)
                        assert not child.worker.is_alive()
                elif scenario == "inspect_column_error":
                    with patch(
                        "loan_tape.column_ui.inspect_column",
                        side_effect=InspectionError("Cannot read column"),
                    ):
                        inspector.column_button.invoke()
                        wait_for_column(app)
                        child = inspector.column_inspector
                        assert child.profile is None
                        assert "No complete column summary" in child.status.get()
                        assert not child.examples.get_children()
                        assert child.jump_button.instate(["disabled"])
                        child.close()
                else:
                    inspector.column_button.invoke()
                    wait_for_column(app)
                    child = inspector.column_inspector
                    profile = child.profile
                    assert profile.total_rows == 500 and profile.filled_cells == 499
                    assert dict(profile.counts)["Blank"] == 1
                    assert profile.numeric_zeroes == 2
                    assert profile.distinct_values == 4 and profile.extra_occurrences == 495
                    assert profile.session is app.session
                    assert profile.table == full
                    assert "inspected all 500" in child.status.get()
                    assert app.session.date_label in child.scope_label.cget("text")
                    app.root.deiconify()
                    child.window.geometry("720x560")
                    child.window.deiconify()
                    root.update()
                    assert child.examples.winfo_height() >= 60
                    assert child.jump_button.winfo_ismapped()
                    child.window.withdraw()
                    item = next(
                        key
                        for key, value in child.entries.items()
                        if key.startswith("example") and value.row == 5400
                    )
                    child.examples.selection_set(item)
                    child._selected()
                    child.jump_button.invoke()
                    assert child.closed
                    wait_for_preview(app)
                    assert inspector.preview.start_row == 5400
                    assert inspector.table.selection() == ("5400",)
                    assert "DZ5400" in inspector.cell_text.get("1.0", "end")
                    assert nav.table_for_inspection() == full
                    # An unsaved edit must never quietly analyze the old saved range.
                    nav.range_text.set("DZ5000:DZ5200")
                    assert inspector.column_button.instate(["disabled"])
                    inspector.inspect_column()
                    assert inspector.column_inspector.closed
                    assert "Save your data set changes" in nav.message.get()
                    nav.range_text.set(full.selection.area.address)
                    for name, area, header, expected in [
                        ("Side", "EA5000:EA5002", "5000", 2),
                        ("Below", "DZ5600:DZ5602", "5600", 2),
                    ]:
                        nav.new_button.invoke()
                        wait_for_preview(app)
                        nav.table_name.set(name)
                        nav.range_text.set(area)
                        nav.header_text.set(header)
                        nav.use_button.invoke()
                        wait_for_preview(app)
                        inspector._select_column(0)
                        inspector.column_button.invoke()
                        wait_for_column(app)
                        child = inspector.column_inspector
                        assert child.profile.table.name == name
                        assert child.profile.total_rows == expected
                        assert child.profile.repeated_values == 0
                        if name == "Below":
                            item = next(
                                key for key, value in child.entries.items() if value.kind == "Blank"
                            )
                            child.examples.selection_set(item)
                            child._selected()
                            child.jump_button.invoke()
                            wait_for_preview(app)
                            assert inspector.table.selection() == ("5602",)
                            assert inspector.table.item("5602", "values") == ("",)
                        else:
                            child.close()
                    assert len(nav.tables) == 3 and nav.tables[0] == full
                    assert Path(app.saved_path.get()).read_bytes() == original
            elif scenario == "inspect_close_busy":
                release = Event()
                done = Event()

                def paused_read(*args, **kwargs):
                    assert release.wait(5)
                    try:
                        return read_parallel_preview(*args, **kwargs)
                    finally:
                        done.set()

                with patch("loan_tape.preview_ui.read_parallel_preview", side_effect=paused_read):
                    app.inspect_button.invoke()
                    app.inspector.close()
                    release.set()
                    assert done.wait(5)
                    root.update()
                    assert app.inspector.closed
                    assert app.inspector.cancelled.is_set()
            elif scenario == "inspect_comparison":
                primary = read_preview(Path(app.saved_path.get()))
                rows = list(primary.rows)
                rows[1] = (PreviewCell("123"), *rows[1][1:])
                secondary = replace(primary, rows=tuple(rows))
                compared = compare_previews(primary, secondary)
                with patch("loan_tape.preview_ui.read_parallel_preview", return_value=compared):
                    app.inspect_button.invoke()
                    wait_for_preview(app)
                inspector = app.inspector
                assert inspector.reader_selector.instate(["readonly"])
                assert "1 cell value differences" in inspector.comparison_note.get()
                assert "difference" in inspector.table.item("2", "tags")
                assert inspector.table.item("2", "values")[0] == "000123"
                inspector.reader_name.set("Excel (xlwings)")
                inspector.reader_selector.event_generate("<<ComboboxSelected>>")
                root.update()
                assert inspector.table.item("2", "values")[0] == "123"
                inspector.table.selection_set("2")
                inspector._selection()
                detail = inspector.cell_text.get("1.0", "end")
                assert "openpyxl: 000123" in detail and "Excel (xlwings): 123" in detail
                inspector.reader_name.set("openpyxl")
                inspector.reader_selector.event_generate("<<ComboboxSelected>>")
                root.update()
                assert inspector.table.item("2", "values")[0] == "000123"
            else:
                app.inspect_button.invoke()
                wait_for_preview(app)
                inspector = app.inspector
                assert inspector is not None
                if scenario in {"inspect_csv", "inspect_excel"}:
                    with patch("loan_tape.preview_ui.open_in_excel") as launch:
                        inspector.open_excel_button.invoke()
                        launch.assert_called_once_with(Path(app.saved_path.get()))
                    with (
                        patch(
                            "loan_tape.preview_ui.open_in_excel",
                            side_effect=OSError("Excel unavailable"),
                        ),
                        patch("loan_tape.preview_ui.messagebox.showerror") as error_dialog,
                    ):
                        inspector.open_excel_button.invoke()
                        error_dialog.assert_called_once_with(
                            "Could not open Excel", "Excel unavailable", parent=inspector.frame
                        )
                    app.root.geometry("980x740" if inspector.navigator else "760x520")
                    app.root.deiconify()
                    root.update()
                    assert inspector.open_excel_button.winfo_ismapped()
                    assert (
                        inspector.open_excel_button.winfo_rootx()
                        + inspector.open_excel_button.winfo_width()
                        <= app.root.winfo_rootx() + app.root.winfo_width()
                    )
                    app.root.withdraw()
                if scenario == "inspect_error":
                    assert inspector.preview is None
                    assert "Could not read this file" in inspector.status.get()
                    assert not inspector.table.get_children()
                elif scenario == "inspect_csv":
                    assert inspector.table.item("2", "values") == ("000123", "0")
                    assert inspector.table.item("1", "values") == ("loan_id", "balance")
                    inspector.table.selection_set("2")
                    inspector._selection()
                    assert "000123" in inspector.cell_text.get("1.0", "end")
                    assert str(inspector.cell_text.cget("state")) == "disabled"
                    app.inspect_button.invoke()
                    assert app.inspector is inspector  # Reuse the same file's preview.
                    inspector.encoding.set("UTF-8")
                    inspector.selectors[1].event_generate("<<ComboboxSelected>>")
                    wait_for_preview(app)
                    assert inspector.preview.rows[1][0].text == "000123"
                elif scenario == "inspect_excel":
                    assert inspector.reader_selector.instate(["disabled"])
                    assert (
                        "skipped" in inspector.comparison_note.get()
                        or "unavailable" in inspector.comparison_note.get()
                    )
                    assert inspector.table.item("2", "values")[4] == "1"
                    inspector.active_column = 4
                    inspector.table.selection_set("2")
                    inspector._selection()
                    assert "=B2+1" in inspector.cell_text.get("1.0", "end")
                    inspector.sheet_name.set("Another sheet")
                    inspector.sheet_selector.event_generate("<<ComboboxSelected>>")
                    wait_for_preview(app)
                    assert inspector.table.item("4", "values")[0] == "source location"
                    assert inspector.table.heading("C", "text") == "C"
                    inspector.sheet_name.set("Empty")
                    inspector.sheet_selector.event_generate("<<ComboboxSelected>>")
                    wait_for_preview(app)
                    assert inspector.status.get() == "No rows to display."
            app.close()
            assert app.closed and app.inspector.closed
        elif scenario == "legacy":
            assert app.original_path.get() == "Not recorded for this earlier copy"
            assert app.original_button.instate(["disabled"])
            assert not app.saved_button.instate(["disabled"])
        else:
            raise AssertionError(f"Unknown scenario: {scenario}")
        assert not callback_errors, callback_errors
    finally:
        if not app.closed:
            wait_for_copy(app)
            app.close()


if __name__ == "__main__":
    run(sys.argv[1], Path(sys.argv[2]))
