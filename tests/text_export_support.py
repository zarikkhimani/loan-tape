"""Synthetic mixed-cell source and real reviewed text/date results for export checks."""

from datetime import date
from pathlib import Path

from openpyxl import Workbook
from recipe_support import approval

from loan_tape.column_definitions import ColumnDefinition, save_column_definition, snapshot_field
from loan_tape.date_analysis import DateBinding, analyze_dates
from loan_tape.date_parser import DateProfile
from loan_tape.date_reader import read_excel_date_columns
from loan_tape.date_runs import save_date_run
from loan_tape.intake import IntakeStore
from loan_tape.packs import PackStore
from loan_tape.ranges import parse_range
from loan_tape.selection import SavedTable, SourceSelection, TableSelections, save_tables
from loan_tape.text_cleanup import TextCleanupSettings, review_excel_text
from loan_tape.text_rule_types import ExactRule, ExactRulesSettings
from loan_tape.transformations import apply_transformation


def prepare_export(work):
    original = work / "fictional-export.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Loans"
    sheet["B2"] = "Fictional portfolio"
    for row_number, values in enumerate(
        [
            ["Borrower", "ID", "Closing", "Amount", "Literal", "Formula", "Error"],
            ["CafÃ© Holdings", "00000000000000000001", "10/18/2026", 0, "=1+1", "=1+1", "#N/A"],
            ["CafÃ© Holdings", "0002", "ambiguous", 1.125, "+CMD", None, None],
            ["Café Group", "0003", "10/19/2026", 100, "'quote", None, None],
            [" Acme ", "0004", None, -12, "@SUM(A1)", None, None],
        ],
        3,
    ):
        for column, value in enumerate(values, 2):
            sheet.cell(row_number, column, value)
    sheet["F4"].data_type = "s"
    sheet.row_dimensions[5].hidden = True
    book.save(original)
    book.close()
    intake = IntakeStore(work / "inputs")
    record = intake.save_path(original)
    source = Path(intake.describe(record)["saved_path"])
    table = SavedTable(
        "1" * 32,
        "Fictional export",
        SourceSelection(record["sha256"], "Loans", parse_range("B2:H8"), 3),
    )
    save_tables(TableSelections(table.selection.workbook_sha256, (table,), table.id))
    store = PackStore(Path(__file__).parents[1] / "packs", work / ".artifacts/packs")
    for name in ("loan.text", "example.loans", "loan.dates"):
        store.activate(name)
    catalog = store.catalog()
    for column, header, field in (
        (2, "Borrower", "loan.text:borrower_name"),
        (3, "ID", "example.loans:loan_id"),
        (4, "Closing", "loan.dates:closing_date"),
    ):
        save_column_definition(
            table,
            ColumnDefinition(
                column,
                header,
                snapshot_field(catalog, catalog.get_field(field)),
                date_formats=("M/D/YYYY",) if column == 4 else (),
            ),
        )
    first = review_excel_text(source, table, 2, TextCleanupSettings(), store)
    parent = apply_transformation(first.snapshot, first.preview, approval(first, ((4, 2),)))
    exact = ExactRulesSettings((ExactRule("Café Holdings", "=Reviewed Name"),))
    second = review_excel_text(source, table, 2, exact, store, parent=parent)
    revision = apply_transformation(
        second.snapshot, second.preview, approval(second, ((4, 2),)), parent=parent
    )
    columns = read_excel_date_columns(source, table.selection, (4,))
    analysis = analyze_dates(
        table,
        columns,
        (DateBinding(4, "loan.dates:closing_date", DateProfile(text_formats=("M/D/YYYY",))),),
        catalog,
        reference_date=date(2026, 10, 18),
    )
    run = save_date_run(analysis, work / ".artifacts/date-runs")
    return original, source, intake, table, store, revision, run


def write_fixture(path, payload):
    """Synthetic output fixture only; application export always uses xlwings."""
    book = Workbook()
    book.remove(book.active)
    for spec in payload["sheets"]:
        sheet = book.create_sheet(spec["name"])
        width = max(map(len, spec["rows"]))
        for r, row in enumerate(spec["rows"], 1):
            for c in range(1, width + 1):
                cell = sheet.cell(r, c)
                cell.value = row[c - 1] if c <= len(row) else None
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                cell.number_format = "@"  # retain selected blank tail dimensions
    book.save(path)
    book.close()
